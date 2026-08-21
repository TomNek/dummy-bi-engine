from __future__ import annotations

import base64 as _base64
import csv as _csv
import datetime as _dt
import gzip as _gzip
import html as _html
import json as _json
import math as _math
import random as _random
import re
import struct as _struct
import difflib as _difflib
import urllib.parse as _urlparse
import uuid as _uuid
import xml.etree.ElementTree as _et
import zlib as _zlib
from dataclasses import dataclass, field
from collections import OrderedDict
from pathlib import Path
from typing import Any, Mapping, Sequence

from .locale import date_parse_formats, datetime_parse_formats, month_names, normalize_culture, number_separators, weekday_names
from .model import MDiagnostic
from .parser import parse_m_query


@dataclass(frozen=True)
class PreviewEvaluationResult:
    columns: list[dict[str, Any]]
    rows: list[list[Any]]
    applied_steps: list[str] = field(default_factory=list)
    blocked_steps: list[dict[str, str]] = field(default_factory=list)
    diagnostics: list[MDiagnostic] = field(default_factory=list)
    result_step_id: str | None = None

    def to_preview_dict(self, *, source_sql: str = "") -> dict[str, Any]:
        data = {
            "sql": source_sql,
            "columns": self.columns,
            "rows": [[_preview_cell(value) for value in row] for row in self.rows],
            "row_count": len(self.rows),
            "applied_steps": self.applied_steps,
            "blocked_steps": self.blocked_steps,
            "diagnostics": [diag.to_dict() for diag in self.diagnostics],
            "execution_scope": "mapped_source_plus_supported_steps" if self.applied_steps else "mapped_source",
        }
        if self.result_step_id:
            data["result_step_id"] = self.result_step_id
        return data


@dataclass(frozen=True)
class MErrorValue:
    message: str


@dataclass(frozen=True)
class MFunctionValue:
    parameters: list[str]
    body: str


@dataclass(frozen=True)
class MQueryValue:
    raw_m: str


@dataclass(frozen=True)
class MBuiltinFunctionValue:
    name: str
    options: dict[str, Any] = field(default_factory=dict)


@dataclass(frozen=True)
class MMetadataValue:
    value: Any
    metadata: dict[str, Any]


class MExpressionError(Exception):
    def __init__(self, error_record: Mapping[str, Any] | str):
        if isinstance(error_record, Mapping):
            self.error_record = dict(error_record)
        else:
            self.error_record = _error_record("Expression.Error", str(error_record), None)
        super().__init__(str(self.error_record.get("Message") or self.error_record.get("Reason") or "Expression error"))


_STATIC_FILE_BASE_PATH_KEY = "__power_query_static_file_base_path__"

_STATIC_M_CONSTANTS: dict[str, Any] = {
    "ACCESSCONTROLKIND.TYPE": "type accesscontrolkind",
    "BINARYENCODING.TYPE": "type binaryencoding",
    "BINARYOCCURRENCE.TYPE": "type binaryoccurrence",
    "BUFFERMODE.TYPE": "type buffermode",
    "BYTEORDER.TYPE": "type byteorder",
    "COMPRESSION.TYPE": "type compression",
    "CSVSTYLE.TYPE": "type csvstyle",
    "DAY.TYPE": "type day",
    "EXTRAVALUES.TYPE": "type extravalues",
    "GROUPKIND.TYPE": "type groupkind",
    "ITEMEXPRESSION.ITEM": "ItemExpression.Item",
    "JOINALGORITHM.TYPE": "type joinalgorithm",
    "JOINKIND.TYPE": "type joinkind",
    "JOINSIDE.TYPE": "type joinside",
    "LIMITCLAUSEKIND.TYPE": "type limitclausekind",
    "MISSINGFIELD.TYPE": "type missingfield",
    "NUMBER.E": _math.e,
    "NUMBER.EPSILON": _math.ulp(1.0),
    "NUMBER.NAN": _math.nan,
    "NUMBER.NEGATIVEINFINITY": -_math.inf,
    "NUMBER.PI": _math.pi,
    "NUMBER.POSITIVEINFINITY": _math.inf,
    "ODataOMITVALUES.TYPE".upper(): "type odataomitvalues",
    "OCCURRENCE.TYPE": "type occurrence",
    "ORDER.TYPE": "type order",
    "PERCENTILEMODE.TYPE": "type percentilemode",
    "PRECISION.TYPE": "type precision",
    "QUOTESTYLE.TYPE": "type quotestyle",
    "RANKKIND.TYPE": "type rankkind",
    "RELATIVEPOSITION.TYPE": "type relativeposition",
    "ROUNDINGMODE.TYPE": "type roundingmode",
    "ROWEXPRESSION.ROW": "RowExpression.Row",
    "SAPBUSINESSEXECUTIONMODE.TYPE": "type sapbusinesswarehouseexecutionmode",
    "SAPBUSINESSWAREHOUSEEXECUTIONMODE.TYPE": "type sapbusinesswarehouseexecutionmode",
    "SAPHANADISTRIBUTION.TYPE": "type saphanadistribution",
    "SAPHANARANGEOPERATOR.TYPE": "type saphanarangeoperator",
    "TEXTENCODING.TYPE": "type textencoding",
    "TRACELEVEL.TYPE": "type tracelevel",
    "WEBMETHOD.TYPE": "type webmethod",
}


def _preview_cell(value: Any) -> Any:
    value = _unwrap_metadata(value)
    if isinstance(value, MErrorValue):
        return {"error": value.message}
    if isinstance(value, (_dt.date, _dt.datetime, _dt.time)):
        return value.isoformat()
    if isinstance(value, _dt.timedelta):
        return str(value)
    if isinstance(value, (bytes, bytearray)):
        raw = bytes(value)
        return {"binary": _base64.b64encode(raw).decode("ascii"), "size": len(raw)}
    return value


def _unwrap_metadata(value: Any) -> Any:
    return value.value if isinstance(value, MMetadataValue) else value


def _metadata_record(value: Any) -> dict[str, Any]:
    return dict(value.metadata) if isinstance(value, MMetadataValue) else {}


def _m_constant_value(name: str) -> Any:
    upper = name.strip().upper()
    if upper == "CULTURE.CURRENT":
        return "en-US"
    if upper == "TIMEZONE.CURRENT":
        return _dt.datetime.now().astimezone().utcoffset() or _dt.timedelta()
    if upper in _STATIC_M_CONSTANTS:
        return _STATIC_M_CONSTANTS[upper]
    return None


def _table_key_metadata(columns: Sequence[Mapping[str, Any]]) -> list[dict[str, Any]]:
    for col in columns:
        keys = col.get("__table_keys")
        if isinstance(keys, list):
            return [dict(item) for item in keys if isinstance(item, Mapping)]
    return []


def _with_table_key_metadata(
    columns: Sequence[Mapping[str, Any]],
    keys: Sequence[Mapping[str, Any]],
) -> list[dict[str, Any]]:
    out = [dict(col) for col in columns]
    if not out:
        out.append({"name": "Value"})
    out[0]["__table_keys"] = [dict(key) for key in keys]
    return out


def _table_key_records_from_columns(columns: Sequence[Mapping[str, Any]]) -> list[dict[str, Any]]:
    return [
        {
            "Columns": list(key.get("Columns") or key.get("columns") or []),
            "Primary": bool(key.get("Primary") if "Primary" in key else key.get("primary")),
        }
        for key in _table_key_metadata(columns)
    ]


def _table_key_records_from_args(columns_value: Any, primary_value: Any = None) -> list[dict[str, Any]]:
    key_columns = [str(item) for item in columns_value] if isinstance(columns_value, list) else []
    return [{"Columns": key_columns, "Primary": bool(primary_value)}] if key_columns else []


def _table_env_lookup(
    env: Mapping[str, tuple[list[dict[str, Any]], list[list[Any]]]],
    expression: str,
) -> tuple[list[dict[str, Any]], list[list[Any]]] | None:
    name = str(expression or "").strip()
    if not name:
        return None
    candidates = [name, _m_identifier_name(name)]
    for candidate in candidates:
        if candidate in env:
            return env[candidate]
    upper_candidates = {candidate.upper() for candidate in candidates}
    for key, value in env.items():
        if key.upper() in upper_candidates:
            return value
    return None


def _binary_content_type(value: Any) -> dict[str, Any]:
    raw = bytes(_unwrap_metadata(value) or b"")
    stripped = raw.lstrip()
    content_type = "application/octet-stream"
    if raw.startswith(b"\x50\x4b\x03\x04"):
        content_type = "application/zip"
    elif raw.startswith(b"\x1f\x8b"):
        content_type = "application/gzip"
    elif raw.startswith(b"\xff\xd8"):
        content_type = "image/jpeg"
    elif raw.startswith(b"\x89PNG\r\n\x1a\n"):
        content_type = "image/png"
    elif raw.startswith(b"%PDF"):
        content_type = "application/pdf"
    elif stripped.startswith((b"{", b"[")):
        content_type = "application/json"
    elif stripped.startswith((b"<?xml", b"<")):
        content_type = "application/xml"
    else:
        try:
            raw.decode("utf-8")
            content_type = "text/plain"
        except Exception:
            pass
    return {"Content.Type": content_type, "Content.Length": len(raw)}


def _binary_format_with_options(fmt: Any, **options: Any) -> MBuiltinFunctionValue:
    if isinstance(fmt, MBuiltinFunctionValue):
        merged = dict(fmt.options)
        merged.update(options)
        return MBuiltinFunctionValue(fmt.name, merged)
    return MBuiltinFunctionValue(str(fmt), options)


def _binary_format_endian(fmt: MBuiltinFunctionValue) -> str:
    text = str(fmt.options.get("byte_order") or fmt.options.get("endian") or "").upper()
    if "BIG" in text:
        return ">"
    return "<"


def _binary_format_read_int(raw: bytes, offset: int, size: int, *, signed: bool, endian: str) -> tuple[int, int]:
    chunk = raw[offset : offset + size]
    if len(chunk) < size:
        raise ValueError("BinaryFormat preview needs more bytes.")
    return int.from_bytes(chunk, "big" if endian == ">" else "little", signed=signed), offset + size


def _binary_format_read_7bit(raw: bytes, offset: int, *, signed: bool) -> tuple[int, int]:
    result = 0
    shift = 0
    while True:
        if offset >= len(raw):
            raise ValueError("BinaryFormat 7-bit integer preview needs more bytes.")
        byte = raw[offset]
        offset += 1
        result |= (byte & 0x7F) << shift
        if not (byte & 0x80):
            break
        shift += 7
    if signed:
        result = (result >> 1) ^ -(result & 1)
    return result, offset


def _binary_format_parse(fmt: Any, raw: bytes, offset: int = 0) -> tuple[Any, int]:
    if not isinstance(fmt, MBuiltinFunctionValue):
        raise ValueError("BinaryFormat preview needs a BinaryFormat function value.")
    name = fmt.name.upper()
    endian = _binary_format_endian(fmt)
    if name == "BINARYFORMAT.NULL":
        return None, offset
    if name == "BINARYFORMAT.BYTE":
        return _binary_format_read_int(raw, offset, 1, signed=False, endian=endian)
    if name == "BINARYFORMAT.SIGNEDINTEGER16":
        return _binary_format_read_int(raw, offset, 2, signed=True, endian=endian)
    if name == "BINARYFORMAT.UNSIGNEDINTEGER16":
        return _binary_format_read_int(raw, offset, 2, signed=False, endian=endian)
    if name == "BINARYFORMAT.SIGNEDINTEGER32":
        return _binary_format_read_int(raw, offset, 4, signed=True, endian=endian)
    if name == "BINARYFORMAT.UNSIGNEDINTEGER32":
        return _binary_format_read_int(raw, offset, 4, signed=False, endian=endian)
    if name == "BINARYFORMAT.SIGNEDINTEGER64":
        return _binary_format_read_int(raw, offset, 8, signed=True, endian=endian)
    if name == "BINARYFORMAT.UNSIGNEDINTEGER64":
        return _binary_format_read_int(raw, offset, 8, signed=False, endian=endian)
    if name == "BINARYFORMAT.7BITENCODEDSIGNEDINTEGER":
        return _binary_format_read_7bit(raw, offset, signed=True)
    if name == "BINARYFORMAT.7BITENCODEDUNSIGNEDINTEGER":
        return _binary_format_read_7bit(raw, offset, signed=False)
    if name in {"BINARYFORMAT.SINGLE", "BINARYFORMAT.DOUBLE"}:
        fmt_code = "f" if name.endswith("SINGLE") else "d"
        size = 4 if fmt_code == "f" else 8
        chunk = raw[offset : offset + size]
        if len(chunk) < size:
            raise ValueError("BinaryFormat floating-point preview needs more bytes.")
        return _struct.unpack(endian + fmt_code, chunk)[0], offset + size
    if name == "BINARYFORMAT.DECIMAL":
        value, new_offset = _binary_format_read_int(raw, offset, 16, signed=True, endian=endian)
        return value, new_offset
    if name == "BINARYFORMAT.BINARY":
        length = fmt.options.get("length")
        if length is None:
            return raw[offset:], len(raw)
        size = max(0, int(_to_number(length) or 0))
        return raw[offset : offset + size], min(len(raw), offset + size)
    if name == "BINARYFORMAT.TEXT":
        length = fmt.options.get("length")
        size = len(raw) - offset if length is None else max(0, int(_to_number(length) or 0))
        encoding = _text_encoding_name(fmt.options.get("encoding"))
        return raw[offset : offset + size].decode(encoding, errors="replace"), min(len(raw), offset + size)
    if name == "BINARYFORMAT.LENGTH":
        size = max(0, int(_to_number(fmt.options.get("length")) or 0))
        return _binary_format_parse(fmt.options.get("format"), raw[offset : offset + size], 0)[0], min(len(raw), offset + size)
    if name == "BINARYFORMAT.LIST":
        item_format = fmt.options.get("format")
        count = fmt.options.get("count")
        limit = max(0, int(_to_number(count) or 0)) if count is not None else None
        out: list[Any] = []
        current = offset
        while current < len(raw) and (limit is None or len(out) < limit):
            value, current = _binary_format_parse(item_format, raw, current)
            out.append(value)
        return out, current
    if name in {"BINARYFORMAT.GROUP", "BINARYFORMAT.RECORD"}:
        fields = fmt.options.get("fields")
        current = offset
        if isinstance(fields, Mapping):
            out_record: dict[str, Any] = {}
            for key, child_format in fields.items():
                out_record[str(key)], current = _binary_format_parse(child_format, raw, current)
            return out_record, current
        out_values: list[Any] = []
        for child_format in fields if isinstance(fields, list) else []:
            value, current = _binary_format_parse(child_format, raw, current)
            out_values.append(value)
        return out_values, current
    if name == "BINARYFORMAT.TRANSFORM":
        value, current = _binary_format_parse(fmt.options.get("format"), raw, offset)
        transform = fmt.options.get("transform")
        if isinstance(transform, (MFunctionValue, MBuiltinFunctionValue)):
            value = _invoke_m_function(transform, [value])
        return value, current
    if name == "BINARYFORMAT.CHOICE":
        selector, current = _binary_format_parse(fmt.options.get("format"), raw, offset)
        choices = fmt.options.get("choices")
        selected = _mapping_get_case_insensitive(choices, str(selector)) if isinstance(choices, Mapping) else None
        if selected is None and isinstance(choices, Mapping):
            selected = choices.get(selector)
        if isinstance(selected, MBuiltinFunctionValue):
            return _binary_format_parse(selected, raw, current)
        return selected, current
    raise ValueError(f"Unsupported BinaryFormat preview function: {fmt.name}")


def _error_record(
    reason: Any,
    message: Any = None,
    detail: Any = None,
    parameters: Any = None,
    error_code: Any = None,
) -> dict[str, Any]:
    message_value = None if message is None else str(message)
    return {
        "Reason": str(reason or "Expression.Error"),
        "Message": message_value,
        "Detail": detail,
        "Message.Format": message_value,
        "Message.Parameters": parameters,
        "ErrorCode": None if error_code is None else str(error_code),
    }


def _column_names(columns: Sequence[Mapping[str, Any]]) -> list[str]:
    return [str(col.get("name") or "").strip() for col in columns]


def _string_literals(text: str) -> list[str]:
    values: list[str] = []
    for match in re.finditer(r'"((?:[^"]|"")*)"', text):
        values.append(match.group(1).replace('""', '"'))
    return values


def _after_first_argument(expression: str) -> str:
    start = expression.find("(")
    if start < 0:
        return expression
    depth = 0
    in_string = False
    i = start + 1
    while i < len(expression):
        char = expression[i]
        if char == '"':
            if in_string and i + 1 < len(expression) and expression[i + 1] == '"':
                i += 2
                continue
            in_string = not in_string
        elif not in_string:
            if char in "([{":
                depth += 1
            elif char in ")]}":
                depth = max(0, depth - 1)
            elif char == "," and depth == 0:
                return expression[i + 1 :]
        i += 1
    return expression


def _string_literals_after_first_arg(expression: str) -> list[str]:
    return _string_literals(_after_first_argument(expression))


def _find_column_index(columns: Sequence[Mapping[str, Any]], name: str) -> int:
    target = name.strip().upper()
    for idx, col in enumerate(columns):
        if str(col.get("name") or "").strip().upper() == target:
            return idx
    raise KeyError(name)


def _find_column_index_or_none(columns: Sequence[Mapping[str, Any]], name: str) -> int | None:
    try:
        return _find_column_index(columns, name)
    except KeyError:
        return None


def _missing_field_mode(expression: str, *, default: str = "error") -> str:
    args = _function_args(expression)
    if len(args) < 3:
        return default
    value = args[2].strip().upper()
    if value in {"MISSINGFIELD.IGNORE", "IGNORE", "1"}:
        return "ignore"
    if value in {"MISSINGFIELD.USENULL", "USENULL", "USE_NULL", "2"}:
        return "use_null"
    return "error"


def _row_record(columns: Sequence[Mapping[str, Any]], row: Sequence[Any]) -> dict[str, Any]:
    return {
        str(col.get("name") or ""): row[idx] if idx < len(row) else None
        for idx, col in enumerate(columns)
    }


def _select_columns(
    columns: list[dict[str, Any]],
    rows: list[list[Any]],
    names: list[str],
    *,
    missing_field: str = "error",
) -> tuple[list[dict[str, Any]], list[list[Any]]]:
    selectors: list[tuple[int | None, str]] = []
    for name in names:
        idx = _find_column_index_or_none(columns, name)
        if idx is None:
            if missing_field == "ignore":
                continue
            if missing_field == "use_null":
                selectors.append((None, name))
                continue
            raise KeyError(name)
        selectors.append((idx, name))
    out_cols = [dict(columns[idx]) if idx is not None else {"name": name} for idx, name in selectors]
    out_rows = [
        [row[idx] if idx is not None and idx < len(row) else None for idx, _name in selectors]
        for row in rows
    ]
    return out_cols, out_rows


def _remove_columns(
    columns: list[dict[str, Any]],
    rows: list[list[Any]],
    names: list[str],
    *,
    missing_field: str = "error",
) -> tuple[list[dict[str, Any]], list[list[Any]]]:
    existing = {str(col.get("name") or "").strip().upper() for col in columns}
    missing = [name for name in names if name.strip().upper() not in existing]
    if missing and missing_field == "error":
        raise KeyError(missing[0])
    remove = {name.strip().upper() for name in names}
    keep = [idx for idx, col in enumerate(columns) if str(col.get("name") or "").strip().upper() not in remove]
    out_cols = [dict(columns[idx]) for idx in keep]
    out_rows = [[row[idx] if idx < len(row) else None for idx in keep] for row in rows]
    return out_cols, out_rows


def _rename_columns(
    columns: list[dict[str, Any]],
    rows: list[list[Any]],
    expression: str,
) -> tuple[list[dict[str, Any]], list[list[Any]]]:
    names = _string_literals_after_first_arg(expression)
    pairs = list(zip(names[0::2], names[1::2]))
    missing_field = _missing_field_mode(expression)
    rename = {old.upper(): new for old, new in pairs}
    out: list[dict[str, Any]] = []
    seen_old = {str(col.get("name") or "").strip().upper() for col in columns}
    for col in columns:
        new_col = dict(col)
        current = str(new_col.get("name") or "")
        if current.upper() in rename:
            new_col["name"] = rename[current.upper()]
        out.append(new_col)
    missing_pairs = [(old, new) for old, new in pairs if old.strip().upper() not in seen_old]
    if missing_pairs and missing_field == "error":
        raise KeyError(missing_pairs[0][0])
    if missing_pairs and missing_field == "use_null":
        out.extend({"name": new} for _old, new in missing_pairs)
        rows = [list(row) + [None for _old, _new in missing_pairs] for row in rows]
    return out, rows


def _reorder_columns(columns: list[dict[str, Any]], rows: list[list[Any]], expression: str) -> tuple[list[dict[str, Any]], list[list[Any]]]:
    names = _string_literals_after_first_arg(expression)
    missing_field = _missing_field_mode(expression)
    selected: list[tuple[int | None, str]] = []
    used_indexes: set[int] = set()
    for name in names:
        idx = _find_column_index_or_none(columns, name)
        if idx is None:
            if missing_field == "ignore":
                continue
            if missing_field == "use_null":
                selected.append((None, name))
                continue
            raise KeyError(name)
        used_indexes.add(idx)
        selected.append((idx, str(columns[idx].get("name") or name)))
    for idx, col in enumerate(columns):
        if idx not in used_indexes:
            selected.append((idx, str(col.get("name") or "")))
    out_cols = [dict(columns[idx]) if idx is not None else {"name": name} for idx, name in selected]
    out_rows = [
        [row[idx] if idx is not None and idx < len(row) else None for idx, _name in selected]
        for row in rows
    ]
    return out_cols, out_rows


def _duplicate_column(columns: list[dict[str, Any]], rows: list[list[Any]], expression: str) -> tuple[list[dict[str, Any]], list[list[Any]]]:
    names = _string_literals_after_first_arg(expression)
    if len(names) < 2:
        raise ValueError("Table.DuplicateColumn preview needs source and new column names.")
    idx = _find_column_index(columns, names[0])
    out_cols = [dict(col) for col in columns] + [{"name": names[1]}]
    out_rows = [list(row) + [row[idx] if idx < len(row) else None] for row in rows]
    return out_cols, out_rows


def _normalize_type_text(type_text: str) -> str:
    normalized = type_text.strip().lower()
    if normalized.startswith("type nullable "):
        normalized = normalized[len("type nullable ") :]
    elif normalized.startswith("type "):
        normalized = normalized[len("type ") :]
    if normalized.endswith(".type"):
        normalized = normalized[: -len(".type")]
    return normalized.strip()


def _m_type_text(value: Any) -> str:
    value = _unwrap_metadata(value)
    if value is None:
        return "type null"
    if isinstance(value, bool):
        return "type logical"
    if isinstance(value, (int, float)):
        return "type number"
    if isinstance(value, str):
        normalized = _normalize_type_text(value)
        if normalized in {
            "any",
            "anynonnull",
            "none",
            "null",
            "logical",
            "number",
            "text",
            "date",
            "datetime",
            "datetimezone",
            "time",
            "duration",
            "binary",
            "list",
            "record",
            "table",
            "type",
            "function",
        }:
            return "type type"
        return "type text"
    if isinstance(value, _dt.datetime):
        return "type datetimezone" if value.tzinfo is not None else "type datetime"
    if isinstance(value, _dt.date):
        return "type date"
    if isinstance(value, _dt.time):
        return "type time"
    if isinstance(value, _dt.timedelta):
        return "type duration"
    if isinstance(value, (bytes, bytearray)):
        return "type binary"
    if isinstance(value, list):
        return "type table" if all(isinstance(item, Mapping) for item in value) else "type list"
    if isinstance(value, Mapping):
        return "type record"
    if isinstance(value, (MFunctionValue, MBuiltinFunctionValue)):
        return "type function"
    return "type any"


def _m_kind_from_type(type_text: Any) -> str:
    normalized = _normalize_type_text(str(type_text or ""))
    mapping = {
        "logical": "Logical",
        "boolean": "Logical",
        "number": "Number",
        "double": "Number",
        "decimal": "Number",
        "currency": "Number",
        "int64": "Number",
        "int32": "Number",
        "integer": "Number",
        "whole number": "Number",
        "text": "Text",
        "string": "Text",
        "date": "Date",
        "datetime": "DateTime",
        "datetimezone": "DateTimeZone",
        "time": "Time",
        "duration": "Duration",
        "binary": "Binary",
        "record": "Record",
        "list": "List",
        "table": "Table",
        "function": "Function",
        "null": "Null",
    }
    return mapping.get(normalized, "Any")


def _m_type_name_from_kind(kind: str) -> str:
    if kind == "Any":
        return "Any.Type"
    if kind == "Null":
        return "Null.Type"
    return f"{kind}.Type"


def _column_type_text(columns: Sequence[Mapping[str, Any]], rows: Sequence[Sequence[Any]], idx: int) -> str:
    if idx < len(columns):
        type_text = str(columns[idx].get("m_type") or "").strip()
        if type_text:
            return type_text
    sample = next((row[idx] for row in rows if idx < len(row) and row[idx] is not None), None)
    return _m_type_text(sample) if sample is not None else "type any"


def _type_compatible(type1: Any, type2: Any) -> bool:
    left_raw = str(type1 or "").strip()
    right_raw = str(type2 or "").strip()
    left_nullable = "nullable" in left_raw.lower()
    right_nullable = "nullable" in right_raw.lower()
    left = _normalize_type_text(left_raw)
    right = _normalize_type_text(right_raw)
    if right == "any":
        return True
    if right == "anynonnull":
        return left != "null" and not left_nullable
    if right_nullable:
        return left == "null" or left == right
    if left_nullable and left == right:
        return False
    return left == right


def _type_record_fields(type_value: Any) -> dict[str, dict[str, Any]]:
    text = str(_unwrap_metadata(type_value) or "").strip()
    match = re.search(r"\[(.*)\]", text, flags=re.DOTALL)
    if not match:
        return {}
    out: dict[str, dict[str, Any]] = {}
    for raw_item in _split_top_level_items(match.group(1)):
        if "=" not in raw_item:
            continue
        name_part, type_part = raw_item.split("=", 1)
        name_part = name_part.strip()
        optional = False
        if name_part.lower().startswith("optional "):
            optional = True
            name_part = name_part[len("optional ") :].strip()
        name = _m_identifier_name(name_part)
        out[name] = {
            "Type": type_part.strip(),
            "Optional": optional,
        }
    return out


def _type_for_record(fields: Any, open_record: Any = None) -> MMetadataValue:
    field_map = fields if isinstance(fields, Mapping) else {}
    parts: list[str] = []
    for name, spec in field_map.items():
        spec_record = spec if isinstance(spec, Mapping) else {}
        optional = bool(spec_record.get("Optional") or spec_record.get("optional"))
        type_text = spec_record.get("Type") or spec_record.get("type") or "type any"
        prefix = "optional " if optional else ""
        parts.append(f"{prefix}{_m_identifier_text(str(name))} = {type_text}")
    text = "type [" + ", ".join(parts) + "]"
    return MMetadataValue(value=text, metadata={"Open": bool(open_record), "RecordFields": dict(field_map)})


def _type_for_function(parameters: Any, return_type: Any = None) -> MMetadataValue:
    params = parameters if isinstance(parameters, Mapping) else {}
    return MMetadataValue(
        value="type function",
        metadata={
            "Parameters": dict(params),
            "ReturnType": return_type or "type any",
        },
    )


def _type_table_schema(type_value: Any) -> list[dict[str, Any]]:
    return [
        {
            "Name": name,
            "Kind": _m_kind_from_type(str(spec.get("Type") or "type any")),
            "TypeName": _m_type_name_from_kind(_m_kind_from_type(str(spec.get("Type") or "type any"))),
            "Optional": bool(spec.get("Optional")),
        }
        for name, spec in _type_record_fields(type_value).items()
    ]


def _type_table_keys(type_value: Any) -> list[dict[str, Any]]:
    metadata = _metadata_record(type_value)
    keys = metadata.get("Keys")
    if isinstance(keys, list):
        return [dict(key) for key in keys if isinstance(key, Mapping)]
    return []


def _value_matches_type(value: Any, type_value: Any) -> bool:
    type_text = str(type_value or "").strip()
    normalized = _normalize_type_text(type_text)
    if normalized == "any":
        return True
    if normalized == "anynonnull":
        return value is not None
    if "nullable" in type_text.lower() and value is None:
        return True
    return _type_compatible(_m_type_text(value), type_text)


def _culture_key(value: Any) -> str:
    return normalize_culture(value)


def _coerce_value(value: Any, type_text: str, culture: Any = None) -> Any:
    if value is None or value == "":
        return None
    normalized = _normalize_type_text(type_text)
    try:
        if normalized in {"number", "double", "decimal", "currency"}:
            return _to_number(value, culture)
        if normalized in {"int64", "int32", "integer", "whole number"}:
            number = _to_number(value, culture)
            return None if number is None else int(number)
        if normalized in {"text", "string"}:
            return str(value)
        if normalized in {"date"}:
            return _to_date(value, culture)
        if normalized in {"datetime"}:
            return _to_datetime(value, culture)
        if normalized in {"datetimezone"}:
            return _to_datetimezone(value, culture)
        if normalized in {"time"}:
            return _to_time(value, culture)
        if normalized in {"logical", "boolean"}:
            return _to_logical(value)
    except Exception:
        return value
    return value


def _transform_column_types(columns: list[dict[str, Any]], rows: list[list[Any]], expression: str) -> tuple[list[dict[str, Any]], list[list[Any]]]:
    args = _function_args(expression)
    culture = _parse_m_value(args[2]) if len(args) > 2 else None
    spec_body = _between_outer(args[1], "{", "}") if len(args) >= 2 else None
    pairs: list[tuple[str, str]] = []
    if spec_body is not None:
        for raw_spec in _split_top_level_items(spec_body):
            body = _between_outer(raw_spec, "{", "}")
            if body is None:
                continue
            parts = _split_top_level_items(body)
            if len(parts) >= 2:
                pairs.append((str(_parse_m_value(parts[0])), parts[1].strip()))
    if not pairs:
        pairs = re.findall(r'\{\s*"([^"]+)"\s*,\s*(type\s+[A-Za-z0-9_ ]+|[A-Za-z0-9_]+\.Type)\s*\}', expression, flags=re.IGNORECASE)
    if not pairs:
        return columns, rows
    indexes = {name: _find_column_index(columns, name) for name, _type in pairs}
    types = {name: type_text for name, type_text in pairs}
    out_cols = [dict(col) for col in columns]
    for name, type_text in types.items():
        idx = indexes[name]
        out_cols[idx]["m_type"] = type_text.strip() if type_text.strip().lower().startswith("type ") else "type " + _normalize_type_text(type_text)
    out_rows: list[list[Any]] = []
    for row in rows:
        out_row = list(row)
        for name, idx in indexes.items():
            if idx < len(out_row):
                out_row[idx] = _coerce_value(out_row[idx], types[name], culture)
        out_rows.append(out_row)
    return out_cols, out_rows


def _apply_value_replace_type(columns: list[dict[str, Any]], rows: list[list[Any]], expression: str) -> tuple[list[dict[str, Any]], list[list[Any]]]:
    args = _function_args(expression)
    if len(args) < 2:
        return columns, rows
    match = re.search(r"\btype\s+table\s*\[(.*)\]\s*$", args[1].strip(), flags=re.IGNORECASE | re.DOTALL)
    if not match:
        return columns, rows
    type_by_column: dict[str, str] = {}
    for item in _split_top_level_items(match.group(1)):
        if "=" not in item:
            continue
        name, type_text = item.split("=", 1)
        column = name.strip().strip('"')
        if column:
            type_by_column[column.upper()] = type_text.strip()
    if not type_by_column:
        return columns, rows
    out_cols = []
    for col in columns:
        new_col = dict(col)
        current = str(new_col.get("name") or "").strip().upper()
        if current in type_by_column:
            type_text = type_by_column[current]
            new_col["m_type"] = type_text if type_text.lower().startswith("type ") else f"type {type_text}"
        out_cols.append(new_col)
    return out_cols, rows


def _parse_literal(raw: str) -> Any:
    text = raw.strip()
    if text.startswith('"') and text.endswith('"'):
        return _decode_m_text(text[1:-1].replace('""', '"'))
    if text.lower() == "null":
        return None
    if text.lower() == "true":
        return True
    if text.lower() == "false":
        return False
    try:
        if "." in text:
            return float(text)
        return int(text)
    except Exception:
        return text


def _decode_m_text(value: str) -> str:
    replacements = {
        "#(tab)": "\t",
        "#(lf)": "\n",
        "#(cr)": "\r",
        "#(cr,lf)": "\r\n",
    }
    out = value
    for token, replacement in replacements.items():
        out = out.replace(token, replacement)
    return out


def _split_top_level_items(text: str, delimiter: str = ",") -> list[str]:
    parts: list[str] = []
    start = 0
    depth = 0
    in_string = False
    i = 0
    while i < len(text):
        ch = text[i]
        nxt = text[i + 1] if i + 1 < len(text) else ""
        if in_string:
            if ch == '"' and nxt == '"':
                i += 2
                continue
            if ch == '"':
                in_string = False
            i += 1
            continue
        if ch == '"':
            in_string = True
            i += 1
            continue
        if ch in "([{":
            depth += 1
        elif ch in ")]}" and depth > 0:
            depth -= 1
        elif ch == delimiter and depth == 0:
            parts.append(text[start:i].strip())
            start = i + 1
        i += 1
    tail = text[start:].strip()
    if tail:
        parts.append(tail)
    return parts


def _between_outer(text: str, opener: str, closer: str) -> str | None:
    stripped = text.strip()
    if not stripped.startswith(opener) or not stripped.endswith(closer):
        return None
    return stripped[1:-1]


def _function_args(expression: str) -> list[str]:
    start = expression.find("(")
    end = expression.rfind(")")
    if start < 0 or end <= start:
        return []
    return _split_top_level_items(expression[start + 1 : end])


def _split_top_level_range(text: str) -> tuple[str, str] | None:
    depth = 0
    in_string = False
    i = 0
    while i < len(text) - 1:
        ch = text[i]
        nxt = text[i + 1]
        if in_string:
            if ch == '"' and nxt == '"':
                i += 2
                continue
            if ch == '"':
                in_string = False
            i += 1
            continue
        if ch == '"':
            in_string = True
            i += 1
            continue
        if ch in "([{":
            depth += 1
        elif ch in ")]}" and depth > 0:
            depth -= 1
        elif depth == 0 and ch == "." and nxt == ".":
            return text[:i].strip(), text[i + 2 :].strip()
        i += 1
    return None


def _m_range_values(left: Any, right: Any) -> list[Any]:
    if isinstance(left, bool) or isinstance(right, bool):
        raise ValueError("M range preview does not support logical values.")
    if isinstance(left, (int, float)) and isinstance(right, (int, float)) and float(left).is_integer() and float(right).is_integer():
        start = int(left)
        end = int(right)
        step = 1 if end >= start else -1
        return list(range(start, end + step, step))
    if isinstance(left, str) and isinstance(right, str) and len(left) == 1 and len(right) == 1:
        start = ord(left)
        end = ord(right)
        step = 1 if end >= start else -1
        return [chr(code) for code in range(start, end + step, step)]
    raise ValueError("M range preview supports integer and single-character ranges only.")


def _parse_m_range_values(text: str) -> list[Any] | None:
    split = _split_top_level_range(text)
    if split is None:
        return None
    left, right = split
    return _m_range_values(_parse_m_value(left), _parse_m_value(right))


def _top_level_arrow_index(text: str) -> int | None:
    depth = 0
    in_string = False
    i = 0
    while i < len(text) - 1:
        ch = text[i]
        nxt = text[i + 1]
        if in_string:
            if ch == '"' and nxt == '"':
                i += 2
                continue
            if ch == '"':
                in_string = False
            i += 1
            continue
        if ch == '"':
            in_string = True
        elif ch in "([{":
            depth += 1
        elif ch in ")]}" and depth > 0:
            depth -= 1
        elif ch == "=" and nxt == ">" and depth == 0:
            return i
        i += 1
    return None


def _function_parameter_names(text: str) -> list[str]:
    left = text.strip()
    if left.startswith("(") and ")" in left:
        params_text = left[1 : left.rfind(")")]
    else:
        params_text = re.sub(r"\s+as\s+.+$", "", left, flags=re.IGNORECASE | re.DOTALL)
    out: list[str] = []
    for raw_part in _split_top_level_items(params_text):
        part = re.sub(r"^optional\s+", "", raw_part.strip(), flags=re.IGNORECASE)
        if not part:
            continue
        name = re.split(r"\s+as\s+|\s+", part, maxsplit=1, flags=re.IGNORECASE)[0].strip()
        out.append(_m_identifier_name(name))
    return [name for name in out if name]


def _parse_m_function_value(raw_m: str) -> MFunctionValue | None:
    text = (raw_m or "").strip()
    arrow_idx = _top_level_arrow_index(text)
    if arrow_idx is None:
        return None
    params = _function_parameter_names(text[:arrow_idx])
    if not params and not re.fullmatch(r"\(\s*\)", text[:arrow_idx].strip()):
        return None
    return MFunctionValue(parameters=params, body=text[arrow_idx + 2 :].strip())


def _strip_top_level_meta(raw_m: str) -> str:
    split = _split_top_level_keyword(raw_m or "", "meta")
    return split[0].strip() if split else (raw_m or "").strip()


def _parse_m_value(raw: str) -> Any:
    text = raw.strip()
    binary_match = re.match(r"^#binary\s*\((.*)\)$", text, flags=re.IGNORECASE | re.DOTALL)
    if binary_match:
        args = _split_top_level_items(binary_match.group(1))
        values = _parse_m_value(args[0]) if args else []
        return _binary_from_list(values)
    ctor_match = re.match(r"^#(date|datetime|datetimezone|time|duration)\s*\((.*)\)$", text, flags=re.IGNORECASE | re.DOTALL)
    if ctor_match:
        values = [_parse_m_value(item) for item in _split_top_level_items(ctor_match.group(2))]
        kind = ctor_match.group(1).lower()
        if kind == "date" and len(values) >= 3:
            return _dt.date(int(values[0]), int(values[1]), int(values[2]))
        if kind == "datetime" and len(values) >= 6:
            return _dt.datetime(int(values[0]), int(values[1]), int(values[2]), int(values[3]), int(values[4]), int(values[5]))
        if kind == "datetimezone" and len(values) >= 8:
            hours = int(values[6])
            minutes = int(values[7])
            total_minutes = hours * 60 + (minutes if hours >= 0 else -abs(minutes))
            tz = _dt.timezone(_dt.timedelta(minutes=total_minutes))
            return _dt.datetime(int(values[0]), int(values[1]), int(values[2]), int(values[3]), int(values[4]), int(values[5]), tzinfo=tz)
        if kind == "time" and len(values) >= 3:
            return _dt.time(int(values[0]), int(values[1]), int(values[2]))
        if kind == "duration" and len(values) >= 4:
            return _dt.timedelta(days=int(values[0]), hours=int(values[1]), minutes=int(values[2]), seconds=float(values[3]))
    range_values = _parse_m_range_values(text)
    if range_values is not None:
        return range_values
    list_body = _between_outer(text, "{", "}")
    if list_body is not None:
        if not list_body.strip():
            return []
        values: list[Any] = []
        for item in _split_top_level_items(list_body):
            range_item = _parse_m_range_values(item)
            if range_item is not None:
                values.extend(range_item)
            else:
                values.append(_parse_m_value(item))
        return values
    record_body = _between_outer(text, "[", "]")
    if record_body is not None:
        record: dict[str, Any] = {}
        for item in _split_top_level_items(record_body):
            if "=" not in item:
                continue
            key, value = item.split("=", 1)
            record[key.strip().strip('"')] = _parse_m_value(value)
        return record
    return _parse_literal(text)


def _split_top_level_equals(text: str) -> tuple[str, str] | None:
    depth = 0
    in_string = False
    i = 0
    while i < len(text):
        ch = text[i]
        nxt = text[i + 1] if i + 1 < len(text) else ""
        if in_string:
            if ch == '"' and nxt == '"':
                i += 2
                continue
            if ch == '"':
                in_string = False
            i += 1
            continue
        if ch == '"':
            in_string = True
        elif ch in "([{":
            depth += 1
        elif ch in ")]}" and depth > 0:
            depth -= 1
        elif ch == "=" and depth == 0:
            return text[:i].strip(), text[i + 1 :].strip()
        i += 1
    return None


def _strip_outer_parens(text: str) -> str:
    stripped = text.strip()
    while stripped.startswith("(") and stripped.endswith(")"):
        depth = 0
        in_string = False
        balanced_outer = True
        for idx, ch in enumerate(stripped):
            nxt = stripped[idx + 1] if idx + 1 < len(stripped) else ""
            if in_string:
                if ch == '"' and nxt == '"':
                    continue
                if ch == '"':
                    in_string = False
                continue
            if ch == '"':
                in_string = True
            elif ch == "(":
                depth += 1
            elif ch == ")":
                depth -= 1
                if depth == 0 and idx != len(stripped) - 1:
                    balanced_outer = False
                    break
        if not balanced_outer:
            break
        stripped = stripped[1:-1].strip()
    return stripped


def _split_top_level_keyword(text: str, keyword: str) -> tuple[str, str] | None:
    pattern = keyword.lower()
    depth = 0
    in_string = False
    i = 0
    while i <= len(text) - len(keyword):
        ch = text[i]
        nxt = text[i + 1] if i + 1 < len(text) else ""
        if in_string:
            if ch == '"' and nxt == '"':
                i += 2
                continue
            if ch == '"':
                in_string = False
            i += 1
            continue
        if ch == '"':
            in_string = True
            i += 1
            continue
        if ch in "([{":
            depth += 1
        elif ch in ")]}" and depth > 0:
            depth -= 1
        elif depth == 0 and text[i : i + len(keyword)].lower() == pattern:
            before = text[i - 1] if i > 0 else " "
            after = text[i + len(keyword)] if i + len(keyword) < len(text) else " "
            if not (before.isalnum() or before == "_") and not (after.isalnum() or after == "_"):
                return text[:i].strip(), text[i + len(keyword) :].strip()
        i += 1
    return None


def _split_top_level_operator(text: str, operators: Sequence[str]) -> tuple[str, str, str] | None:
    depth = 0
    in_string = False
    i = len(text) - 1
    ordered = sorted(operators, key=len, reverse=True)
    while i >= 0:
        ch = text[i]
        if in_string:
            if ch == '"':
                if i > 0 and text[i - 1] == '"':
                    i -= 2
                    continue
                in_string = False
            i -= 1
            continue
        if ch == '"':
            in_string = True
            i -= 1
            continue
        if ch in ")]}":
            depth += 1
            i -= 1
            continue
        if ch in "([{":
            depth = max(0, depth - 1)
            i -= 1
            continue
        if depth == 0:
            for op in ordered:
                start = i - len(op) + 1
                if start < 0:
                    continue
                if text[start : i + 1] != op:
                    continue
                if op in {"+", "-"} and (start == 0 or text[start - 1] in "(,+-*/<>=&"):
                    continue
                return text[:start].strip(), op, text[i + 1 :].strip()
        i -= 1
    return None


def _to_number(value: Any, culture: Any = None) -> float | None:
    value = _unwrap_metadata(value)
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return 1.0 if value else 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    decimal, thousands = number_separators(culture)
    for separator in thousands:
        if separator and separator != decimal:
            text = text.replace(separator, "")
    if decimal != ".":
        text = text.replace(decimal, ".")
    try:
        return float(text)
    except Exception:
        raise ValueError(f"Cannot convert {value!r} to number")


def _number_text(value: Any, format_text: Any = None, culture: Any = None) -> str | None:
    number = _to_number(value)
    if number is None:
        return None
    fmt = str(format_text or "").strip()
    if not fmt:
        text = str(int(number)) if float(number).is_integer() else str(number)
    else:
        lower = fmt.lower()
        if lower.startswith("p"):
            digits = int(lower[1:] or "2") if lower[1:].isdigit() else 2
            text = f"{number * 100:.{digits}f} %"
        elif lower.startswith("e"):
            mantissa, exponent = f"{number:.6e}".split("e", 1)
            sign = "+" if int(exponent) >= 0 else "-"
            text = f"{mantissa}e{sign}{abs(int(exponent)):03d}"
        else:
            text = str(int(number)) if float(number).is_integer() else str(number)
    if _culture_key(culture).startswith(("de", "fr", "es", "it", "pt")):
        text = text.replace(".", ",")
    return text


def _math_number(value: Any) -> float | None:
    return _to_number(value)


def _int_number(value: Any) -> int:
    number = _to_number(value)
    return 0 if number is None else int(number)


def _to_date(value: Any, culture: Any = None) -> _dt.date | None:
    value = _unwrap_metadata(value)
    if isinstance(value, _dt.datetime):
        return value.date()
    if isinstance(value, _dt.date):
        return value
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    culture_formats = date_parse_formats(culture)
    for fmt in (*culture_formats, "%Y-%m-%d", "%Y/%m/%d", "%m/%d/%Y", "%d.%m.%Y", "%d/%m/%Y"):
        try:
            sample_len = len(_dt.datetime(2000, 12, 31).strftime(fmt))
            return _dt.datetime.strptime(text[:sample_len], fmt).date()
        except Exception:
            pass
    try:
        return _dt.datetime.fromisoformat(text.replace("Z", "+00:00")).date()
    except Exception:
        raise ValueError(f"Cannot convert {value!r} to date")


def _to_datetime(value: Any, culture: Any = None) -> _dt.datetime | None:
    value = _unwrap_metadata(value)
    if isinstance(value, _dt.datetime):
        return value
    if isinstance(value, _dt.date):
        return _dt.datetime.combine(value, _dt.time())
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    culture_formats = datetime_parse_formats(culture)
    for fmt in (*culture_formats, "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S"):
        try:
            sample_len = len(_dt.datetime(2000, 12, 31, 23, 59, 59).strftime(fmt))
            return _dt.datetime.strptime(text[:sample_len], fmt)
        except Exception:
            pass
    try:
        return _dt.datetime.fromisoformat(text.replace("Z", "+00:00")).replace(tzinfo=None)
    except Exception:
        raise ValueError(f"Cannot convert {value!r} to datetime")


def _to_time(value: Any, culture: Any = None) -> _dt.time | None:
    value = _unwrap_metadata(value)
    if isinstance(value, _dt.datetime):
        return value.timetz() if value.tzinfo is not None else value.time()
    if isinstance(value, _dt.time):
        return value
    if isinstance(value, _dt.date):
        return _dt.time()
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    if text.startswith("#time"):
        parsed = _parse_m_value(text)
        if isinstance(parsed, _dt.time):
            return parsed
    try:
        return _dt.time.fromisoformat(text)
    except Exception:
        parsed = _to_datetime(value, culture)
        return parsed.time() if parsed is not None else None


def _timezone_from_parts(hours: Any = 0, minutes: Any = 0) -> _dt.timezone:
    hour_value = int(_to_number(hours) or 0)
    minute_value = int(_to_number(minutes) or 0)
    total_minutes = hour_value * 60 + (minute_value if hour_value >= 0 else -abs(minute_value))
    return _dt.timezone(_dt.timedelta(minutes=total_minutes))


def _to_datetimezone(value: Any, culture: Any = None) -> _dt.datetime | None:
    value = _unwrap_metadata(value)
    if value is None or value == "":
        return None
    if isinstance(value, _dt.datetime):
        return value if value.tzinfo is not None else value.replace(tzinfo=_dt.timezone.utc)
    if isinstance(value, _dt.date):
        return _dt.datetime.combine(value, _dt.time(), tzinfo=_dt.timezone.utc)
    text = str(value).strip()
    if not text:
        return None
    if text.startswith("#datetimezone"):
        parsed = _parse_m_value(text)
        if isinstance(parsed, _dt.datetime):
            return parsed if parsed.tzinfo is not None else parsed.replace(tzinfo=_dt.timezone.utc)
    try:
        parsed = _dt.datetime.fromisoformat(text.replace("Z", "+00:00"))
        return parsed if parsed.tzinfo is not None else parsed.replace(tzinfo=_dt.timezone.utc)
    except Exception:
        parsed = _to_datetime(value, culture)
        return parsed.replace(tzinfo=_dt.timezone.utc) if parsed is not None else None


def _to_duration(value: Any) -> _dt.timedelta | None:
    value = _unwrap_metadata(value)
    if value is None or value == "":
        return None
    if isinstance(value, _dt.timedelta):
        return value
    if isinstance(value, (int, float)):
        return _dt.timedelta(days=float(value))
    text = str(value).strip()
    if not text:
        return None
    if text.startswith("#duration"):
        parsed = _parse_m_value(text)
        if isinstance(parsed, _dt.timedelta):
            return parsed
    parts = text.split(":")
    try:
        if len(parts) == 3:
            return _dt.timedelta(hours=int(parts[0]), minutes=int(parts[1]), seconds=float(parts[2]))
        return _dt.timedelta(days=float(text))
    except Exception:
        raise ValueError(f"Cannot convert {value!r} to duration")


def _to_logical(value: Any) -> bool | None:
    value = _unwrap_metadata(value)
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return value
    text = str(value).strip().lower()
    if text in {"true", "1", "yes", "y", "wahr"}:
        return True
    if text in {"false", "0", "no", "n", "falsch"}:
        return False
    raise ValueError(f"Cannot convert {value!r} to logical")


def _text_delimiter_index(value: str, delimiter: str, occurrence: int = 0) -> int:
    start = 0
    found = -1
    for _idx in range(max(0, occurrence) + 1):
        found = value.find(delimiter, start)
        if found < 0:
            return -1
        start = found + len(delimiter)
    return found


def _comparer_ignore_case(value: Any) -> bool:
    text = str(value or "").strip().upper()
    return text.endswith("ORDINALIGNORECASE") or text.endswith("FROMCULTUREIGNORECASE")


def _text_character_set(value: Any) -> set[str]:
    chars: list[str] = []

    def collect(item: Any) -> None:
        if item is None:
            return
        if isinstance(item, list):
            for child in item:
                collect(child)
            return
        chars.extend(str(item))

    collect(value)
    return set(chars)


def _trim_text(value: Any, trim_value: Any = None, *, side: str = "both") -> str | None:
    if value is None:
        return None
    text = str(value)
    if trim_value is None:
        if side == "start":
            return text.lstrip()
        if side == "end":
            return text.rstrip()
        return text.strip()
    chars = "".join(sorted(_text_character_set(trim_value)))
    if side == "start":
        return text.lstrip(chars)
    if side == "end":
        return text.rstrip(chars)
    return text.strip(chars)


_MONTH_NAMES_DE = [
    "Januar",
    "Februar",
    "März",
    "April",
    "Mai",
    "Juni",
    "Juli",
    "August",
    "September",
    "Oktober",
    "November",
    "Dezember",
]

_WEEKDAY_NAMES_DE = ["Montag", "Dienstag", "Mittwoch", "Donnerstag", "Freitag", "Samstag", "Sonntag"]


def _localized_month_names(culture: Any = None) -> list[str]:
    return month_names(culture)


def _localized_weekday_names(culture: Any = None) -> list[str]:
    return weekday_names(culture)


def _temporal_format(value: _dt.date | _dt.datetime, format_text: str | None, culture: Any = None) -> str:
    if not format_text:
        return value.isoformat()
    months = _localized_month_names(culture)
    weekdays = _localized_weekday_names(culture)

    def token_value(token: str) -> str:
        char = token[0]
        length = len(token)
        if char == "y":
            return f"{value.year % 100:02d}" if length == 2 else f"{value.year:04d}"
        if char == "M":
            if length >= 4:
                return months[value.month - 1]
            if length == 3:
                return months[value.month - 1][:3]
            return f"{value.month:02d}" if length == 2 else str(value.month)
        if char == "d":
            if length >= 4:
                return weekdays[value.weekday()]
            if length == 3:
                return weekdays[value.weekday()][:3]
            return f"{value.day:02d}" if length == 2 else str(value.day)
        if isinstance(value, _dt.datetime):
            if char in {"H", "h"}:
                return f"{value.hour:02d}" if length >= 2 else str(value.hour)
            if char == "m":
                return f"{value.minute:02d}" if length >= 2 else str(value.minute)
            if char == "s":
                return f"{value.second:02d}" if length >= 2 else str(value.second)
        return token

    out: list[str] = []
    i = 0
    while i < len(format_text):
        char = format_text[i]
        if char in {"y", "M", "d", "H", "h", "m", "s"}:
            j = i + 1
            while j < len(format_text) and format_text[j] == char:
                j += 1
            out.append(token_value(format_text[i:j]))
            i = j
            continue
        out.append(char)
        i += 1
    return "".join(out)


def _date_to_text(value: Any, options: Any = None, culture: Any = None) -> str | None:
    date_value = _to_date(value)
    if date_value is None:
        return None
    format_text: str | None = None
    if isinstance(options, Mapping):
        format_text = options.get("Format") or options.get("format")
        culture = options.get("Culture") or options.get("culture") or culture
    elif options is not None:
        format_text = str(options)
    return _temporal_format(date_value, format_text or "yyyy-MM-dd", culture)


_MONTH_NAMES_EN = [
    "January",
    "February",
    "March",
    "April",
    "May",
    "June",
    "July",
    "August",
    "September",
    "October",
    "November",
    "December",
]

_WEEKDAY_NAMES_EN = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]


def _datetime_to_text(value: Any, options: Any = None, culture: Any = None) -> str | None:
    datetime_value = _to_datetime(value, culture)
    if datetime_value is None:
        return None
    format_text: str | None = None
    if isinstance(options, Mapping):
        format_text = options.get("Format") or options.get("format")
        culture = options.get("Culture") or options.get("culture") or culture
    elif options is not None:
        format_text = str(options)
    return _temporal_format(datetime_value, format_text or "yyyy-MM-ddTHH:mm:ss", culture)


def _time_to_text(value: Any, options: Any = None, culture: Any = None) -> str | None:
    time_value = _to_time(value, culture)
    if time_value is None:
        return None
    format_text: str | None = None
    if isinstance(options, Mapping):
        format_text = options.get("Format") or options.get("format")
        culture = options.get("Culture") or options.get("culture") or culture
    elif options is not None:
        format_text = str(options)
    if format_text == "T":
        format_text = "HH:mm:ss"
    datetime_value = _dt.datetime.combine(_dt.date(2000, 1, 1), time_value)
    return _temporal_format(datetime_value, format_text or "HH:mm:ss", culture)


def _timezone_offset_text(value: _dt.datetime) -> str:
    offset = value.utcoffset() or _dt.timedelta()
    total_minutes = int(offset.total_seconds() // 60)
    sign = "+" if total_minutes >= 0 else "-"
    total_minutes = abs(total_minutes)
    return f"{sign}{total_minutes // 60:02d}:{total_minutes % 60:02d}"


def _datetimezone_to_text(value: Any, options: Any = None, culture: Any = None) -> str | None:
    datetime_value = _to_datetimezone(value, culture)
    if datetime_value is None:
        return None
    format_text: str | None = None
    if isinstance(options, Mapping):
        format_text = options.get("Format") or options.get("format")
        culture = options.get("Culture") or options.get("culture") or culture
    elif options is not None:
        format_text = str(options)
    if format_text == "O":
        return f"{datetime_value:%Y-%m-%dT%H:%M:%S}.0000000{_timezone_offset_text(datetime_value)}"
    format_text = format_text or "yyyy-MM-ddTHH:mm:ss zzz"
    offset_text = _timezone_offset_text(datetime_value)
    return _temporal_format(datetime_value.replace(tzinfo=None), format_text.replace("zzz", "__PQ_TZ__"), culture).replace("__PQ_TZ__", offset_text)


def _duration_to_text(value: Any, format_text: Any = None) -> str | None:
    if format_text is not None:
        raise ValueError("Duration.ToText preview supports the default format only.")
    duration = _to_duration(value)
    if duration is None:
        return None
    total_seconds = duration.total_seconds()
    sign = "-" if total_seconds < 0 else ""
    total_seconds = abs(total_seconds)
    days = int(total_seconds // 86400)
    remainder = total_seconds - (days * 86400)
    hours = int(remainder // 3600)
    remainder -= hours * 3600
    minutes = int(remainder // 60)
    seconds = remainder - minutes * 60
    if float(seconds).is_integer():
        seconds_text = f"{int(seconds):02d}"
    else:
        seconds_text = f"{seconds:06.3f}".rstrip("0").rstrip(".")
    return f"{sign}{days}.{hours:02d}:{minutes:02d}:{seconds_text}"


def _duration_from_text(value: Any) -> _dt.timedelta | None:
    if value is None or value == "":
        return None
    text = str(value).strip()
    if text.startswith("#duration"):
        return _to_duration(text)
    sign = -1 if text.startswith("-") else 1
    if text[:1] in {"-", "+"}:
        text = text[1:]
    days = 0
    if "." in text and text.find(".") < text.find(":"):
        day_text, text = text.split(".", 1)
        days = int(day_text)
    parts = text.split(":")
    if len(parts) != 3:
        return _to_duration(value)
    hours, minutes, seconds = int(parts[0]), int(parts[1]), float(parts[2])
    return sign * _dt.timedelta(days=days, hours=hours, minutes=minutes, seconds=seconds)


def _duration_record(value: Any) -> dict[str, Any]:
    duration = _to_duration(value) or _dt.timedelta()
    total_seconds = duration.total_seconds()
    sign = -1 if total_seconds < 0 else 1
    total_seconds = abs(total_seconds)
    days = int(total_seconds // 86400)
    remainder = total_seconds - (days * 86400)
    hours = int(remainder // 3600)
    remainder -= hours * 3600
    minutes = int(remainder // 60)
    seconds = remainder - (minutes * 60)
    return {
        "Days": sign * days,
        "Hours": sign * hours,
        "Minutes": sign * minutes,
        "Seconds": sign * seconds,
    }


def _date_record(value: Any) -> dict[str, Any]:
    date_value = _to_date(value)
    if date_value is None:
        return {"Year": None, "Month": None, "Day": None}
    return {"Year": date_value.year, "Month": date_value.month, "Day": date_value.day}


def _time_record(value: Any) -> dict[str, Any]:
    time_value = _to_time(value)
    if time_value is None:
        return {"Hour": None, "Minute": None, "Second": None}
    second = time_value.second + (time_value.microsecond / 1_000_000)
    return {"Hour": time_value.hour, "Minute": time_value.minute, "Second": second}


def _datetime_record(value: Any) -> dict[str, Any]:
    datetime_value = _to_datetime(value)
    if datetime_value is None:
        return {"Year": None, "Month": None, "Day": None, "Hour": None, "Minute": None, "Second": None}
    second = datetime_value.second + (datetime_value.microsecond / 1_000_000)
    return {
        "Year": datetime_value.year,
        "Month": datetime_value.month,
        "Day": datetime_value.day,
        "Hour": datetime_value.hour,
        "Minute": datetime_value.minute,
        "Second": second,
    }


def _datetimezone_record(value: Any) -> dict[str, Any]:
    datetime_value = _to_datetimezone(value)
    record = _datetime_record(datetime_value)
    offset = datetime_value.utcoffset() if datetime_value is not None else None
    total_minutes = int(offset.total_seconds() // 60) if offset is not None else 0
    hours = int(total_minutes / 60)
    minutes = abs(total_minutes) % 60
    record["ZoneHours"] = hours
    record["ZoneMinutes"] = minutes
    return record


def _datetime_from_filetime(value: Any, *, timezone: bool = False) -> _dt.datetime | None:
    number = _to_number(value)
    if number is None:
        return None
    base = _dt.datetime(1601, 1, 1, tzinfo=_dt.timezone.utc)
    result = base + _dt.timedelta(microseconds=int(number) / 10)
    return result if timezone else result.replace(tzinfo=None)


def _day_first_index(first_day: Any = None) -> int:
    text = str(first_day or "Day.Sunday").strip().lower()
    mapping = {
        "day.monday": 0,
        "monday": 0,
        "1": 0,
        "day.tuesday": 1,
        "tuesday": 1,
        "2": 1,
        "day.wednesday": 2,
        "wednesday": 2,
        "3": 2,
        "day.thursday": 3,
        "thursday": 3,
        "4": 3,
        "day.friday": 4,
        "friday": 4,
        "5": 4,
        "day.saturday": 5,
        "saturday": 5,
        "6": 5,
        "day.sunday": 6,
        "sunday": 6,
        "0": 6,
    }
    return mapping.get(text, 6)


def _week_index(value: _dt.date, period_start: _dt.date, first_day: Any = None) -> int:
    first_idx = _day_first_index(first_day)
    offset = (period_start.weekday() - first_idx) % 7
    first_week_start = period_start - _dt.timedelta(days=offset)
    return ((value - first_week_start).days // 7) + 1


def _date_quarter(value: _dt.date) -> int:
    return ((value.month - 1) // 3) + 1


def _month_delta(left: _dt.date, right: _dt.date) -> int:
    return (left.year - right.year) * 12 + (left.month - right.month)


def _quarter_delta(left: _dt.date, right: _dt.date) -> int:
    return (left.year - right.year) * 4 + (_date_quarter(left) - _date_quarter(right))


def _week_start(value: _dt.date, first_day: Any = None) -> _dt.date:
    return value - _dt.timedelta(days=((value.weekday() - _day_first_index(first_day)) % 7))


def _date_is_in_period(fn: str, value: Any, n: Any = None) -> bool:
    date_value = _to_date(value)
    if date_value is None:
        return False
    today = _dt.date.today()
    upper = fn.upper()
    if upper.endswith("CURRENTDAY"):
        return date_value == today
    if upper.endswith("CURRENTWEEK"):
        return _week_start(date_value) == _week_start(today)
    if upper.endswith("CURRENTMONTH"):
        return date_value.year == today.year and date_value.month == today.month
    if upper.endswith("CURRENTQUARTER"):
        return date_value.year == today.year and _date_quarter(date_value) == _date_quarter(today)
    if upper.endswith("CURRENTYEAR"):
        return date_value.year == today.year
    if upper.endswith("NEXTDAY"):
        return date_value == today + _dt.timedelta(days=1)
    if upper.endswith("NEXTWEEK"):
        return _week_start(date_value) == _week_start(today) + _dt.timedelta(days=7)
    if upper.endswith("NEXTMONTH"):
        return _month_delta(date_value, today) == 1
    if upper.endswith("NEXTQUARTER"):
        return _quarter_delta(date_value, today) == 1
    if upper.endswith("NEXTYEAR"):
        return date_value.year - today.year == 1
    if upper.endswith("PREVIOUSDAY"):
        return date_value == today - _dt.timedelta(days=1)
    if upper.endswith("PREVIOUSWEEK"):
        return _week_start(date_value) == _week_start(today) - _dt.timedelta(days=7)
    if upper.endswith("PREVIOUSMONTH"):
        return _month_delta(today, date_value) == 1
    if upper.endswith("PREVIOUSQUARTER"):
        return _quarter_delta(today, date_value) == 1
    if upper.endswith("PREVIOUSYEAR"):
        return today.year - date_value.year == 1
    if upper.endswith("NEXTNDAYS"):
        count = max(0, int(_to_number(n) or 0))
        return today < date_value <= today + _dt.timedelta(days=count)
    if upper.endswith("NEXTNWEEKS"):
        count = max(0, int(_to_number(n) or 0))
        delta = (_week_start(date_value) - _week_start(today)).days // 7
        return 1 <= delta <= count
    if upper.endswith("NEXTNMONTHS"):
        return 1 <= _month_delta(date_value, today) <= max(0, int(_to_number(n) or 0))
    if upper.endswith("NEXTNQUARTERS"):
        return 1 <= _quarter_delta(date_value, today) <= max(0, int(_to_number(n) or 0))
    if upper.endswith("NEXTNYEARS"):
        return 1 <= (date_value.year - today.year) <= max(0, int(_to_number(n) or 0))
    if upper.endswith("PREVIOUSNDAYS"):
        count = max(0, int(_to_number(n) or 0))
        return today - _dt.timedelta(days=count) <= date_value < today
    if upper.endswith("PREVIOUSNWEEKS"):
        count = max(0, int(_to_number(n) or 0))
        delta = (_week_start(today) - _week_start(date_value)).days // 7
        return 1 <= delta <= count
    if upper.endswith("PREVIOUSNMONTHS"):
        return 1 <= _month_delta(today, date_value) <= max(0, int(_to_number(n) or 0))
    if upper.endswith("PREVIOUSNQUARTERS"):
        return 1 <= _quarter_delta(today, date_value) <= max(0, int(_to_number(n) or 0))
    if upper.endswith("PREVIOUSNYEARS"):
        return 1 <= (today.year - date_value.year) <= max(0, int(_to_number(n) or 0))
    if upper.endswith("YEARTODATE"):
        return date_value.year == today.year and date_value <= today
    return False


def _datetime_is_in_period(fn: str, value: Any, n: Any = None) -> bool:
    datetime_value = _to_datetime(value)
    if datetime_value is None:
        return False
    now = _dt.datetime.now()
    upper = fn.upper()
    if upper.endswith("CURRENTHOUR"):
        return datetime_value.replace(minute=0, second=0, microsecond=0) == now.replace(minute=0, second=0, microsecond=0)
    if upper.endswith("CURRENTMINUTE"):
        return datetime_value.replace(second=0, microsecond=0) == now.replace(second=0, microsecond=0)
    if upper.endswith("CURRENTSECOND"):
        return datetime_value.replace(microsecond=0) == now.replace(microsecond=0)
    for label, unit in (("HOUR", "hours"), ("MINUTE", "minutes"), ("SECOND", "seconds")):
        if upper.endswith(f"NEXT{label}"):
            return now < datetime_value <= now + _dt.timedelta(**{unit: 1})
        if upper.endswith(f"PREVIOUS{label}"):
            return now - _dt.timedelta(**{unit: 1}) <= datetime_value < now
        if upper.endswith(f"NEXTN{label}S"):
            return now < datetime_value <= now + _dt.timedelta(**{unit: max(0, int(_to_number(n) or 0))})
        if upper.endswith(f"PREVIOUSN{label}S"):
            return now - _dt.timedelta(**{unit: max(0, int(_to_number(n) or 0))}) <= datetime_value < now
    return False


def _eval_add_subtract(lval: Any, op: str, rval: Any) -> Any:
    if lval is None or rval is None:
        return None
    if isinstance(lval, _dt.datetime) and isinstance(rval, _dt.datetime) and op == "-":
        return lval - rval
    if isinstance(lval, _dt.date) and isinstance(rval, _dt.date) and op == "-":
        return lval - rval
    if isinstance(lval, (_dt.date, _dt.datetime)) and isinstance(rval, _dt.timedelta):
        return lval + rval if op == "+" else lval - rval
    if isinstance(lval, _dt.timedelta) and isinstance(rval, (_dt.date, _dt.datetime)) and op == "+":
        return rval + lval
    if isinstance(lval, _dt.timedelta) and isinstance(rval, _dt.timedelta):
        return lval + rval if op == "+" else lval - rval
    left_number = _to_number(lval)
    right_number = _to_number(rval)
    if left_number is None or right_number is None:
        return None
    return left_number + right_number if op == "+" else left_number - right_number


def _eval_function_call(
    name: str,
    args: list[Any],
    *,
    variables: Mapping[str, Any] | None = None,
) -> Any:
    fn = name.strip().upper()
    if fn == "CSV.DOCUMENT":
        return _csv_document_records(args[0] if args else None, args[1] if len(args) > 1 else None)
    if fn == "JSON.DOCUMENT":
        return _json_document_value(args[0] if args else None)
    if fn == "JSON.FROMVALUE":
        return _json_from_value(args[0] if args else None, args[1] if len(args) > 1 else None)
    if fn == "XML.TABLES":
        return _xml_tables_value(args[0] if args else None)
    if fn == "XML.DOCUMENT":
        return _xml_tables_value(args[0] if args else None)
    if fn == "HTML.TABLE":
        return _html_table_value(args[0] if args else "", args[1] if len(args) > 1 else [])
    if fn == "ACCESSCONTROLENTRY.CONDITIONTOIDENTITIES":
        return []
    if fn == "IDENTITY.FROM":
        return _identity_value(args[0] if args else None, args[1] if len(args) > 1 else None)
    if fn == "IDENTITY.ISMEMBEROF":
        return _identity_is_member_of(args[0] if args else None, args[1] if len(args) > 1 else None)
    if fn == "IDENTITYPROVIDER.DEFAULT":
        return {"Kind": "IdentityProvider", "Name": "Default"}
    if fn == "ODBC.INFEROPTIONS":
        return {
            "Kind": "OdbcOptions",
            "ConnectionString": args[0] if args else None,
            "Options": args[1] if len(args) > 1 and isinstance(args[1], Mapping) else {},
            "Inferred": True,
        }
    if fn == "WEB.HEADERS":
        return {"Url": args[0] if args else None, "Headers": {}, "NetworkAccess": "not_performed"}
    if fn == "EXCEL.WORKBOOK":
        return _excel_workbook_value(args[0] if args else None, args[1] if len(args) > 1 else None)
    if fn == "ERROR.RECORD":
        return _error_record(
            args[0] if args else "Expression.Error",
            args[1] if len(args) > 1 else None,
            args[2] if len(args) > 2 else None,
            args[3] if len(args) > 3 else None,
            args[4] if len(args) > 4 else None,
        )
    if fn == "EXPRESSION.IDENTIFIER":
        return _m_identifier_text(str(args[0] if args else ""))
    if fn == "EXPRESSION.CONSTANT":
        return _m_literal_text(args[0] if args else None)
    if fn == "EXPRESSION.EVALUATE":
        expression_text = str(args[0] if args else "")
        env = args[1] if len(args) > 1 and isinstance(args[1], Mapping) else None
        return _eval_m_expression([], [], expression_text, variables=env)
    if fn == "FUNCTION.INVOKE" or fn == "FUNCTION.INVOKEAFTER" or fn == "FUNCTION.INVOKEWITHERRORCONTEXT":
        function_value = args[0] if args else None
        call_args = args[1] if len(args) > 1 and isinstance(args[1], list) else []
        if isinstance(function_value, (MFunctionValue, MBuiltinFunctionValue)):
            return _invoke_m_function(function_value, call_args, variables=variables)
        raise ValueError(f"{fn} preview needs a function value.")
    if fn == "FUNCTION.FROM":
        function_value = args[1] if len(args) > 1 else None
        if isinstance(function_value, (MFunctionValue, MBuiltinFunctionValue)):
            return function_value
        raise ValueError("Function.From preview needs a function value.")
    if fn == "FUNCTION.ISDATASOURCE":
        return False
    if fn == "FUNCTION.SCALARVECTOR":
        function_value = args[2] if len(args) > 2 else (args[0] if args else None)
        if isinstance(function_value, (MFunctionValue, MBuiltinFunctionValue)):
            return function_value
        raise ValueError("Function.ScalarVector preview needs a scalar function value.")
    if fn == "ITEMEXPRESSION.FROM":
        return {"Kind": "ItemExpression", "Expression": args[0] if args else None}
    if fn == "ROWEXPRESSION.FROM":
        return {"Kind": "RowExpression", "Expression": args[0] if args else None}
    if fn == "ROWEXPRESSION.COLUMN":
        return {"Kind": "RowExpression", "Column": args[0] if args else None}
    if fn == "DIAGNOSTICS.ACTIVITYID":
        return str(_uuid.uuid4())
    if fn == "DIAGNOSTICS.CORRELATIONID":
        return str(_uuid.uuid4())
    if fn == "DIAGNOSTICS.TRACE":
        return args[2] if len(args) > 2 else None
    if fn == "CHARACTER.FROMNUMBER":
        return chr(_int_number(args[0] if args else 0))
    if fn == "CHARACTER.TONUMBER":
        text = str(args[0] if args else "")
        return ord(text[0]) if text else None
    if fn == "LINES.FROMTEXT":
        return _lines_from_text(args[0] if args else "", args[1] if len(args) > 1 else None, args[2] if len(args) > 2 else None)
    if fn == "LINES.TOTEXT":
        return _lines_to_text(args[0] if args else [], args[1] if len(args) > 1 else None)
    if fn == "LINES.FROMBINARY":
        text = _text_from_binary(args[0] if args else None, args[3] if len(args) > 3 else None)
        return _lines_from_text(text, args[1] if len(args) > 1 else None, args[2] if len(args) > 2 else None)
    if fn == "LINES.TOBINARY":
        text = _lines_to_text(args[0] if args else [], args[1] if len(args) > 1 else None)
        return _text_to_binary(text, args[2] if len(args) > 2 else None, args[3] if len(args) > 3 else None)
    if fn == "URI.ESCAPEDATASTRING":
        return _urlparse.quote(str(args[0] if args else ""), safe="")
    if fn == "URI.UNESCAPEDATASTRING":
        return _urlparse.unquote(str(args[0] if args else ""))
    if fn == "URI.PARTS":
        return _uri_parts(args[0] if args else "")
    if fn == "URI.BUILDQUERYSTRING":
        return _uri_build_query_string(args[0] if args else {})
    if fn == "URI.COMBINE":
        return _uri_combine(args[0] if args else "", args[1] if len(args) > 1 else "")
    if fn in {"GEOGRAPHY.FROMWELLKNOWNTEXT", "GEOMETRY.FROMWELLKNOWNTEXT"}:
        return _wkt_record(args[0] if args else "", kind=fn.split(".", 1)[0].title())
    if fn in {"GEOGRAPHY.TOWELLKNOWNTEXT", "GEOMETRY.TOWELLKNOWNTEXT"}:
        return _wkt_from_record(args[0] if args else None)
    if fn == "GEOGRAPHYPOINT.FROM":
        return {"Kind": "Geography", "Type": "POINT", "Latitude": args[0] if args else None, "Longitude": args[1] if len(args) > 1 else None}
    if fn == "GEOMETRYPOINT.FROM":
        return {"Kind": "Geometry", "Type": "POINT", "X": args[0] if args else None, "Y": args[1] if len(args) > 1 else None}
    if fn in {"VALUE.VIEWERROR", "BINARY.VIEWERROR", "TABLE.VIEWERROR"}:
        return {"Kind": fn, "Error": args[0] if args else None}
    if fn in {"VALUE.VIEWFUNCTION", "BINARY.VIEWFUNCTION", "TABLE.VIEWFUNCTION"}:
        return args[0] if args else None
    if fn == "COMPARER.ORDINAL":
        return "Comparer.Ordinal"
    if fn == "COMPARER.ORDINALIGNORECASE":
        return "Comparer.OrdinalIgnoreCase"
    if fn == "COMPARER.FROMCULTURE":
        ignore_case = bool(args[1]) if len(args) > 1 else False
        return "Comparer.FromCultureIgnoreCase" if ignore_case else "Comparer.FromCulture"
    if fn == "COMPARER.EQUALS":
        comparer = args[0] if args else None
        left = "" if len(args) <= 1 or args[1] is None else str(args[1])
        right = "" if len(args) <= 2 or args[2] is None else str(args[2])
        if _comparer_ignore_case(comparer):
            return left.lower() == right.lower()
        return left == right
    if fn == "REPLACER.REPLACEVALUE":
        value = args[0] if args else None
        old_value = args[1] if len(args) > 1 else None
        new_value = args[2] if len(args) > 2 else None
        return new_value if value == old_value else value
    if fn == "REPLACER.REPLACETEXT":
        value = args[0] if args else None
        old_value = "" if len(args) <= 1 or args[1] is None else str(args[1])
        new_value = "" if len(args) <= 2 or args[2] is None else str(args[2])
        return None if value is None else str(value).replace(old_value, new_value)
    if fn == "SPLITTER.SPLITBYNOTHING":
        return MBuiltinFunctionValue(fn)
    if fn == "SPLITTER.SPLITTEXTBYDELIMITER":
        return MBuiltinFunctionValue(fn, {"delimiter": args[0] if args else "", "quote_style": args[1] if len(args) > 1 else None})
    if fn == "SPLITTER.SPLITTEXTBYANYDELIMITER":
        return MBuiltinFunctionValue(fn, {
            "delimiters": args[0] if args else [],
            "quote_style": args[1] if len(args) > 1 else None,
            "start_at_end": args[2] if len(args) > 2 else False,
        })
    if fn == "SPLITTER.SPLITTEXTBYEACHDELIMITER":
        return MBuiltinFunctionValue(fn, {
            "delimiters": args[0] if args else [],
            "quote_style": args[1] if len(args) > 1 else None,
            "start_at_end": args[2] if len(args) > 2 else False,
        })
    if fn == "SPLITTER.SPLITTEXTBYLENGTHS":
        return MBuiltinFunctionValue(fn, {"lengths": args[0] if args else [], "start_at_end": args[1] if len(args) > 1 else False})
    if fn == "SPLITTER.SPLITTEXTBYPOSITIONS":
        return MBuiltinFunctionValue(fn, {"positions": args[0] if args else [], "start_at_end": args[1] if len(args) > 1 else False})
    if fn == "SPLITTER.SPLITTEXTBYRANGES":
        return MBuiltinFunctionValue(fn, {"ranges": args[0] if args else [], "start_at_end": args[1] if len(args) > 1 else False})
    if fn == "SPLITTER.SPLITTEXTBYREPEATEDLENGTHS":
        return MBuiltinFunctionValue(fn, {"length": args[0] if args else 1, "start_at_end": args[1] if len(args) > 1 else False})
    if fn == "SPLITTER.SPLITTEXTBYWHITESPACE":
        return MBuiltinFunctionValue(fn, {"quote_style": args[0] if args else None})
    if fn == "SPLITTER.SPLITTEXTBYCHARACTERTRANSITION":
        return MBuiltinFunctionValue(fn, {"before": args[0] if args else None, "after": args[1] if len(args) > 1 else None})
    if fn == "COMBINER.COMBINETEXTBYDELIMITER":
        return MBuiltinFunctionValue(fn, {"delimiter": args[0] if args else "", "quote_style": args[1] if len(args) > 1 else None})
    if fn == "COMBINER.COMBINETEXTBYEACHDELIMITER":
        return MBuiltinFunctionValue(fn, {
            "delimiters": args[0] if args else [],
            "quote_style": args[1] if len(args) > 1 else None,
            "start_at_end": args[2] if len(args) > 2 else False,
        })
    if fn == "COMBINER.COMBINETEXTBYLENGTHS":
        return MBuiltinFunctionValue(fn, {"lengths": args[0] if args else []})
    if fn == "COMBINER.COMBINETEXTBYPOSITIONS":
        return MBuiltinFunctionValue(fn, {"positions": args[0] if args else []})
    if fn == "COMBINER.COMBINETEXTBYRANGES":
        return MBuiltinFunctionValue(fn, {"ranges": args[0] if args else []})
    if fn == "BINARY.BUFFER":
        return args[0] if args else b""
    if fn == "BINARY.FROM":
        return _binary_from_value(args[0] if args else None)
    if fn == "BINARY.FROMLIST":
        return _binary_from_list(args[0] if args else [])
    if fn == "BINARY.FROMTEXT":
        return _binary_from_text(args[0] if args else "", args[1] if len(args) > 1 else None)
    if fn == "BINARY.TOTEXT":
        return _binary_to_text(args[0] if args else None, args[1] if len(args) > 1 else None)
    if fn == "BINARY.TOLIST":
        return list(bytes(_unwrap_metadata(args[0] if args else b"") or b""))
    if fn == "BINARY.LENGTH":
        value = _unwrap_metadata(args[0] if args else b"")
        return len(bytes(value or b""))
    if fn == "BINARY.INFERCONTENTTYPE":
        return _binary_content_type(args[0] if args else b"")
    if fn == "BINARY.APPROXIMATELENGTH":
        value = _unwrap_metadata(args[0] if args else b"")
        return len(bytes(value or b""))
    if fn == "BINARY.RANGE":
        return _binary_range(args[0] if args else b"", args[1] if len(args) > 1 else 0, args[2] if len(args) > 2 else None)
    if fn == "BINARY.SPLIT":
        return _binary_split(args[0] if args else b"", args[1] if len(args) > 1 else 1)
    if fn == "BINARY.COMBINE":
        parts = args[0] if args and isinstance(args[0], list) else []
        return b"".join(bytes(_unwrap_metadata(item) or b"") for item in parts)
    if fn == "BINARY.COMPRESS":
        return _binary_compress(args[0] if args else b"", args[1] if len(args) > 1 else None)
    if fn == "BINARY.DECOMPRESS":
        return _binary_decompress(args[0] if args else b"", args[1] if len(args) > 1 else None)
    if fn == "BINARY.VIEW":
        handlers = args[1] if len(args) > 1 and isinstance(args[1], Mapping) else {}
        get_stream = handlers.get("GetStream") if isinstance(handlers, Mapping) else None
        if isinstance(get_stream, (MFunctionValue, MBuiltinFunctionValue)):
            return _invoke_m_function(get_stream, [], variables=variables)
        return args[0] if args else b""
    if fn.startswith("BINARYFORMAT."):
        if fn == "BINARYFORMAT.BYTEORDER":
            return _binary_format_with_options(args[0] if args else MBuiltinFunctionValue("BinaryFormat.Null"), byte_order=args[1] if len(args) > 1 else None)
        if fn == "BINARYFORMAT.BINARY":
            return MBuiltinFunctionValue(fn, {"length": args[0] if args else None})
        if fn == "BINARYFORMAT.TEXT":
            return MBuiltinFunctionValue(fn, {"length": args[0] if args else None, "encoding": args[1] if len(args) > 1 else None})
        if fn == "BINARYFORMAT.LENGTH":
            return MBuiltinFunctionValue(fn, {"format": args[0] if args else MBuiltinFunctionValue("BinaryFormat.Null"), "length": args[1] if len(args) > 1 else None})
        if fn == "BINARYFORMAT.LIST":
            return MBuiltinFunctionValue(fn, {"format": args[0] if args else MBuiltinFunctionValue("BinaryFormat.Null"), "count": args[1] if len(args) > 1 else None})
        if fn in {"BINARYFORMAT.GROUP", "BINARYFORMAT.RECORD"}:
            return MBuiltinFunctionValue(fn, {"fields": args[0] if args else []})
        if fn == "BINARYFORMAT.TRANSFORM":
            return MBuiltinFunctionValue(fn, {"format": args[0] if args else MBuiltinFunctionValue("BinaryFormat.Null"), "transform": args[1] if len(args) > 1 else None})
        if fn == "BINARYFORMAT.CHOICE":
            return MBuiltinFunctionValue(fn, {
                "format": args[0] if args else MBuiltinFunctionValue("BinaryFormat.Null"),
                "choices": args[1] if len(args) > 1 else {},
                "default": args[2] if len(args) > 2 else None,
            })
        return MBuiltinFunctionValue(fn)
    if fn == "TEXT.TOBINARY":
        return _text_to_binary(args[0] if args else None, args[1] if len(args) > 1 else None, args[2] if len(args) > 2 else None)
    if fn == "TEXT.FROMBINARY":
        return _text_from_binary(args[0] if args else None, args[1] if len(args) > 1 else None)
    if fn == "TABLE.PROMOTEHEADERS":
        return _promote_headers_records(args[0] if args else [])
    if fn == "TABLE.COLUMNNAMES":
        records = args[0] if args and isinstance(args[0], list) else []
        names: list[str] = []
        seen: set[str] = set()
        for record in records:
            if not isinstance(record, Mapping):
                continue
            for key in record.keys():
                name = str(key)
                upper = name.upper()
                if upper not in seen:
                    seen.add(upper)
                    names.append(name)
        return names
    if fn == "TABLE.ROWCOUNT":
        records = args[0] if args and isinstance(args[0], list) else []
        return len(records)
    if fn == "TABLE.APPROXIMATEROWCOUNT":
        records = args[0] if args and isinstance(args[0], list) else []
        return len(records)
    if fn == "TABLE.HASCOLUMNS":
        cols, _rows = _table_value_columns_rows(args[0] if args else [])
        available = {str(col.get("name") or "").strip().upper() for col in cols}
        wanted_arg = args[1] if len(args) > 1 else []
        wanted = [wanted_arg] if isinstance(wanted_arg, str) else list(wanted_arg if isinstance(wanted_arg, list) else [])
        return all(str(name).strip().upper() in available for name in wanted)
    if fn == "TABLE.ISDISTINCT":
        records = [dict(record) for record in args[0] if isinstance(record, Mapping)] if args and isinstance(args[0], list) else []
        criteria = args[1] if len(args) > 1 else None
        if isinstance(criteria, str):
            keys = [criteria]
        elif isinstance(criteria, list) and all(isinstance(item, str) for item in criteria):
            keys = list(criteria)
        else:
            keys = []
        seen: set[Any] = set()
        for record in records:
            value = tuple(_mapping_get_case_insensitive(record, key) if any(str(existing).upper() == key.upper() for existing in record.keys()) else None for key in keys) if keys else tuple((key, _value_key(value)) for key, value in record.items())
            marker = _value_key(value)
            if marker in seen:
                return False
            seen.add(marker)
        return True
    if fn == "TABLE.FIRSTVALUE":
        cols, rows = _table_value_columns_rows(args[0] if args else [])
        default = args[1] if len(args) > 1 else None
        if not cols or not rows:
            return default
        return rows[0][0] if rows[0] else default
    if fn == "TABLE.SINGLEROW":
        records = [dict(record) for record in args[0] if isinstance(record, Mapping)] if args and isinstance(args[0], list) else []
        if len(records) != 1:
            raise ValueError("Table.SingleRow preview needs exactly one row.")
        return records[0]
    if fn == "TABLE.PREFIXCOLUMNS":
        records = [dict(record) for record in args[0] if isinstance(record, Mapping)] if args and isinstance(args[0], list) else []
        prefix = str(args[1] if len(args) > 1 and args[1] is not None else "")
        return [
            {f"{prefix}.{key}": value for key, value in record.items()}
            for record in records
        ]
    if fn == "TABLE.FROMVALUE":
        value = args[0] if args else None
        return [{"Value": value}]
    if fn == "TABLE.VIEW":
        handlers = args[1] if len(args) > 1 and isinstance(args[1], Mapping) else {}
        get_rows = handlers.get("GetRows") if isinstance(handlers, Mapping) else None
        if isinstance(get_rows, (MFunctionValue, MBuiltinFunctionValue)):
            return _invoke_m_function(get_rows, [], variables=variables)
        return args[0] if args and isinstance(args[0], list) else []
    if fn == "TABLE.SPLIT":
        records = [dict(record) for record in args[0] if isinstance(record, Mapping)] if args and isinstance(args[0], list) else []
        size = max(1, int(_to_number(args[1] if len(args) > 1 else 1) or 1))
        return [records[idx : idx + size] for idx in range(0, len(records), size)]
    if fn == "TABLE.SPLITAT":
        records = [dict(record) for record in args[0] if isinstance(record, Mapping)] if args and isinstance(args[0], list) else []
        offset = max(0, int(_to_number(args[1] if len(args) > 1 else 0) or 0))
        return [records[:offset], records[offset:]]
    if fn == "TABLE.ADDJOINCOLUMN":
        left_cols, left_rows = _table_value_columns_rows(args[0] if args else [])
        right_cols, right_rows = _table_value_columns_rows(args[2] if len(args) > 2 else [])
        expression = "Table.NestedJoin(_, {left_keys}, _, {right_keys}, {nested_name}, {join_kind})".format(
            left_keys=_m_literal_text(_join_key_names(args[1] if len(args) > 1 else [])),
            right_keys=_m_literal_text(_join_key_names(args[3] if len(args) > 3 else [])),
            nested_name=_m_literal_text(args[4] if len(args) > 4 else "Joined"),
            join_kind=_m_literal_text(args[5] if len(args) > 5 else "JoinKind.LeftOuter"),
        )
        out_cols, out_rows = _nested_join_tables((left_cols, left_rows), (right_cols, right_rows), expression)
        return _table_records(out_cols, out_rows)
    if fn == "TABLE.COMBINECOLUMNSTORECORD":
        records = [dict(record) for record in args[0] if isinstance(record, Mapping)] if args and isinstance(args[0], list) else []
        source_arg = args[1] if len(args) > 1 else []
        source_names = [source_arg] if isinstance(source_arg, str) else [str(item) for item in source_arg] if isinstance(source_arg, list) else []
        output_name = str(args[2] if len(args) > 2 and args[2] is not None else "Record")
        source_lookup = {name.upper() for name in source_names}
        out: list[dict[str, Any]] = []
        for record in records:
            combined = {name: _mapping_get_case_insensitive(record, name) for name in source_names}
            row = {
                key: value
                for key, value in record.items()
                if str(key).upper() not in source_lookup
            }
            row[output_name] = combined
            out.append(row)
        return out
    if fn == "TABLE.ADDRANKCOLUMN":
        cols, rows = _table_value_columns_rows(args[0] if args else [])
        output_name = str(args[1] if len(args) > 1 and args[1] is not None else "Rank")
        criteria = args[2] if len(args) > 2 else []
        specs = _sort_specs_from_criteria(criteria)
        if not specs:
            raise ValueError("Table.AddRankColumn preview needs comparison criteria.")
        indexed_rows = [(idx, list(row)) for idx, row in enumerate(rows)]
        for name, descending in reversed(specs):
            col_idx = _find_column_index(cols, name)
            indexed_rows.sort(
                key=lambda item: (
                    item[1][col_idx] is None if col_idx < len(item[1]) else True,
                    _sortable_key(item[1][col_idx]) if col_idx < len(item[1]) else None,
                ),
                reverse=descending,
            )
        rank_indexes = [(name, _find_column_index(cols, name)) for name, _descending in specs]
        ranks = [0 for _row in rows]
        last_key: Any = object()
        current_rank = 0
        for sorted_pos, (original_idx, sorted_row) in enumerate(indexed_rows, start=1):
            key = tuple(_value_key(sorted_row[idx] if idx < len(sorted_row) else None) for _name, idx in rank_indexes)
            if key != last_key:
                current_rank = sorted_pos
                last_key = key
            ranks[original_idx] = current_rank
        return _table_records(
            [dict(col) for col in cols] + [{"name": output_name}],
            [list(row) + [ranks[idx]] for idx, row in enumerate(rows)],
        )
    if fn == "TABLE.PARTITION":
        records = [dict(record) for record in args[0] if isinstance(record, Mapping)] if args and isinstance(args[0], list) else []
        column = str(args[1] if len(args) > 1 else "")
        group_count = max(1, int(_to_number(args[2] if len(args) > 2 else 1) or 1))
        hash_fn = args[3] if len(args) > 3 else None
        partitions: list[list[dict[str, Any]]] = [[] for _idx in range(group_count)]
        for record in records:
            value = _mapping_get_case_insensitive(record, column) if column else record
            hashed = _invoke_m_function(hash_fn, [value], variables=variables) if isinstance(hash_fn, (MFunctionValue, MBuiltinFunctionValue)) else value
            try:
                bucket = int(_to_number(hashed) or 0) % group_count
            except Exception:
                bucket = abs(hash(_value_key(hashed))) % group_count
            partitions[bucket].append(record)
        return partitions
    if fn == "TABLE.FROMPARTITIONS":
        partition_column = str(args[0] if args else "Partition")
        partitions = args[1] if len(args) > 1 and isinstance(args[1], list) else []
        out: list[dict[str, Any]] = []
        for idx, partition in enumerate(partitions):
            records = [dict(record) for record in partition if isinstance(record, Mapping)] if isinstance(partition, list) else []
            for record in records:
                row = dict(record)
                row.setdefault(partition_column, idx)
                out.append(row)
        return out
    if fn == "TABLES.GETRELATIONSHIPS":
        return []
    if fn == "TABLE.KEYS":
        return []
    if fn == "TABLE.REPLACEKEYS":
        return args[0] if args else []
    if fn == "TABLE.PARTITIONVALUES":
        records = [dict(record) for record in args[0] if isinstance(record, Mapping)] if args and isinstance(args[0], list) else []
        column = str(args[1] if len(args) > 1 else "")
        seen: set[Any] = set()
        out: list[Any] = []
        for record in records:
            value = record.get(column)
            key = _value_key(value)
            if key not in seen:
                seen.add(key)
                out.append(value)
        return out
    if fn in {"TABLE.PARTIONKEY", "TABLE.REPLACEPARTITIONKEY", "TABLE.REPLACERELATIONSHIPIDENTITY", "TABLE.WITHERRORCONTEXT"}:
        return args[0] if args else []
    if fn == "TABLE.CONFORMTOPAGEREADER":
        return args[0] if args else []
    if fn == "LIST.CONFORMTOPAGEREADER":
        return list(args[0]) if args and isinstance(args[0], list) else []
    if fn == "TABLE.COLUMNCOUNT":
        cols, _rows = _table_value_columns_rows(args[0] if args else [])
        return len(cols)
    if fn == "TABLE.COLUMNSOFTYPE":
        cols, rows = _table_value_columns_rows(args[0] if args else [])
        types = args[1] if len(args) > 1 and isinstance(args[1], list) else []
        out: list[str] = []
        for idx, col in enumerate(cols):
            type_text = _column_type_text(cols, rows, idx)
            if any(_type_compatible(type_text, wanted) for wanted in types):
                out.append(str(col.get("name") or ""))
        return out
    if fn == "TABLE.SCHEMA":
        return _table_schema_records(args[0] if args else [])
    if fn == "TABLE.PROFILE":
        return _table_profile_records(args[0] if args else [])
    if fn == "TABLE.ISEMPTY":
        records = args[0] if args and isinstance(args[0], list) else []
        return len(records) == 0
    if fn == "TABLE.FIRST":
        records = args[0] if args and isinstance(args[0], list) else []
        return records[0] if records else (args[1] if len(args) > 1 else None)
    if fn == "TABLE.LAST":
        records = args[0] if args and isinstance(args[0], list) else []
        return records[-1] if records else (args[1] if len(args) > 1 else None)
    if fn == "TABLE.FROMRECORDS":
        return [dict(record) for record in args[0] if isinstance(record, Mapping)] if args and isinstance(args[0], list) else []
    if fn == "TABLE.FROMROWS":
        cols, rows = _constant_table_from_rows(
            args[0] if args else [],
            args[1] if len(args) > 1 else [],
        )
        return _table_records(cols, rows)
    if fn == "TABLE.FROMCOLUMNS":
        cols, rows = _constant_table_from_columns(
            args[0] if args else [],
            args[1] if len(args) > 1 else [],
        )
        return _table_records(cols, rows)
    if fn == "TABLE.FROMLIST":
        cols, rows = _constant_table_from_list(
            args[0] if args else [],
            args[2] if len(args) > 2 else [],
        )
        return _table_records(cols, rows)
    if fn == "TABLE.SELECTCOLUMNS":
        cols, rows = _table_value_columns_rows(args[0] if args else [])
        names = [str(item) for item in args[1]] if len(args) > 1 and isinstance(args[1], list) else []
        out_cols, out_rows = _select_columns(cols, rows, names, missing_field=_missing_field_mode_from_value(args[2] if len(args) > 2 else None))
        return _table_records(out_cols, out_rows)
    if fn == "TABLE.REMOVECOLUMNS":
        cols, rows = _table_value_columns_rows(args[0] if args else [])
        names = [str(item) for item in args[1]] if len(args) > 1 and isinstance(args[1], list) else []
        out_cols, out_rows = _remove_columns(cols, rows, names, missing_field=_missing_field_mode_from_value(args[2] if len(args) > 2 else None))
        return _table_records(out_cols, out_rows)
    if fn == "TABLE.REORDERCOLUMNS":
        cols, rows = _table_value_columns_rows(args[0] if args else [])
        names = [str(item) for item in args[1]] if len(args) > 1 and isinstance(args[1], list) else []
        selected: list[tuple[int, str]] = []
        used: set[int] = set()
        for col_name in names:
            idx = _find_column_index_or_none(cols, col_name)
            if idx is None:
                continue
            used.add(idx)
            selected.append((idx, str(cols[idx].get("name") or col_name)))
        for idx, col in enumerate(cols):
            if idx not in used:
                selected.append((idx, str(col.get("name") or "")))
        out_cols = [dict(cols[idx]) for idx, _name in selected]
        out_rows = [[row[idx] if idx < len(row) else None for idx, _name in selected] for row in rows]
        return _table_records(out_cols, out_rows)
    if fn == "TABLE.RENAMECOLUMNS":
        records = [dict(record) for record in args[0] if isinstance(record, Mapping)] if args and isinstance(args[0], list) else []
        pairs = args[1] if len(args) > 1 and isinstance(args[1], list) else []
        rename: dict[str, str] = {}
        for pair in pairs:
            if isinstance(pair, list) and len(pair) >= 2:
                rename[str(pair[0]).upper()] = str(pair[1])
        out: list[dict[str, Any]] = []
        for record in records:
            out.append({rename.get(str(key).upper(), str(key)): value for key, value in record.items()})
        return out
    if fn == "TABLE.DISTINCT":
        cols, rows = _table_value_columns_rows(args[0] if args else [])
        out_cols, out_rows = _distinct_rows(cols, rows)
        return _table_records(out_cols, out_rows)
    if fn == "TABLE.SORT":
        cols, rows = _table_value_columns_rows(args[0] if args else [])
        criteria = args[1] if len(args) > 1 else []
        out_cols, out_rows = _sort_rows_by_criteria(cols, rows, criteria)
        return _table_records(out_cols, out_rows)
    if fn == "TABLE.FINDTEXT":
        records = [dict(record) for record in args[0] if isinstance(record, Mapping)] if args and isinstance(args[0], list) else []
        needle = str(args[1] if len(args) > 1 and args[1] is not None else "")
        return [
            record for record in records
            if any(needle in str(value) for value in record.values() if value is not None)
        ]
    if fn == "TABLE.CONTAINS":
        records = [dict(record) for record in args[0] if isinstance(record, Mapping)] if args and isinstance(args[0], list) else []
        wanted = args[1] if len(args) > 1 and isinstance(args[1], Mapping) else {}
        criteria = args[2] if len(args) > 2 else None
        return any(_table_row_matches(record, wanted, criteria) for record in records)
    if fn == "TABLE.CONTAINSANY":
        records = [dict(record) for record in args[0] if isinstance(record, Mapping)] if args and isinstance(args[0], list) else []
        wanted_rows = args[1] if len(args) > 1 and isinstance(args[1], list) else []
        criteria = args[2] if len(args) > 2 else None
        return any(
            isinstance(wanted, Mapping) and any(_table_row_matches(record, wanted, criteria) for record in records)
            for wanted in wanted_rows
        )
    if fn == "TABLE.CONTAINSALL":
        records = [dict(record) for record in args[0] if isinstance(record, Mapping)] if args and isinstance(args[0], list) else []
        wanted_rows = args[1] if len(args) > 1 and isinstance(args[1], list) else []
        criteria = args[2] if len(args) > 2 else None
        return all(
            isinstance(wanted, Mapping) and any(_table_row_matches(record, wanted, criteria) for record in records)
            for wanted in wanted_rows
        )
    if fn == "TABLE.POSITIONOF":
        records = [dict(record) for record in args[0] if isinstance(record, Mapping)] if args and isinstance(args[0], list) else []
        wanted = args[1] if len(args) > 1 and isinstance(args[1], Mapping) else {}
        occurrence = args[2] if len(args) > 2 else None
        criteria = args[3] if len(args) > 3 else None
        return _table_positions(records, [wanted], occurrence, criteria)
    if fn == "TABLE.POSITIONOFANY":
        records = [dict(record) for record in args[0] if isinstance(record, Mapping)] if args and isinstance(args[0], list) else []
        wanted_rows = [dict(record) for record in args[1] if isinstance(record, Mapping)] if len(args) > 1 and isinstance(args[1], list) else []
        occurrence = args[2] if len(args) > 2 else None
        criteria = args[3] if len(args) > 3 else None
        return _table_positions(records, wanted_rows, occurrence, criteria)
    if fn == "TABLE.REMOVEMATCHINGROWS":
        records = [dict(record) for record in args[0] if isinstance(record, Mapping)] if args and isinstance(args[0], list) else []
        wanted_rows = [dict(record) for record in args[1] if isinstance(record, Mapping)] if len(args) > 1 and isinstance(args[1], list) else []
        criteria = args[2] if len(args) > 2 else None
        return [
            record for record in records
            if not any(_table_row_matches(record, wanted, criteria) for wanted in wanted_rows)
        ]
    if fn == "TABLE.REPLACEMATCHINGROWS":
        records = [dict(record) for record in args[0] if isinstance(record, Mapping)] if args and isinstance(args[0], list) else []
        replacements = args[1] if len(args) > 1 and isinstance(args[1], list) else []
        criteria = args[2] if len(args) > 2 else None
        specs: list[tuple[dict[str, Any], dict[str, Any]]] = []
        for spec in replacements:
            if isinstance(spec, list) and len(spec) >= 2 and isinstance(spec[0], Mapping) and isinstance(spec[1], Mapping):
                specs.append((dict(spec[0]), dict(spec[1])))
        out: list[dict[str, Any]] = []
        for record in records:
            replacement = next((new for old, new in specs if _table_row_matches(record, old, criteria)), None)
            out.append(dict(replacement) if replacement is not None else record)
        return out
    if fn == "TABLE.INSERTROWS":
        records = [dict(record) for record in args[0] if isinstance(record, Mapping)] if args and isinstance(args[0], list) else []
        offset = max(0, min(len(records), int(_to_number(args[1]) or 0))) if len(args) > 1 else 0
        inserted = [dict(record) for record in args[2] if isinstance(record, Mapping)] if len(args) > 2 and isinstance(args[2], list) else []
        return records[:offset] + inserted + records[offset:]
    if fn == "TABLE.REVERSEROWS":
        records = [dict(record) for record in args[0] if isinstance(record, Mapping)] if args and isinstance(args[0], list) else []
        return list(reversed(records))
    if fn == "TABLE.REPEAT":
        records = [dict(record) for record in args[0] if isinstance(record, Mapping)] if args and isinstance(args[0], list) else []
        count = max(0, int(_to_number(args[1]) or 0)) if len(args) > 1 else 0
        return [dict(record) for _idx in range(count) for record in records]
    if fn == "TABLE.TOLIST":
        cols, rows = _table_value_columns_rows(args[0] if args else [])
        combiner = args[1] if len(args) > 1 else None
        out: list[Any] = []
        for row in rows:
            values = [row[idx] if idx < len(row) else None for idx, _col in enumerate(cols)]
            if combiner is None:
                out.append(values)
            elif isinstance(combiner, (MFunctionValue, MBuiltinFunctionValue)):
                out.append(_invoke_m_function(combiner, [values], variables=variables))
            else:
                out.append("|".join("" if item is None else str(item) for item in values))
        return out
    if fn == "TABLE.TRANSPOSE":
        cols, rows = _table_value_columns_rows(args[0] if args else [])
        width = max([len(cols), *(len(row) for row in rows)] or [0])
        transposed_rows = [
            {f"Column{row_idx + 1}": rows[row_idx][col_idx] if col_idx < len(rows[row_idx]) else None for row_idx in range(len(rows))}
            for col_idx in range(width)
        ]
        return transposed_rows
    if fn == "TABLE.FIRSTN":
        records = [dict(record) for record in args[0] if isinstance(record, Mapping)] if args and isinstance(args[0], list) else []
        count = max(0, int(_to_number(args[1]) or 0)) if len(args) > 1 else 1
        return records[:count]
    if fn == "TABLE.LASTN":
        records = [dict(record) for record in args[0] if isinstance(record, Mapping)] if args and isinstance(args[0], list) else []
        count = max(0, int(_to_number(args[1]) or 0)) if len(args) > 1 else 1
        return records[-count:] if count > 0 else []
    if fn in {"TABLE.SKIP", "TABLE.REMOVEFIRSTN"}:
        records = [dict(record) for record in args[0] if isinstance(record, Mapping)] if args and isinstance(args[0], list) else []
        count = max(0, int(_to_number(args[1]) or 0)) if len(args) > 1 else 1
        return records[count:]
    if fn == "TABLE.REMOVELASTN":
        records = [dict(record) for record in args[0] if isinstance(record, Mapping)] if args and isinstance(args[0], list) else []
        count = max(0, int(_to_number(args[1]) or 0)) if len(args) > 1 else 1
        return records[:-count] if count > 0 else records
    if fn == "TABLE.RANGE":
        records = [dict(record) for record in args[0] if isinstance(record, Mapping)] if args and isinstance(args[0], list) else []
        offset = max(0, int(_to_number(args[1]) or 0)) if len(args) > 1 else 0
        count = max(0, int(_to_number(args[2]) or 0)) if len(args) > 2 else len(records) - offset
        return records[offset : offset + count]
    if fn == "TABLE.REMOVEROWS":
        records = [dict(record) for record in args[0] if isinstance(record, Mapping)] if args and isinstance(args[0], list) else []
        offset = max(0, int(_to_number(args[1]) or 0)) if len(args) > 1 else 0
        count = 1 if len(args) <= 2 or args[2] is None else max(0, int(_to_number(args[2]) or 0))
        return records[:offset] + records[offset + count :]
    if fn == "TABLE.ALTERNATEROWS":
        records = [dict(record) for record in args[0] if isinstance(record, Mapping)] if args and isinstance(args[0], list) else []
        offset = max(0, int(_to_number(args[1]) or 0)) if len(args) > 1 else 0
        skip = max(0, int(_to_number(args[2]) or 0)) if len(args) > 2 else 0
        take = max(0, int(_to_number(args[3]) or 0)) if len(args) > 3 else 0
        return _alternate_sequence(records, offset, skip, take)
    if fn == "TABLE.REPLACEROWS":
        records = [dict(record) for record in args[0] if isinstance(record, Mapping)] if args and isinstance(args[0], list) else []
        offset = max(0, int(_to_number(args[1]) or 0)) if len(args) > 1 else 0
        count = max(0, int(_to_number(args[2]) or 0)) if len(args) > 2 else 0
        replacements = [dict(record) for record in args[3] if isinstance(record, Mapping)] if len(args) > 3 and isinstance(args[3], list) else []
        return records[:offset] + replacements + records[offset + count :]
    if fn == "TEXT.UPPER":
        return None if args[0] is None else str(args[0]).upper()
    if fn == "TEXT.LOWER":
        return None if args[0] is None else str(args[0]).lower()
    if fn == "TEXT.TRIM":
        return _trim_text(args[0], args[1] if len(args) > 1 else None)
    if fn == "TEXT.TRIMSTART":
        return _trim_text(args[0], args[1] if len(args) > 1 else None, side="start")
    if fn == "TEXT.TRIMEND":
        return _trim_text(args[0], args[1] if len(args) > 1 else None, side="end")
    if fn == "TEXT.CLEAN":
        return None if args[0] is None else "".join(ch for ch in str(args[0]) if ch.isprintable())
    if fn == "TEXT.PROPER":
        return None if args[0] is None else str(args[0]).title()
    if fn == "TEXT.LENGTH":
        return None if args[0] is None else len(str(args[0]))
    if fn == "TEXT.REPLACE":
        return None if args[0] is None else str(args[0]).replace(str(args[1]), str(args[2]))
    if fn == "TEXT.POSITIONOF":
        if args[0] is None:
            return -1
        value = str(args[0])
        needle = str(args[1]) if len(args) > 1 else ""
        occurrence = str(args[2]).strip().upper() if len(args) > 2 and args[2] is not None else "OCCURRENCE.FIRST"
        if len(args) > 3 and _comparer_ignore_case(args[3]):
            value = value.lower()
            needle = needle.lower()
        if occurrence.endswith("LAST"):
            return value.rfind(needle)
        return value.find(needle)
    if fn == "TEXT.POSITIONOFANY":
        if args[0] is None:
            return -1
        value = str(args[0])
        chars = [str(item) for item in args[1]] if len(args) > 1 and isinstance(args[1], list) else []
        occurrence = str(args[2]).strip().upper() if len(args) > 2 and args[2] is not None else "OCCURRENCE.FIRST"
        positions = [idx for idx, char in enumerate(value) if char in chars]
        if occurrence.endswith("ALL"):
            return positions
        if occurrence.endswith("LAST"):
            return positions[-1] if positions else -1
        return positions[0] if positions else -1
    if fn == "TEXT.REPEAT":
        return None if args[0] is None else str(args[0]) * max(0, int(args[1]))
    if fn == "TEXT.INSERT":
        if args[0] is None:
            return None
        value = str(args[0])
        offset = max(0, min(len(value), int(args[1])))
        return value[:offset] + str(args[2]) + value[offset:]
    if fn == "TEXT.REPLACERANGE":
        if args[0] is None:
            return None
        value = str(args[0])
        offset = max(0, min(len(value), int(args[1])))
        count = max(0, int(args[2]))
        return value[:offset] + str(args[3]) + value[offset + count :]
    if fn == "TEXT.REMOVERANGE":
        if args[0] is None:
            return None
        value = str(args[0])
        offset = max(0, min(len(value), int(args[1])))
        if len(args) > 2 and args[2] is not None:
            count = max(0, int(args[2]))
            return value[:offset] + value[offset + count :]
        return value[:offset]
    if fn == "TEXT.SELECT":
        if args[0] is None:
            return None
        allowed = _text_character_set(args[1] if len(args) > 1 else [])
        return "".join(ch for ch in str(args[0]) if ch in allowed)
    if fn == "TEXT.REMOVE":
        if args[0] is None:
            return None
        removed = _text_character_set(args[1] if len(args) > 1 else [])
        return "".join(ch for ch in str(args[0]) if ch not in removed)
    if fn == "TEXT.TOLIST":
        return [] if args[0] is None else list(str(args[0]))
    if fn == "TEXT.REVERSE":
        return None if args[0] is None else str(args[0])[::-1]
    if fn == "TEXT.COMBINE":
        values = args[0] if isinstance(args[0], list) else []
        delimiter = str(args[1]) if len(args) > 1 and args[1] is not None else ""
        return delimiter.join(str(item) for item in values if item is not None)
    if fn == "TEXT.CONTAINS":
        if args[0] is None:
            return False
        value = str(args[0])
        needle = str(args[1]) if len(args) > 1 else ""
        if len(args) > 2 and _comparer_ignore_case(args[2]):
            return needle.lower() in value.lower()
        return needle in value
    if fn == "TEXT.STARTSWITH":
        if args[0] is None:
            return False
        value = str(args[0])
        prefix = str(args[1]) if len(args) > 1 else ""
        if len(args) > 2 and _comparer_ignore_case(args[2]):
            return value.lower().startswith(prefix.lower())
        return value.startswith(prefix)
    if fn == "TEXT.ENDSWITH":
        if args[0] is None:
            return False
        value = str(args[0])
        suffix = str(args[1]) if len(args) > 1 else ""
        if len(args) > 2 and _comparer_ignore_case(args[2]):
            return value.lower().endswith(suffix.lower())
        return value.endswith(suffix)
    if fn == "TEXT.SPLIT":
        return [] if args[0] is None else str(args[0]).split(str(args[1]))
    if fn == "TEXT.SPLITANY":
        if args[0] is None:
            return []
        separators = re.escape(str(args[1]) if len(args) > 1 else "")
        if not separators:
            return [str(args[0])]
        return re.split(f"[{separators}]", str(args[0]))
    if fn == "TEXT.BEFOREDELIMITER":
        if args[0] is None:
            return None
        value = str(args[0])
        delimiter = str(args[1])
        occurrence = int(args[2]) if len(args) > 2 and args[2] is not None else 0
        idx = _text_delimiter_index(value, delimiter, occurrence)
        return value if idx < 0 else value[:idx]
    if fn == "TEXT.AFTERDELIMITER":
        if args[0] is None:
            return None
        value = str(args[0])
        delimiter = str(args[1])
        occurrence = int(args[2]) if len(args) > 2 and args[2] is not None else 0
        idx = _text_delimiter_index(value, delimiter, occurrence)
        return "" if idx < 0 else value[idx + len(delimiter) :]
    if fn == "TEXT.BETWEENDELIMITERS":
        if args[0] is None:
            return None
        value = str(args[0])
        first = str(args[1])
        second = str(args[2])
        start = value.find(first)
        if start < 0:
            return ""
        tail = value[start + len(first) :]
        end = tail.find(second)
        return tail if end < 0 else tail[:end]
    if fn == "TEXT.START":
        return None if args[0] is None else str(args[0])[: int(args[1])]
    if fn == "TEXT.END":
        return None if args[0] is None else str(args[0])[-int(args[1]) :]
    if fn == "TEXT.AT":
        if args[0] is None:
            return None
        value = str(args[0])
        idx = int(args[1])
        return value[idx]
    if fn == "TEXT.MIDDLE":
        if args[0] is None:
            return None
        value = str(args[0])
        offset = max(0, int(args[1]))
        count = int(args[2]) if len(args) > 2 and args[2] is not None else None
        return value[offset:] if count is None else value[offset : offset + max(0, count)]
    if fn == "TEXT.RANGE":
        if args[0] is None:
            return None
        offset = int(args[1])
        count = int(args[2]) if len(args) > 2 and args[2] is not None else None
        return str(args[0])[offset:] if count is None else str(args[0])[offset : offset + count]
    if fn == "TEXT.PADSTART":
        pad = str(args[2]) if len(args) > 2 and args[2] is not None else " "
        return None if args[0] is None else str(args[0]).rjust(int(args[1]), pad[:1] or " ")
    if fn == "TEXT.PADEND":
        pad = str(args[2]) if len(args) > 2 and args[2] is not None else " "
        return None if args[0] is None else str(args[0]).ljust(int(args[1]), pad[:1] or " ")
    if fn == "TEXT.FORMAT":
        return _text_format(args[0] if args else None, args[1] if len(args) > 1 else [], args[2] if len(args) > 2 else None)
    if fn == "GUID.FROM":
        return _guid_from_value(args[0] if args else None)
    if fn == "TEXT.NEWGUID":
        return str(_uuid.uuid4())
    if fn == "TEXT.INFERNUMBERTYPE":
        try:
            number = _to_number(args[0] if args else None, args[1] if len(args) > 1 else None)
        except Exception:
            return "type text"
        if number is None:
            return "type nullable number"
        return "Int64.Type" if float(number).is_integer() else "Double.Type"
    if fn == "TEXT.FROM":
        return None if args[0] is None else str(args[0])
    if fn == "NUMBER.FROM":
        return _to_number(args[0], args[1] if len(args) > 1 else None)
    if fn == "NUMBER.FROMTEXT":
        return _to_number(args[0], args[1] if len(args) > 1 else None)
    if fn in {"BYTE.FROM", "INT8.FROM", "INT16.FROM", "INT32.FROM", "INT64.FROM"}:
        number = _to_number(args[0] if args else None, args[1] if len(args) > 1 else None)
        if number is None:
            return None
        value = int(number)
        ranges = {
            "BYTE.FROM": (0, 255),
            "INT8.FROM": (-128, 127),
            "INT16.FROM": (-32768, 32767),
            "INT32.FROM": (-2147483648, 2147483647),
            "INT64.FROM": (-9223372036854775808, 9223372036854775807),
        }
        low, high = ranges[fn]
        if value < low or value > high:
            raise ValueError(f"{fn} preview value is outside the supported range.")
        return value
    if fn in {"CURRENCY.FROM", "DECIMAL.FROM", "DOUBLE.FROM", "SINGLE.FROM"}:
        return _to_number(args[0] if args else None, args[1] if len(args) > 1 else None)
    if fn == "PERCENTAGE.FROM":
        value = args[0] if args else None
        if isinstance(value, str) and value.strip().endswith("%"):
            number = _to_number(value.strip()[:-1], args[1] if len(args) > 1 else None)
            return None if number is None else number / 100
        return _to_number(value, args[1] if len(args) > 1 else None)
    if fn == "NUMBER.TOTEXT":
        return _number_text(args[0] if args else None, args[1] if len(args) > 1 else None, args[2] if len(args) > 2 else None)
    if fn == "NUMBER.ROUND":
        number = _to_number(args[0])
        if number is None:
            return None
        digits = int(args[1]) if len(args) > 1 and args[1] is not None else 0
        return round(number, digits)
    if fn == "NUMBER.ABS":
        number = _to_number(args[0])
        return abs(number) if number is not None else None
    if fn == "NUMBER.SIGN":
        number = _to_number(args[0])
        if number is None:
            return None
        return 1 if number > 0 else -1 if number < 0 else 0
    if fn == "NUMBER.ROUNDUP":
        number = _to_number(args[0])
        if number is None:
            return None
        digits = int(args[1]) if len(args) > 1 and args[1] is not None else 0
        factor = 10**digits
        return _math.ceil(number * factor) / factor
    if fn == "NUMBER.ROUNDDOWN":
        number = _to_number(args[0])
        if number is None:
            return None
        digits = int(args[1]) if len(args) > 1 and args[1] is not None else 0
        factor = 10**digits
        return _math.floor(number * factor) / factor
    if fn == "NUMBER.ROUNDTOWARDZERO":
        number = _to_number(args[0])
        if number is None:
            return None
        digits = int(args[1]) if len(args) > 1 and args[1] is not None else 0
        factor = 10**digits
        rounded = _math.trunc(number * factor) / factor
        return int(rounded) if digits == 0 else rounded
    if fn == "NUMBER.ROUNDAWAYFROMZERO":
        number = _to_number(args[0])
        if number is None:
            return None
        digits = int(args[1]) if len(args) > 1 and args[1] is not None else 0
        factor = 10**digits
        rounded = (_math.ceil(number * factor) if number >= 0 else _math.floor(number * factor)) / factor
        return int(rounded) if digits == 0 else rounded
    if fn == "NUMBER.ISEVEN":
        number = _to_number(args[0])
        return False if number is None else int(number) % 2 == 0
    if fn == "NUMBER.ISODD":
        number = _to_number(args[0])
        return False if number is None else int(number) % 2 != 0
    if fn == "NUMBER.MOD":
        left = _to_number(args[0])
        right = _to_number(args[1])
        if left is None or right in {None, 0}:
            return None
        return left % right
    if fn == "NUMBER.POWER":
        left = _to_number(args[0])
        right = _to_number(args[1])
        return left**right if left is not None and right is not None else None
    if fn == "NUMBER.SQRT":
        number = _to_number(args[0])
        return _math.sqrt(number) if number is not None and number >= 0 else None
    if fn == "NUMBER.EXP":
        number = _math_number(args[0] if args else None)
        return _math.exp(number) if number is not None else None
    if fn == "NUMBER.LN":
        number = _math_number(args[0] if args else None)
        return _math.log(number) if number is not None and number > 0 else None
    if fn == "NUMBER.LOG":
        number = _math_number(args[0] if args else None)
        base = _math_number(args[1]) if len(args) > 1 else 10
        return _math.log(number, base) if number is not None and number > 0 and base not in {None, 0, 1} else None
    if fn == "NUMBER.LOG10":
        number = _math_number(args[0] if args else None)
        return _math.log10(number) if number is not None and number > 0 else None
    if fn == "NUMBER.SIN":
        number = _math_number(args[0] if args else None)
        return _math.sin(number) if number is not None else None
    if fn == "NUMBER.COS":
        number = _math_number(args[0] if args else None)
        return _math.cos(number) if number is not None else None
    if fn == "NUMBER.TAN":
        number = _math_number(args[0] if args else None)
        return _math.tan(number) if number is not None else None
    if fn == "NUMBER.ASIN":
        number = _math_number(args[0] if args else None)
        return _math.asin(number) if number is not None and -1 <= number <= 1 else None
    if fn == "NUMBER.ACOS":
        number = _math_number(args[0] if args else None)
        return _math.acos(number) if number is not None and -1 <= number <= 1 else None
    if fn == "NUMBER.ATAN":
        number = _math_number(args[0] if args else None)
        return _math.atan(number) if number is not None else None
    if fn == "NUMBER.ATAN2":
        y = _math_number(args[0] if args else None)
        x = _math_number(args[1] if len(args) > 1 else None)
        return _math.atan2(y, x) if y is not None and x is not None else None
    if fn == "NUMBER.SINH":
        number = _math_number(args[0] if args else None)
        return _math.sinh(number) if number is not None else None
    if fn == "NUMBER.COSH":
        number = _math_number(args[0] if args else None)
        return _math.cosh(number) if number is not None else None
    if fn == "NUMBER.TANH":
        number = _math_number(args[0] if args else None)
        return _math.tanh(number) if number is not None else None
    if fn == "NUMBER.FACTORIAL":
        number = _math_number(args[0] if args else None)
        return _math.factorial(int(number)) if number is not None and number >= 0 and float(number).is_integer() else None
    if fn == "NUMBER.COMBINATIONS":
        n = _int_number(args[0] if args else 0)
        k = _int_number(args[1] if len(args) > 1 else 0)
        return _math.comb(n, k) if 0 <= k <= n else 0
    if fn == "NUMBER.PERMUTATIONS":
        n = _int_number(args[0] if args else 0)
        k = _int_number(args[1] if len(args) > 1 else 0)
        return _math.perm(n, k) if 0 <= k <= n else 0
    if fn == "NUMBER.ISNAN":
        number = _math_number(args[0] if args else None)
        return False if number is None else _math.isnan(number)
    if fn == "NUMBER.ISINFINITY":
        number = _math_number(args[0] if args else None)
        return False if number is None else _math.isinf(number)
    if fn == "NUMBER.BITWISEAND":
        return _int_number(args[0] if args else 0) & _int_number(args[1] if len(args) > 1 else 0)
    if fn == "NUMBER.BITWISEOR":
        return _int_number(args[0] if args else 0) | _int_number(args[1] if len(args) > 1 else 0)
    if fn == "NUMBER.BITWISEXOR":
        return _int_number(args[0] if args else 0) ^ _int_number(args[1] if len(args) > 1 else 0)
    if fn == "NUMBER.BITWISENOT":
        return ~_int_number(args[0] if args else 0)
    if fn == "NUMBER.BITWISESHIFTLEFT":
        return _int_number(args[0] if args else 0) << _int_number(args[1] if len(args) > 1 else 0)
    if fn == "NUMBER.BITWISESHIFTRIGHT":
        return _int_number(args[0] if args else 0) >> _int_number(args[1] if len(args) > 1 else 0)
    if fn == "NUMBER.INTEGERDIVIDE":
        left = _to_number(args[0])
        right = _to_number(args[1])
        if left is None or right in {None, 0}:
            return None
        return int(left // right)
    if fn == "NUMBER.RANDOM":
        return _random.random()
    if fn == "NUMBER.RANDOMBETWEEN":
        bottom = int(_to_number(args[0] if args else 0) or 0)
        top = int(_to_number(args[1] if len(args) > 1 else bottom) or bottom)
        if top < bottom:
            bottom, top = top, bottom
        return _random.randint(bottom, top)
    if fn == "DATE.FROM":
        return _to_date(args[0], args[1] if len(args) > 1 else None)
    if fn == "DATE.FROMTEXT":
        return _to_date(args[0], args[1] if len(args) > 1 else None)
    if fn.startswith("DATE.ISIN"):
        return _date_is_in_period(fn, args[0] if args else None, args[1] if len(args) > 1 else None)
    if fn == "DATETIME.FROM":
        return _to_datetime(args[0], args[1] if len(args) > 1 else None)
    if fn == "DATETIME.FROMTEXT":
        return _to_datetime(args[0], args[1] if len(args) > 1 else None)
    if fn == "DATETIME.LOCALNOW" or fn == "DATETIME.FIXEDLOCALNOW":
        return _dt.datetime.now().replace(microsecond=0)
    if fn.startswith("DATETIME.ISIN"):
        return _datetime_is_in_period(fn, args[0] if args else None, args[1] if len(args) > 1 else None)
    if fn == "DATETIME.TOTEXT":
        return _datetime_to_text(args[0] if args else None, args[1] if len(args) > 1 else None, args[2] if len(args) > 2 else None)
    if fn == "DATETIME.TORECORD":
        return _datetime_record(args[0] if args else None)
    if fn == "DATETIME.ADDZONE":
        value = _to_datetime(args[0] if args else None)
        return None if value is None else value.replace(tzinfo=_timezone_from_parts(args[1] if len(args) > 1 else 0, args[2] if len(args) > 2 else 0))
    if fn == "DATETIME.FROMFILETIME":
        return _datetime_from_filetime(args[0] if args else None)
    if fn == "DATETIME.DATE":
        return _to_date(args[0] if args else None)
    if fn == "DATETIME.TIME":
        return _to_time(args[0] if args else None)
    if fn == "DATETIMEZONE.FROM":
        return _to_datetimezone(args[0], args[1] if len(args) > 1 else None)
    if fn == "DATETIMEZONE.FROMTEXT":
        return _to_datetimezone(args[0], args[1] if len(args) > 1 else None)
    if fn == "DATETIMEZONE.LOCALNOW" or fn == "DATETIMEZONE.FIXEDLOCALNOW":
        return _dt.datetime.now().astimezone().replace(microsecond=0)
    if fn == "DATETIMEZONE.UTCNOW" or fn == "DATETIMEZONE.FIXEDUTCNOW":
        return _dt.datetime.now(_dt.timezone.utc).replace(microsecond=0)
    if fn == "DATETIMEZONE.TOTEXT":
        return _datetimezone_to_text(args[0] if args else None, args[1] if len(args) > 1 else None, args[2] if len(args) > 2 else None)
    if fn == "DATETIMEZONE.TORECORD":
        return _datetimezone_record(args[0] if args else None)
    if fn == "DATETIMEZONE.FROMFILETIME":
        return _datetime_from_filetime(args[0] if args else None, timezone=True)
    if fn == "DATETIMEZONE.TOUTC":
        value = _to_datetimezone(args[0] if args else None)
        return value.astimezone(_dt.timezone.utc) if value is not None else None
    if fn == "DATETIMEZONE.TOLOCAL":
        value = _to_datetimezone(args[0] if args else None)
        return value.astimezone() if value is not None else None
    if fn == "DATETIMEZONE.SWITCHZONE":
        value = _to_datetimezone(args[0] if args else None)
        if value is None:
            return None
        return value.astimezone(_timezone_from_parts(args[1] if len(args) > 1 else 0, args[2] if len(args) > 2 else 0))
    if fn == "DATETIMEZONE.REMOVEZONE":
        value = _to_datetimezone(args[0] if args else None)
        return value.replace(tzinfo=None) if value is not None else None
    if fn == "DATETIMEZONE.ZONEHOURS":
        value = _to_datetimezone(args[0] if args else None)
        if value is None or value.utcoffset() is None:
            return None
        return int(value.utcoffset().total_seconds() // 3600)
    if fn == "DATETIMEZONE.ZONEMINUTES":
        value = _to_datetimezone(args[0] if args else None)
        if value is None or value.utcoffset() is None:
            return None
        minutes = int(abs(value.utcoffset().total_seconds()) // 60)
        return minutes % 60
    if fn == "LOGICAL.FROM":
        return _to_logical(args[0])
    if fn == "LOGICAL.FROMTEXT":
        return _to_logical(args[0])
    if fn == "LOGICAL.TOTEXT":
        value = _to_logical(args[0] if args else None)
        return None if value is None else ("true" if value else "false")
    if fn == "VALUE.ADD":
        return _eval_add_subtract(args[0] if args else None, "+", args[1] if len(args) > 1 else None)
    if fn == "VALUE.SUBTRACT":
        return _eval_add_subtract(args[0] if args else None, "-", args[1] if len(args) > 1 else None)
    if fn == "VALUE.MULTIPLY":
        left = _to_number(args[0] if args else None)
        right = _to_number(args[1] if len(args) > 1 else None)
        return None if left is None or right is None else left * right
    if fn == "VALUE.DIVIDE":
        left = _to_number(args[0] if args else None)
        right = _to_number(args[1] if len(args) > 1 else None)
        return None if left is None or right in {None, 0} else left / right
    if fn == "VALUE.COMPARE":
        left = args[0] if args else None
        right = args[1] if len(args) > 1 else None
        if left == right:
            return 0
        try:
            return -1 if left < right else 1
        except Exception:
            return -1 if str(left) < str(right) else 1
    if fn == "VALUE.AS":
        value = args[0] if args else None
        type_value = args[1] if len(args) > 1 else "type any"
        if _value_matches_type(value, type_value):
            return value
        raise ValueError(f"Value.As preview cannot coerce value to {type_value!r}.")
    if fn == "VALUE.FROMTEXT":
        text_value = args[0] if args else None
        if text_value is None:
            return None
        text = str(text_value).strip()
        for parser in (_to_logical, _to_number, _to_datetime, _to_date, _to_time):
            try:
                parsed = parser(text)
                if parsed is not None:
                    return parsed
            except Exception:
                pass
        return text
    if fn == "VALUE.EXPRESSION":
        return _m_literal_text(args[0] if args else None)
    if fn == "VALUE.OPTIMIZE":
        return args[0] if args else None
    if fn == "VALUE.ALTERNATES":
        primary = args[0] if args else None
        alternates = args[1] if len(args) > 1 and isinstance(args[1], list) else []
        return [primary] + alternates
    if fn == "VALUE.TRAITS":
        return _metadata_record(args[0] if args else None).get("Traits", {})
    if fn == "VALUE.VERSIONIDENTITY":
        return _metadata_record(args[0] if args else None).get("VersionIdentity")
    if fn == "VALUE.VERSIONS":
        return _metadata_record(args[0] if args else None).get("Versions", [])
    if fn == "VALUE.LINEAGE":
        return _metadata_record(args[0] if args else None).get("Lineage")
    if fn == "VALUE.FIREWALL":
        return args[0] if args else None
    if fn == "ACTION.WITHERRORCONTEXT":
        return args[0] if args else None
    if fn == "DIRECTQUERYCAPABILITIES.FROM":
        return dict(args[0]) if args and isinstance(args[0], Mapping) else {}
    if fn == "EMBEDDED.VALUE":
        return args[0] if args else None
    if fn == "EXCEL.SHAPETABLE":
        return args[0] if args and isinstance(args[0], list) else []
    if fn == "GRAPH.NODES":
        return args[0] if args and isinstance(args[0], list) else []
    if fn == "MODULE.VERSIONS":
        return []
    if fn == "PROGRESS.DATASOURCEPROGRESS":
        return args[0] if args else None
    if fn == "SQLEXPRESSION.SCHEMAFROM":
        return dict(args[0]) if args and isinstance(args[0], Mapping) else {"Expression": args[0] if args else None}
    if fn == "SQLEXPRESSION.TOEXPRESSION":
        return str(args[0] if args else "")
    if fn == "VARIABLE.VALUE":
        name_value = str(args[0] if args else "")
        if variables and name_value in variables:
            return variables[name_value]
        raise KeyError(name_value)
    if fn == "VARIABLE.VALUEORDEFAULT":
        name_value = str(args[0] if args else "")
        if variables and name_value in variables:
            return variables[name_value]
        return args[1] if len(args) > 1 else None
    if fn == "VALUE.METADATA":
        return _metadata_record(args[0] if args else None)
    if fn == "VALUE.REPLACEMETADATA":
        value = _unwrap_metadata(args[0] if args else None)
        metadata = args[1] if len(args) > 1 and isinstance(args[1], Mapping) else {}
        return MMetadataValue(value=value, metadata=dict(metadata))
    if fn == "VALUE.REMOVEMETADATA":
        value = args[0] if args else None
        if not isinstance(value, MMetadataValue):
            return value
        if len(args) <= 1 or args[1] is None:
            return value.value
        remove_fields = _record_field_names(args[1])
        metadata = {
            key: item
            for key, item in value.metadata.items()
            if key not in remove_fields
        }
        return MMetadataValue(value=value.value, metadata=metadata) if metadata else value.value
    if fn == "VALUE.TYPE":
        return _m_type_text(args[0] if args else None)
    if fn == "VALUE.IS":
        return _value_matches_type(args[0] if args else None, args[1] if len(args) > 1 else "type any")
    if fn == "TYPE.IS":
        return _type_compatible(args[0] if args else "type none", args[1] if len(args) > 1 else "type any")
    if fn == "TYPE.NONNULLABLE":
        return str(args[0] if args else "type any").replace("type nullable ", "type ")
    if fn == "TYPE.ISNULLABLE":
        return "nullable" in str(args[0] if args else "").lower() or str(args[0] if args else "").strip().lower() in {"type any", "any.type"}
    if fn == "TYPE.CLOSEDRECORD":
        return MMetadataValue(value=str(args[0] if args else "type []"), metadata={**_metadata_record(args[0] if args else None), "Open": False})
    if fn == "TYPE.OPENRECORD":
        return MMetadataValue(value=str(args[0] if args else "type []"), metadata={**_metadata_record(args[0] if args else None), "Open": True})
    if fn == "TYPE.ISOPENRECORD":
        metadata = _metadata_record(args[0] if args else None)
        return bool(metadata.get("Open", True))
    if fn == "TYPE.RECORDFIELDS":
        return _type_record_fields(args[0] if args else "type []")
    if fn == "TYPE.FORRECORD":
        return _type_for_record(args[0] if args else {}, args[1] if len(args) > 1 else False)
    if fn == "TYPE.FORFUNCTION":
        return _type_for_function(args[0] if args else {}, args[1] if len(args) > 1 else "type any")
    if fn == "TYPE.FUNCTIONPARAMETERS":
        return _metadata_record(args[0] if args else None).get("Parameters", {})
    if fn == "TYPE.FUNCTIONREQUIREDPARAMETERS":
        params = _metadata_record(args[0] if args else None).get("Parameters", {})
        return sum(1 for spec in params.values() if not (isinstance(spec, Mapping) and bool(spec.get("Optional") or spec.get("optional")))) if isinstance(params, Mapping) else 0
    if fn == "TYPE.FUNCTIONRETURN":
        return _metadata_record(args[0] if args else None).get("ReturnType", "type any")
    if fn == "TYPE.LISTITEM":
        text = str(args[0] if args else "type list")
        match = re.search(r"\{(.+)\}", text)
        return match.group(1).strip() if match else "type any"
    if fn == "TYPE.TABLECOLUMN":
        fields = _type_record_fields(args[0] if args else "type table []")
        name = str(args[1] if len(args) > 1 else "")
        return (fields.get(name) or {}).get("Type")
    if fn == "TYPE.TABLEROW":
        fields = _type_record_fields(args[0] if args else "type table []")
        return _type_for_record(fields, False)
    if fn == "TYPE.TABLESCHEMA":
        return _type_table_schema(args[0] if args else "type table []")
    if fn == "TYPE.TABLEKEYS":
        return _type_table_keys(args[0] if args else None)
    if fn == "TYPE.ADDTABLEKEY":
        metadata = _metadata_record(args[0] if args else None)
        keys = list(metadata.get("Keys") or [])
        keys.extend(_table_key_records_from_args(args[1] if len(args) > 1 else [], args[2] if len(args) > 2 else False))
        return MMetadataValue(value=_unwrap_metadata(args[0] if args else "type table []"), metadata={**metadata, "Keys": keys})
    if fn == "TYPE.REPLACETABLEKEYS":
        metadata = _metadata_record(args[0] if args else None)
        keys = args[1] if len(args) > 1 and isinstance(args[1], list) else []
        return MMetadataValue(value=_unwrap_metadata(args[0] if args else "type table []"), metadata={**metadata, "Keys": keys})
    if fn == "TYPE.TABLEPARTITIONKEY":
        return _metadata_record(args[0] if args else None).get("PartitionKey")
    if fn == "TYPE.REPLACETABLEPARTITIONKEY":
        metadata = _metadata_record(args[0] if args else None)
        return MMetadataValue(value=_unwrap_metadata(args[0] if args else "type table []"), metadata={**metadata, "PartitionKey": args[1] if len(args) > 1 else None})
    if fn == "TYPE.FACETS":
        return _metadata_record(args[0] if args else None).get("Facets", {})
    if fn == "TYPE.REPLACEFACETS":
        metadata = _metadata_record(args[0] if args else None)
        return MMetadataValue(value=_unwrap_metadata(args[0] if args else "type any"), metadata={**metadata, "Facets": args[1] if len(args) > 1 and isinstance(args[1], Mapping) else {}})
    if fn == "TYPE.UNION":
        types = args[0] if args and isinstance(args[0], list) else []
        return "type union(" + ", ".join(str(item) for item in types) + ")"
    if fn == "VALUE.EQUALS":
        return (args[0] if args else None) == (args[1] if len(args) > 1 else None)
    if fn == "VALUE.NULLABLEEQUALS":
        if not args or len(args) < 2 or args[0] is None or args[1] is None:
            return None
        return args[0] == args[1]
    if fn == "TIME.HOUR":
        value = _to_time(args[0] if args else None)
        return value.hour if value else None
    if fn == "TIME.MINUTE":
        value = _to_time(args[0] if args else None)
        return value.minute if value else None
    if fn == "TIME.SECOND":
        value = _to_time(args[0] if args else None)
        return value.second if value else None
    if fn == "TIME.FROM":
        return _to_time(args[0] if args else None, args[1] if len(args) > 1 else None)
    if fn == "TIME.FROMTEXT":
        return _to_time(args[0] if args else None, args[1] if len(args) > 1 else None)
    if fn == "TIME.TOTEXT":
        return _time_to_text(args[0] if args else None, args[1] if len(args) > 1 else None, args[2] if len(args) > 2 else None)
    if fn == "TIME.TORECORD":
        return _time_record(args[0] if args else None)
    if fn == "TIME.STARTOFHOUR":
        value = _to_time(args[0] if args else None)
        return value.replace(minute=0, second=0, microsecond=0) if value else None
    if fn == "TIME.ENDOFHOUR":
        value = _to_time(args[0] if args else None)
        return value.replace(minute=59, second=59, microsecond=999999) if value else None
    if fn == "DATE.YEAR":
        value = _to_date(args[0])
        return value.year if value else None
    if fn == "DATE.MONTH":
        value = _to_date(args[0])
        return value.month if value else None
    if fn == "DATE.DAY":
        value = _to_date(args[0])
        return value.day if value else None
    if fn == "DATE.TOTEXT":
        return _date_to_text(args[0] if args else None, args[1] if len(args) > 1 else None, args[2] if len(args) > 2 else None)
    if fn == "DATE.TORECORD":
        return _date_record(args[0] if args else None)
    if fn == "DATE.QUARTEROFYEAR":
        value = _to_date(args[0])
        return ((value.month - 1) // 3 + 1) if value else None
    if fn == "DATE.DAYOFWEEK":
        value = _to_date(args[0])
        return ((value.weekday() - _day_first_index(args[1] if len(args) > 1 else None)) % 7) if value else None
    if fn == "DATE.MONTHNAME":
        value = _to_date(args[0])
        return _localized_month_names(args[1] if len(args) > 1 else None)[value.month - 1] if value else None
    if fn == "DATE.DAYOFWEEKNAME":
        value = _to_date(args[0])
        return _localized_weekday_names(args[1] if len(args) > 1 else None)[value.weekday()] if value else None
    if fn == "DATE.DAYOFYEAR":
        value = _to_date(args[0])
        return value.timetuple().tm_yday if value else None
    if fn == "DATE.WEEKOFYEAR":
        value = _to_date(args[0])
        return _week_index(value, _dt.date(value.year, 1, 1), args[1] if len(args) > 1 else None) if value else None
    if fn == "DATE.WEEKOFMONTH":
        value = _to_date(args[0])
        return _week_index(value, _dt.date(value.year, value.month, 1), args[1] if len(args) > 1 else None) if value else None
    if fn == "DATE.ADDDAYS":
        value = _to_date(args[0])
        return value + _dt.timedelta(days=int(args[1])) if value else None
    if fn == "DATE.ADDWEEKS":
        value = _to_date(args[0])
        return value + _dt.timedelta(days=7 * int(args[1])) if value else None
    if fn == "DATE.ADDMONTHS":
        value = _to_date(args[0])
        if not value:
            return None
        month_index = value.month - 1 + int(args[1])
        year = value.year + month_index // 12
        month = month_index % 12 + 1
        month_lengths = [31, 29 if year % 4 == 0 and (year % 100 != 0 or year % 400 == 0) else 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31]
        return _dt.date(year, month, min(value.day, month_lengths[month - 1]))
    if fn == "DATE.ADDQUARTERS":
        return _eval_function_call("Date.AddMonths", [args[0] if args else None, int(args[1]) * 3 if len(args) > 1 else 0])
    if fn == "DATE.ADDYEARS":
        value = _to_date(args[0])
        if not value:
            return None
        year = value.year + int(args[1])
        try:
            return value.replace(year=year)
        except ValueError:
            return value.replace(year=year, day=28)
    if fn == "DATE.DAYSINMONTH":
        value = _to_date(args[0] if args else None)
        if not value:
            return None
        next_month = _eval_function_call("Date.AddMonths", [value.replace(day=1), 1])
        return (next_month - _dt.timedelta(days=1)).day
    if fn == "DATE.ISLEAPYEAR":
        value = _to_date(args[0] if args else None)
        return False if not value else value.year % 4 == 0 and (value.year % 100 != 0 or value.year % 400 == 0)
    if fn == "DATE.STARTOFDAY":
        value = args[0] if args else None
        if isinstance(value, _dt.datetime):
            return value.replace(hour=0, minute=0, second=0, microsecond=0)
        return _to_date(value)
    if fn == "DATE.ENDOFDAY":
        value = args[0] if args else None
        if isinstance(value, _dt.datetime):
            return value.replace(hour=23, minute=59, second=59, microsecond=999999)
        return _to_date(value)
    if fn == "DATE.STARTOFWEEK":
        value = _to_date(args[0])
        return value - _dt.timedelta(days=((value.weekday() - _day_first_index(args[1] if len(args) > 1 else None)) % 7)) if value else None
    if fn == "DATE.ENDOFWEEK":
        value = _to_date(args[0])
        if not value:
            return None
        start = _eval_function_call("Date.StartOfWeek", [value, args[1] if len(args) > 1 else None])
        return start + _dt.timedelta(days=6)
    if fn == "DATE.STARTOFMONTH":
        value = _to_date(args[0])
        return _dt.date(value.year, value.month, 1) if value else None
    if fn == "DATE.ENDOFMONTH":
        value = _to_date(args[0])
        if not value:
            return None
        next_month = _eval_function_call("Date.AddMonths", [value.replace(day=1), 1])
        return next_month - _dt.timedelta(days=1)
    if fn == "DATE.STARTOFQUARTER":
        value = _to_date(args[0])
        if not value:
            return None
        month = ((value.month - 1) // 3) * 3 + 1
        return _dt.date(value.year, month, 1)
    if fn == "DATE.ENDOFQUARTER":
        value = _to_date(args[0])
        if not value:
            return None
        start = _eval_function_call("Date.StartOfQuarter", [value])
        next_quarter = _eval_function_call("Date.AddMonths", [start, 3])
        return next_quarter - _dt.timedelta(days=1)
    if fn == "DATE.STARTOFYEAR":
        value = _to_date(args[0])
        return _dt.date(value.year, 1, 1) if value else None
    if fn == "DATE.ENDOFYEAR":
        value = _to_date(args[0])
        return _dt.date(value.year, 12, 31) if value else None
    if fn == "DURATION.FROM":
        return _to_duration(args[0] if args else None)
    if fn == "DURATION.FROMTEXT":
        return _duration_from_text(args[0] if args else None)
    if fn == "DURATION.TORECORD":
        return _duration_record(args[0] if args else None)
    if fn == "DURATION.DAYS":
        value = _to_duration(args[0] if args else None)
        return value.days if value is not None else None
    if fn == "DURATION.HOURS":
        value = _to_duration(args[0] if args else None)
        return (value.seconds // 3600) if value is not None else None
    if fn == "DURATION.MINUTES":
        value = _to_duration(args[0] if args else None)
        return ((value.seconds % 3600) // 60) if value is not None else None
    if fn == "DURATION.SECONDS":
        value = _to_duration(args[0] if args else None)
        return ((value.seconds % 60) + value.microseconds / 1_000_000) if value is not None else None
    if fn == "DURATION.TOTALDAYS":
        value = _to_duration(args[0] if args else None)
        return value.total_seconds() / 86400 if value is not None else None
    if fn == "DURATION.TOTALHOURS":
        value = _to_duration(args[0] if args else None)
        return value.total_seconds() / 3600 if value is not None else None
    if fn == "DURATION.TOTALMINUTES":
        value = _to_duration(args[0] if args else None)
        return value.total_seconds() / 60 if value is not None else None
    if fn == "DURATION.TOTALSECONDS":
        value = _to_duration(args[0] if args else None)
        return value.total_seconds() if value is not None else None
    if fn == "DURATION.TOTEXT":
        return _duration_to_text(args[0] if args else None, args[1] if len(args) > 1 else None)
    if fn == "LIST.DATES":
        start = _to_date(args[0] if args else None)
        count = int(_to_number(args[1]) or 0) if len(args) > 1 else 0
        step = _to_duration(args[2] if len(args) > 2 else _dt.timedelta(days=1)) or _dt.timedelta(days=1)
        if start is None or count <= 0:
            return []
        return [start + (step * idx) for idx in range(count)]
    if fn == "LIST.DATETIMES":
        start = _to_datetime(args[0] if args else None)
        count = int(_to_number(args[1]) or 0) if len(args) > 1 else 0
        step = _to_duration(args[2] if len(args) > 2 else _dt.timedelta(days=1)) or _dt.timedelta(days=1)
        if start is None or count <= 0:
            return []
        return [start + (step * idx) for idx in range(count)]
    if fn == "LIST.DATETIMEZONES":
        start = _to_datetimezone(args[0] if args else None)
        count = int(_to_number(args[1]) or 0) if len(args) > 1 else 0
        step = _to_duration(args[2] if len(args) > 2 else _dt.timedelta(days=1)) or _dt.timedelta(days=1)
        if start is None or count <= 0:
            return []
        return [start + (step * idx) for idx in range(count)]
    if fn == "LIST.DURATIONS":
        start = _to_duration(args[0] if args else None)
        count = int(_to_number(args[1]) or 0) if len(args) > 1 else 0
        step = _to_duration(args[2] if len(args) > 2 else _dt.timedelta(days=1)) or _dt.timedelta(days=1)
        if start is None or count <= 0:
            return []
        return [start + (step * idx) for idx in range(count)]
    if fn == "LIST.TIMES":
        start = _to_time(args[0] if args else None)
        count = int(_to_number(args[1]) or 0) if len(args) > 1 else 0
        step = _to_duration(args[2] if len(args) > 2 else _dt.timedelta(hours=1)) or _dt.timedelta(hours=1)
        if start is None or count <= 0:
            return []
        base = _dt.datetime.combine(_dt.date(2000, 1, 1), start)
        return [(base + (step * idx)).time() for idx in range(count)]
    if fn == "LIST.NUMBERS":
        start = _to_number(args[0] if args else 0) or 0
        count = int(_to_number(args[1]) or 0) if len(args) > 1 else 0
        increment = _to_number(args[2]) if len(args) > 2 else 1
        increment = 1 if increment is None else increment
        values: list[Any] = []
        for idx in range(max(0, count)):
            value = start + (increment * idx)
            values.append(int(value) if float(value).is_integer() else value)
        return values
    if fn == "LIST.RANDOM":
        count = max(0, int(_to_number(args[0] if args else 0) or 0))
        return [_random.random() for _idx in range(count)]
    if fn == "LIST.COUNT":
        return len(args[0]) if isinstance(args[0], list) else 0
    if fn == "LIST.NONNULLCOUNT":
        return sum(1 for item in args[0] if item is not None) if args and isinstance(args[0], list) else 0
    if fn == "LIST.SUM":
        return sum(_numeric_values(args[0] if isinstance(args[0], list) else []))
    if fn == "LIST.PRODUCT":
        values = _numeric_values(args[0] if isinstance(args[0], list) else [])
        if not values:
            return None
        product = 1.0
        for value in values:
            product *= value
        return product
    if fn == "LIST.MIN":
        values = _numeric_values(args[0] if isinstance(args[0], list) else [])
        return min(values) if values else None
    if fn == "LIST.MAX":
        values = _numeric_values(args[0] if isinstance(args[0], list) else [])
        return max(values) if values else None
    if fn == "LIST.AVERAGE":
        values = _numeric_values(args[0] if isinstance(args[0], list) else [])
        return sum(values) / len(values) if values else None
    if fn == "LIST.MEDIAN":
        return _list_median(args[0] if args and isinstance(args[0], list) else [])
    if fn == "LIST.PERCENTILE":
        return _list_percentile(args[0] if args and isinstance(args[0], list) else [], args[1] if len(args) > 1 else 0.5)
    if fn == "LIST.MODE":
        return _list_mode(args[0] if args and isinstance(args[0], list) else [], return_all=False)
    if fn == "LIST.MODES":
        return _list_mode(args[0] if args and isinstance(args[0], list) else [], return_all=True)
    if fn == "LIST.STANDARDDEVIATION":
        values = _numeric_values(args[0] if isinstance(args[0], list) else [])
        if len(values) < 2:
            return None
        average = sum(values) / len(values)
        variance = sum((value - average) ** 2 for value in values) / (len(values) - 1)
        return _math.sqrt(variance)
    if fn == "LIST.VARIANCE":
        values = _numeric_values(args[0] if isinstance(args[0], list) else [])
        if len(values) < 2:
            return None
        average = sum(values) / len(values)
        return sum((value - average) ** 2 for value in values) / (len(values) - 1)
    if fn == "LIST.COVARIANCE":
        left_values = _numeric_values(args[0] if args and isinstance(args[0], list) else [])
        right_values = _numeric_values(args[1] if len(args) > 1 and isinstance(args[1], list) else [])
        count = min(len(left_values), len(right_values))
        if count < 2:
            return None
        left_values = left_values[:count]
        right_values = right_values[:count]
        left_average = sum(left_values) / count
        right_average = sum(right_values) / count
        return sum((left_values[idx] - left_average) * (right_values[idx] - right_average) for idx in range(count)) / (count - 1)
    if fn == "LIST.DISTINCT":
        seen: set[Any] = set()
        out: list[Any] = []
        for item in args[0] if isinstance(args[0], list) else []:
            key = _value_key(item)
            if key in seen:
                continue
            seen.add(key)
            out.append(item)
        return out
    if fn == "LIST.DIFFERENCE":
        values = list(args[0]) if args and isinstance(args[0], list) else []
        excluded = {_value_key(item) for item in args[1]} if len(args) > 1 and isinstance(args[1], list) else set()
        return [item for item in values if _value_key(item) not in excluded]
    if fn == "LIST.INTERSECT":
        lists = [item for item in args[0] if isinstance(item, list)] if args and isinstance(args[0], list) else []
        if not lists:
            return []
        common = set(_value_key(item) for item in lists[0])
        for values in lists[1:]:
            common &= {_value_key(item) for item in values}
        out: list[Any] = []
        seen: set[Any] = set()
        for item in lists[0]:
            key = _value_key(item)
            if key in common and key not in seen:
                seen.add(key)
                out.append(item)
        return out
    if fn == "LIST.UNION":
        lists = [item for item in args[0] if isinstance(item, list)] if args and isinstance(args[0], list) else []
        out: list[Any] = []
        seen: set[Any] = set()
        for values in lists:
            for item in values:
                key = _value_key(item)
                if key in seen:
                    continue
                seen.add(key)
                out.append(item)
        return out
    if fn == "LIST.ISDISTINCT":
        values = args[0] if args and isinstance(args[0], list) else []
        seen: set[Any] = set()
        for item in values:
            key = _value_key(item)
            if key in seen:
                return False
            seen.add(key)
        return True
    if fn == "LIST.REMOVENULLS":
        return [item for item in args[0] if item is not None] if isinstance(args[0], list) else []
    if fn == "LIST.COMBINE":
        lists = args[0] if args and isinstance(args[0], list) else []
        out: list[Any] = []
        for item in lists:
            if isinstance(item, list):
                out.extend(item)
        return out
    if fn == "LIST.ZIP":
        lists = [item for item in args[0] if isinstance(item, list)] if args and isinstance(args[0], list) else []
        width = max((len(item) for item in lists), default=0)
        return [[item[idx] if idx < len(item) else None for item in lists] for idx in range(width)]
    if fn == "LIST.ISEMPTY":
        return not bool(args[0]) if args and isinstance(args[0], list) else True
    if fn == "LIST.POSITIONS":
        values = args[0] if args and isinstance(args[0], list) else []
        return list(range(len(values)))
    if fn == "LIST.POSITIONOF":
        values = args[0] if args and isinstance(args[0], list) else []
        wanted = args[1] if len(args) > 1 else None
        wanted_key = _value_key(wanted)
        positions = [idx for idx, item in enumerate(values) if _value_key(item) == wanted_key]
        occurrence = str(args[2]).strip().upper() if len(args) > 2 and args[2] is not None else "OCCURRENCE.FIRST"
        if occurrence.endswith("ALL"):
            return positions
        if occurrence.endswith("LAST"):
            return positions[-1] if positions else -1
        return positions[0] if positions else -1
    if fn == "LIST.POSITIONOFANY":
        values = args[0] if args and isinstance(args[0], list) else []
        candidates = {_value_key(item) for item in args[1]} if len(args) > 1 and isinstance(args[1], list) else set()
        positions = [idx for idx, item in enumerate(values) if _value_key(item) in candidates]
        occurrence = str(args[2]).strip().upper() if len(args) > 2 and args[2] is not None else "OCCURRENCE.FIRST"
        if occurrence.endswith("ALL"):
            return positions
        if occurrence.endswith("LAST"):
            return positions[-1] if positions else -1
        return positions[0] if positions else -1
    if fn == "LIST.BUFFER":
        return list(args[0]) if args and isinstance(args[0], list) else []
    if fn == "LIST.REVERSE":
        return list(reversed(args[0])) if args and isinstance(args[0], list) else []
    if fn == "LIST.SORT":
        values = list(args[0]) if args and isinstance(args[0], list) else []
        order = str(args[1]).strip().upper() if len(args) > 1 and args[1] is not None else "ORDER.ASCENDING"
        return sorted(values, key=lambda item: (item is None, item), reverse=order.endswith("DESCENDING"))
    if fn == "LIST.ALTERNATE":
        values = list(args[0]) if args and isinstance(args[0], list) else []
        skip_count = max(0, int(_to_number(args[1]) or 0)) if len(args) > 1 else 0
        repeat_interval = int(_to_number(args[2]) or 0) if len(args) > 2 and args[2] is not None else None
        offset = max(0, int(_to_number(args[3]) or 0)) if len(args) > 3 and args[3] is not None else 0
        out = values[:offset]
        idx = offset + skip_count
        if repeat_interval is None:
            out.extend(values[idx:])
            return out
        take = max(0, repeat_interval)
        while idx < len(values):
            out.extend(values[idx : idx + take])
            idx += take + skip_count
        return out
    if fn == "LIST.FINDTEXT":
        values = args[0] if args and isinstance(args[0], list) else []
        needle = str(args[1] if len(args) > 1 and args[1] is not None else "")
        out: list[Any] = []
        for item in values:
            haystacks = item.values() if isinstance(item, Mapping) else [item]
            if any(needle in str(value) for value in haystacks if value is not None):
                out.append(item)
        return out
    if fn == "LIST.FIRST":
        values = args[0] if isinstance(args[0], list) else []
        return values[0] if values else (args[1] if len(args) > 1 else None)
    if fn == "LIST.LAST":
        values = args[0] if isinstance(args[0], list) else []
        return values[-1] if values else (args[1] if len(args) > 1 else None)
    if fn == "LIST.SINGLE":
        values = args[0] if args and isinstance(args[0], list) else []
        if len(values) == 1:
            return values[0]
        raise ValueError("List.Single preview needs exactly one item.")
    if fn == "LIST.SINGLEORDEFAULT":
        values = args[0] if args and isinstance(args[0], list) else []
        default_value = args[1] if len(args) > 1 else None
        if not values:
            return default_value
        if len(values) == 1:
            return values[0]
        raise ValueError("List.SingleOrDefault preview needs zero or one item.")
    if fn == "LIST.CONTAINS":
        return args[1] in args[0] if isinstance(args[0], list) and len(args) > 1 else False
    if fn == "LIST.CONTAINSANY":
        values = args[0] if args and isinstance(args[0], list) else []
        candidates = args[1] if len(args) > 1 and isinstance(args[1], list) else []
        return any(item in values for item in candidates)
    if fn == "LIST.CONTAINSALL":
        values = args[0] if args and isinstance(args[0], list) else []
        candidates = args[1] if len(args) > 1 and isinstance(args[1], list) else []
        return all(item in values for item in candidates)
    if fn == "LIST.FIRSTN":
        values = args[0] if args and isinstance(args[0], list) else []
        count = int(_to_number(args[1]) or 0) if len(args) > 1 else 1
        return values[: max(0, count)]
    if fn == "LIST.LASTN":
        values = args[0] if args and isinstance(args[0], list) else []
        count = int(_to_number(args[1]) or 0) if len(args) > 1 else 1
        return values[-count:] if count > 0 else []
    if fn in {"LIST.SKIP", "LIST.REMOVEFIRSTN"}:
        values = args[0] if args and isinstance(args[0], list) else []
        count = int(_to_number(args[1]) or 0) if len(args) > 1 else 1
        return values[max(0, count) :]
    if fn == "LIST.REMOVELASTN":
        values = args[0] if args and isinstance(args[0], list) else []
        count = int(_to_number(args[1]) or 0) if len(args) > 1 else 1
        return values[:-count] if count > 0 else list(values)
    if fn == "LIST.RANGE":
        values = args[0] if args and isinstance(args[0], list) else []
        offset = max(0, int(_to_number(args[1]) or 0)) if len(args) > 1 else 0
        if len(args) > 2 and args[2] is not None:
            count = max(0, int(_to_number(args[2]) or 0))
            return values[offset : offset + count]
        return values[offset:]
    if fn == "LIST.INSERTRANGE":
        values = list(args[0]) if args and isinstance(args[0], list) else []
        offset = max(0, min(len(values), int(_to_number(args[1]) or 0))) if len(args) > 1 else 0
        inserted = args[2] if len(args) > 2 and isinstance(args[2], list) else []
        return values[:offset] + list(inserted) + values[offset:]
    if fn == "LIST.REMOVERANGE":
        values = list(args[0]) if args and isinstance(args[0], list) else []
        offset = max(0, min(len(values), int(_to_number(args[1]) or 0))) if len(args) > 1 else 0
        count = len(values) - offset if len(args) <= 2 or args[2] is None else max(0, int(_to_number(args[2]) or 0))
        return values[:offset] + values[offset + count :]
    if fn == "LIST.REPLACERANGE":
        values = list(args[0]) if args and isinstance(args[0], list) else []
        offset = max(0, min(len(values), int(_to_number(args[1]) or 0))) if len(args) > 1 else 0
        count = max(0, int(_to_number(args[2]) or 0)) if len(args) > 2 else 0
        replacement = args[3] if len(args) > 3 and isinstance(args[3], list) else []
        return values[:offset] + list(replacement) + values[offset + count :]
    if fn == "LIST.REMOVEITEMS":
        values = list(args[0]) if args and isinstance(args[0], list) else []
        remove = args[1] if len(args) > 1 and isinstance(args[1], list) else []
        remove_keys = {_value_key(item) for item in remove}
        return [item for item in values if _value_key(item) not in remove_keys]
    if fn == "LIST.REMOVEMATCHINGITEMS":
        values = list(args[0]) if args and isinstance(args[0], list) else []
        remove = args[1] if len(args) > 1 and isinstance(args[1], list) else []
        remove_keys = {_value_key(item) for item in remove}
        return [item for item in values if _value_key(item) not in remove_keys]
    if fn == "LIST.REPLACEMATCHINGITEMS":
        values = list(args[0]) if args and isinstance(args[0], list) else []
        replacements = args[1] if len(args) > 1 and isinstance(args[1], list) else []
        lookup: dict[Any, Any] = {}
        for spec in replacements:
            if isinstance(spec, list) and len(spec) >= 2:
                lookup[_value_key(spec[0])] = spec[1]
        return [lookup.get(_value_key(item), item) for item in values]
    if fn == "LIST.REPLACEVALUE":
        values = list(args[0]) if args and isinstance(args[0], list) else []
        old_value = args[1] if len(args) > 1 else None
        new_value = args[2] if len(args) > 2 else None
        replacer = args[3] if len(args) > 3 else "Replacer.ReplaceValue"
        if str(replacer).upper().endswith("REPLACETEXT"):
            return [None if item is None else str(item).replace(str(old_value), str(new_value)) for item in values]
        return [new_value if item == old_value else item for item in values]
    if fn == "LIST.REPEAT":
        values = list(args[0]) if args and isinstance(args[0], list) else []
        count = max(0, int(_to_number(args[1]) or 0)) if len(args) > 1 else 0
        return values * count
    if fn == "LIST.SPLIT":
        values = list(args[0]) if args and isinstance(args[0], list) else []
        page_size = max(1, int(_to_number(args[1]) or 1)) if len(args) > 1 else 1
        return [values[idx : idx + page_size] for idx in range(0, len(values), page_size)]
    if fn == "LIST.ANYTRUE":
        values = args[0] if args and isinstance(args[0], list) else []
        return any(bool(item) for item in values)
    if fn == "LIST.ALLTRUE":
        values = args[0] if args and isinstance(args[0], list) else []
        return all(bool(item) for item in values)
    if fn == "TABLE.COLUMN":
        records = args[0] if args and isinstance(args[0], list) else []
        column_name = str(args[1]) if len(args) > 1 else ""
        values: list[Any] = []
        found = False
        for record in records:
            if not isinstance(record, Mapping):
                values.append(None)
                continue
            try:
                values.append(_mapping_get_case_insensitive(record, column_name))
                found = True
            except KeyError:
                values.append(None)
        if records and not found:
            raise KeyError(column_name)
        return values
    if fn == "TABLE.TORECORDS":
        return [dict(record) for record in args[0] if isinstance(record, Mapping)] if args and isinstance(args[0], list) else []
    if fn == "TABLE.TOROWS":
        records = [record for record in args[0] if isinstance(record, Mapping)] if args and isinstance(args[0], list) else []
        names: list[str] = []
        seen: set[str] = set()
        for record in records:
            for key in record.keys():
                key_text = str(key)
                upper = key_text.upper()
                if upper not in seen:
                    seen.add(upper)
                    names.append(key_text)
        return [[record.get(name) for name in names] for record in records]
    if fn == "TABLE.TOCOLUMNS":
        records = [record for record in args[0] if isinstance(record, Mapping)] if args and isinstance(args[0], list) else []
        names: list[str] = []
        seen: set[str] = set()
        for record in records:
            for key in record.keys():
                key_text = str(key)
                upper = key_text.upper()
                if upper not in seen:
                    seen.add(upper)
                    names.append(key_text)
        return [[record.get(name) for record in records] for name in names]
    if fn == "RECORD.FIELD":
        record = args[0] if isinstance(args[0], Mapping) else {}
        return record.get(str(args[1]))
    if fn == "RECORD.FIELDORDEFAULT":
        record = args[0] if isinstance(args[0], Mapping) else {}
        default = args[2] if len(args) > 2 else None
        return record.get(str(args[1]), default)
    if fn == "RECORD.HASFIELDS":
        record = args[0] if isinstance(args[0], Mapping) else {}
        fields = args[1] if isinstance(args[1], list) else [args[1]]
        return all(str(field) in record for field in fields)
    if fn == "RECORD.FIELDCOUNT":
        record = args[0] if isinstance(args[0], Mapping) else {}
        return len(record)
    if fn == "RECORD.FIELDNAMES":
        record = args[0] if isinstance(args[0], Mapping) else {}
        return [str(key) for key in record.keys()]
    if fn == "RECORD.FIELDVALUES":
        record = args[0] if isinstance(args[0], Mapping) else {}
        return list(record.values())
    if fn == "RECORD.TOLIST":
        record = args[0] if isinstance(args[0], Mapping) else {}
        return list(record.values())
    if fn == "RECORD.COMBINE":
        records = args[0] if args and isinstance(args[0], list) else []
        out: dict[str, Any] = {}
        for record in records:
            if not isinstance(record, Mapping):
                raise ValueError("Record.Combine preview needs a list of records.")
            out.update(record)
        return out
    if fn == "RECORD.ADDFIELD":
        record = dict(args[0]) if args and isinstance(args[0], Mapping) else {}
        field = str(args[1]) if len(args) > 1 else ""
        record[field] = args[2] if len(args) > 2 else None
        return record
    if fn == "RECORD.FROMLIST":
        values = args[0] if args and isinstance(args[0], list) else []
        fields = args[1] if len(args) > 1 and isinstance(args[1], list) else []
        return {str(field): values[idx] if idx < len(values) else None for idx, field in enumerate(fields)}
    if fn == "RECORD.FROMTABLE":
        records = args[0] if args and isinstance(args[0], list) else []
        out: dict[str, Any] = {}
        for record in records:
            if not isinstance(record, Mapping):
                continue
            try:
                name = _mapping_get_case_insensitive(record, "Name")
                value = _mapping_get_case_insensitive(record, "Value")
            except KeyError:
                continue
            out[str(name)] = value
        return out
    if fn == "RECORD.TOTABLE":
        record = args[0] if isinstance(args[0], Mapping) else {}
        return [{"Name": str(key), "Value": value} for key, value in record.items()]
    if fn == "RECORD.SELECTFIELDS":
        missing_field = _missing_field_mode_from_value(args[2] if len(args) > 2 else None)
        return _record_select_fields(
            args[0] if args else {},
            args[1] if len(args) > 1 else [],
            missing_field,
        )
    if fn == "RECORD.REMOVEFIELDS":
        missing_field = _missing_field_mode_from_value(args[2] if len(args) > 2 else None)
        return _record_remove_fields(
            args[0] if args else {},
            args[1] if len(args) > 1 else [],
            missing_field,
        )
    if fn == "RECORD.RENAMEFIELDS":
        missing_field = _missing_field_mode_from_value(args[2] if len(args) > 2 else None)
        return _record_rename_fields(
            args[0] if args else {},
            args[1] if len(args) > 1 else [],
            missing_field,
        )
    if fn == "RECORD.REORDERFIELDS":
        missing_field = _missing_field_mode_from_value(args[2] if len(args) > 2 else None)
        record = dict(args[0]) if args and isinstance(args[0], Mapping) else {}
        ordered: dict[str, Any] = {}
        used: set[str] = set()
        for field in _record_field_names(args[1] if len(args) > 1 else []):
            if field in record:
                ordered[field] = record[field]
                used.add(field)
            elif missing_field == "use_null":
                ordered[field] = None
            elif missing_field != "ignore":
                raise KeyError(field)
        for key, value in record.items():
            if key not in used:
                ordered[key] = value
        return ordered
    if fn == "RECORD.TRANSFORMFIELDNAMES":
        record = dict(args[0]) if args and isinstance(args[0], Mapping) else {}
        transform = args[1] if len(args) > 1 else None
        out: dict[str, Any] = {}
        for key, value in record.items():
            new_name = _eval_unary_transform(transform, str(key), variables=variables)
            out[str(new_name)] = value
        return out
    if fn == "RECORD.TRANSFORMFIELDS":
        missing_field = _missing_field_mode_from_value(args[2] if len(args) > 2 else None)
        return _record_transform_fields(
            args[0] if args else {},
            _record_transform_specs_from_value(args[1] if len(args) > 1 else []),
            missing_field=missing_field,
            variables=variables,
        )
    raise ValueError(f"Unsupported M expression function: {name}")


def _unique_csv_headers(headers: Sequence[Any]) -> list[str]:
    seen: dict[str, int] = {}
    out: list[str] = []
    for idx, header in enumerate(headers):
        base = str(header) if str(header) else f"Column{idx + 1}"
        count = seen.get(base, 0)
        seen[base] = count + 1
        out.append(base if count == 0 else f"{base}.{count}")
    return out


def _promote_headers_records(value: Any) -> list[dict[str, Any]]:
    records = value if isinstance(value, list) else []
    clean_records = [record for record in records if isinstance(record, Mapping)]
    if not clean_records:
        return []
    original_headers = list(clean_records[0].keys())
    promoted_headers = _unique_csv_headers([clean_records[0].get(header) for header in original_headers])
    out: list[dict[str, Any]] = []
    for record in clean_records[1:]:
        out.append({
            promoted: record.get(original)
            for original, promoted in zip(original_headers, promoted_headers)
        })
    return out


def _static_file_base_path(variables: Mapping[str, Any] | None) -> Path | None:
    if not variables or _STATIC_FILE_BASE_PATH_KEY not in variables:
        return None
    raw = variables.get(_STATIC_FILE_BASE_PATH_KEY)
    if raw is None or raw == "":
        return None
    return raw if isinstance(raw, Path) else Path(str(raw))


def _file_contents_binary_ref(value: Any, variables: Mapping[str, Any] | None = None) -> dict[str, Any]:
    raw_path = str(value or "").strip()
    if not raw_path:
        raise ValueError("File.Contents preview needs a static path string.")
    path = Path(raw_path)
    base_path = _static_file_base_path(variables)
    resolved = path if path.is_absolute() or base_path is None else base_path / path
    ref: dict[str, Any] = {
        "kind": "binary_ref",
        "path": str(resolved),
        "source_path": raw_path,
        "source": "File.Contents",
    }
    try:
        if resolved.exists() and resolved.is_file():
            ref["size"] = resolved.stat().st_size
    except OSError:
        pass
    return ref


def _binary_ref_path(value: Any, function_name: str) -> Path:
    if not isinstance(value, Mapping) or str(value.get("kind") or "").lower() != "binary_ref":
        raise ValueError(f"{function_name} preview currently expects a File.Contents or folder [Content] binary reference.")
    path_raw = value.get("path")
    if not path_raw:
        raise ValueError(f"{function_name} preview binary reference has no path.")
    path = Path(str(path_raw))
    if not path.exists() or not path.is_file():
        raise ValueError(f"{function_name} preview file not found: {path}")
    return path


def _json_document_value(value: Any) -> Any:
    if isinstance(value, str):
        return _json.loads(value)
    if isinstance(value, (bytes, bytearray)):
        return _json.loads(bytes(value).decode("utf-8-sig"))
    path = _binary_ref_path(value, "Json.Document")
    with path.open("r", encoding="utf-8-sig") as handle:
        return _json.load(handle)


def _json_ready_value(value: Any) -> Any:
    value = _unwrap_metadata(value)
    if isinstance(value, Mapping):
        return {str(key): _json_ready_value(child) for key, child in value.items()}
    if isinstance(value, list):
        return [_json_ready_value(child) for child in value]
    if isinstance(value, tuple):
        return [_json_ready_value(child) for child in value]
    if isinstance(value, (_dt.datetime, _dt.date, _dt.time)):
        return value.isoformat()
    if isinstance(value, _dt.timedelta):
        return value.total_seconds()
    if isinstance(value, (bytes, bytearray)):
        return _base64.b64encode(bytes(value)).decode("ascii")
    if isinstance(value, MFunctionValue):
        return {"Kind": "Function", "Parameters": value.parameters, "Body": value.body}
    if isinstance(value, MBuiltinFunctionValue):
        return {"Kind": "BuiltinFunction", "Name": value.name}
    if isinstance(value, MErrorValue):
        return {"Kind": "Error", "Message": value.message}
    return value


def _json_from_value(value: Any, encoding: Any = None) -> bytes | None:
    if value is None:
        return None
    text = _json.dumps(_json_ready_value(value), ensure_ascii=False, separators=(",", ":"), default=str)
    return _text_to_binary(text, encoding or "TextEncoding.Utf8")


def _html_plain_text(value: str) -> str:
    text = re.sub(r"(?is)<(script|style)\b[^>]*>.*?</\1>", "", value or "")
    text = re.sub(r"(?s)<[^>]+>", " ", text)
    return " ".join(_html.unescape(text).split())


def _html_selector_tag(selector: Any) -> str | None:
    text = str(selector or "").strip()
    if not text:
        return None
    candidate = re.split(r"\s+", text)[-1]
    candidate = re.split(r"[.#\[:]", candidate, maxsplit=1)[0]
    return candidate.lower() if re.match(r"^[A-Za-z][A-Za-z0-9]*$", candidate) else None


def _html_selector_specs(specs: Any) -> list[tuple[str, str | None]]:
    out: list[tuple[str, str | None]] = []
    values = specs if isinstance(specs, list) else []
    for idx, spec in enumerate(values):
        if isinstance(spec, Mapping):
            name = str(spec.get("Name") or spec.get("name") or spec.get("Column") or spec.get("column") or f"Column{idx + 1}")
            selector = spec.get("Selector") or spec.get("selector")
            out.append((name, _html_selector_tag(selector)))
        elif isinstance(spec, list) and spec:
            name = str(spec[0] if spec[0] is not None else f"Column{idx + 1}")
            selector = spec[1] if len(spec) > 1 else None
            out.append((name, _html_selector_tag(selector)))
    return out


def _html_table_value(value: Any, specs: Any) -> list[dict[str, Any]]:
    html_text = bytes(value).decode("utf-8", errors="replace") if isinstance(value, (bytes, bytearray)) else str(value or "")
    selectors = _html_selector_specs(specs)
    if not selectors:
        return []
    values_by_column: list[tuple[str, list[str]]] = []
    for name, tag in selectors:
        if not tag:
            values_by_column.append((name, []))
            continue
        pattern = re.compile(rf"(?is)<{re.escape(tag)}\b[^>]*>(.*?)</{re.escape(tag)}>")
        values_by_column.append((name, [_html_plain_text(match.group(1)) for match in pattern.finditer(html_text)]))
    row_count = max((len(values) for _name, values in values_by_column), default=0)
    rows: list[dict[str, Any]] = []
    for row_idx in range(row_count):
        row: dict[str, Any] = {}
        for name, values in values_by_column:
            row[name] = values[row_idx] if row_idx < len(values) else None
        rows.append(row)
    return rows


def _identity_value(value: Any, provider: Any = None) -> dict[str, Any]:
    if isinstance(value, Mapping):
        name = value.get("Name") or value.get("name") or value.get("Id") or value.get("id")
    else:
        name = value
    out: dict[str, Any] = {"Kind": "Identity", "Name": str(name or "")}
    if provider is not None:
        out["Provider"] = provider
    return out


def _identity_name(value: Any) -> str:
    if isinstance(value, Mapping):
        for key in ("Name", "name", "Id", "id"):
            if key in value:
                return str(value.get(key) or "").upper()
    return str(value or "").upper()


def _identity_is_member_of(identity: Any, collection: Any) -> bool:
    target = _identity_name(identity)
    if not target:
        return False
    if isinstance(collection, Mapping):
        collection = collection.get("Members") or collection.get("members") or collection.get("Identities") or collection.get("identities")
    values = collection if isinstance(collection, list) else [collection]
    return any(_identity_name(value) == target for value in values)


def _xml_local_name(name: str) -> str:
    if "}" in name:
        return name.rsplit("}", 1)[1]
    return name


def _xml_direct_text(element: _et.Element, *, include_child_tails: bool = True) -> str | None:
    segments: list[str] = []
    text = (element.text or "").strip()
    if text:
        segments.append(text)
    if include_child_tails:
        for child in list(element):
            tail = (child.tail or "").strip()
            if tail:
                segments.append(tail)
    return " ".join(segments) if segments else None


def _xml_scalar_text(element: _et.Element) -> str | None:
    return _xml_direct_text(element, include_child_tails=False)


def _xml_record_put(record: dict[str, Any], name: str, value: Any) -> None:
    base = name or "Value"
    candidate = base
    idx = 1
    existing = {str(key).upper() for key in record}
    while candidate.upper() in existing:
        candidate = f"{base}.{idx}"
        idx += 1
    record[candidate] = value


def _xml_attribute_record(element: _et.Element) -> dict[str, Any]:
    record: dict[str, Any] = {}
    for key, value in element.attrib.items():
        _xml_record_put(record, _xml_local_name(key), value)
    return record


def _xml_element_record(element: _et.Element) -> dict[str, Any]:
    record = _xml_attribute_record(element)
    children = list(element)
    if not children:
        text = _xml_scalar_text(element)
        if text is not None:
            _xml_record_put(record, _xml_local_name(element.tag), text)
        return record
    text = _xml_direct_text(element)
    if text is not None:
        _xml_record_put(record, "#text", text)
    grouped: dict[str, list[Any]] = {}
    for child in children:
        name = _xml_local_name(child.tag)
        child_children = list(child)
        if child_children or child.attrib:
            value: Any = _xml_element_record(child)
        else:
            value = _xml_scalar_text(child)
        grouped.setdefault(name, []).append(value)
    for name, values in grouped.items():
        _xml_record_put(record, name, values[0] if len(values) == 1 else values)
    return record


def _xml_tables_value(value: Any) -> list[dict[str, Any]]:
    if isinstance(value, str):
        root = _et.fromstring(value)
    elif isinstance(value, (bytes, bytearray)):
        root = _et.fromstring(bytes(value))
    else:
        path = _binary_ref_path(value, "Xml.Tables")
        root = _et.parse(path).getroot()
    children = list(root)
    if children:
        root_attributes = _xml_attribute_record(root)
        if not root_attributes:
            return [_xml_element_record(child) for child in children]
        rows: list[dict[str, Any]] = []
        for child in children:
            row = dict(root_attributes)
            for key, child_value in _xml_element_record(child).items():
                _xml_record_put(row, key, child_value)
            rows.append(row)
        return rows
    return [_xml_element_record(root)]


def _binary_from_text(value: Any, encoding: Any = None) -> bytes:
    text = "" if value is None else str(value)
    mode = str(encoding or "BinaryEncoding.Base64").strip().upper()
    if mode.endswith(".BASE64") or mode == "BASE64":
        return _base64.b64decode(text)
    if mode.endswith(".HEX") or mode == "HEX":
        return bytes.fromhex(text)
    raise ValueError(f"Binary.FromText preview does not support {encoding!r}.")


def _binary_to_text(value: Any, encoding: Any = None) -> str | None:
    if value is None:
        return None
    raw = bytes(_unwrap_metadata(value) or b"")
    mode = str(encoding or "BinaryEncoding.Base64").strip().upper()
    if mode.endswith(".BASE64") or mode == "BASE64":
        return _base64.b64encode(raw).decode("ascii")
    if mode.endswith(".HEX") or mode == "HEX":
        return raw.hex()
    raise ValueError(f"Binary.ToText preview does not support {encoding!r}.")


def _text_encoding_name(encoding: Any, *, decode: bool = False) -> str:
    mode = str(encoding or "TextEncoding.Utf8").strip().upper()
    if mode.endswith(".UTF16BE") or mode == "UTF16BE":
        return "utf-16-be"
    if mode.endswith(".UTF16") or mode == "UTF16":
        return "utf-16" if decode else "utf-16-le"
    if mode.endswith(".ASCII") or mode == "ASCII":
        return "ascii"
    if mode.endswith(".UTF8") or mode == "UTF8" or encoding is None:
        return "utf-8"
    raise ValueError(f"Text binary preview does not support {encoding!r}.")


def _text_to_binary(value: Any, encoding: Any = None, include_bom: Any = None) -> bytes | None:
    if value is None:
        return None
    codec = _text_encoding_name(encoding)
    raw = str(value).encode(codec)
    if bool(include_bom):
        mode = str(encoding or "TextEncoding.Utf8").strip().upper()
        if mode.endswith(".UTF16BE") or mode == "UTF16BE":
            return b"\xfe\xff" + raw
        if mode.endswith(".UTF16") or mode == "UTF16":
            return b"\xff\xfe" + raw
        if mode.endswith(".UTF8") or mode == "UTF8" or encoding is None:
            return b"\xef\xbb\xbf" + raw
    return raw


def _text_from_binary(value: Any, encoding: Any = None) -> str | None:
    if value is None:
        return None
    raw = bytes(_unwrap_metadata(value) or b"")
    codec = _text_encoding_name(encoding, decode=True)
    return raw.decode(codec)


def _lines_from_text(value: Any, quote_style: Any = None, include_line_separators: Any = None) -> list[str]:
    if value is None:
        return []
    text = str(value)
    include = bool(include_line_separators)
    if include:
        return text.splitlines(keepends=True)
    return text.splitlines()


def _lines_to_text(lines: Any, line_separator: Any = None) -> str:
    values = ["" if item is None else str(item) for item in lines] if isinstance(lines, list) else []
    separator = "\r\n" if line_separator is None else str(line_separator)
    return separator.join(values)


def _uri_parts(value: Any) -> dict[str, Any]:
    parsed = _urlparse.urlsplit(str(value or ""))
    query = {key: vals[-1] if vals else "" for key, vals in _urlparse.parse_qs(parsed.query, keep_blank_values=True).items()}
    return {
        "Scheme": parsed.scheme,
        "Host": parsed.hostname,
        "Port": parsed.port,
        "Path": parsed.path,
        "Query": query,
        "Fragment": parsed.fragment,
        "UserName": parsed.username,
        "Password": parsed.password,
    }


def _uri_build_query_string(value: Any) -> str:
    if not isinstance(value, Mapping):
        return ""
    return _urlparse.urlencode({
        str(key): "" if item is None else str(item)
        for key, item in value.items()
    })


def _uri_combine(base: Any, relative: Any) -> str:
    return _urlparse.urljoin(str(base or ""), str(relative or ""))


def _text_format(template: Any, arguments: Any, culture: Any = None) -> str | None:
    if template is None:
        return None
    text = str(template)
    if arguments is None:
        arguments = []

    def replacement(match: re.Match[str]) -> str:
        key = match.group(1).strip()
        value: Any = None
        found = False
        if isinstance(arguments, Mapping):
            for existing_key, existing_value in arguments.items():
                if str(existing_key).upper() == key.upper():
                    value = existing_value
                    found = True
                    break
        elif isinstance(arguments, list):
            try:
                idx = int(key)
            except Exception:
                idx = -1
            if 0 <= idx < len(arguments):
                value = arguments[idx]
                found = True
        if not found:
            return match.group(0)
        converted = _eval_function_call("Text.From", [value, culture] if culture is not None else [value])
        return "" if converted is None else str(converted)

    return re.sub(r"#\{([^}]+)\}", replacement, text)


def _m_literal_text(value: Any) -> str:
    value = _unwrap_metadata(value)
    if value is None:
        return "null"
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (int, float)):
        if isinstance(value, float) and _math.isnan(value):
            return "Number.NaN"
        if isinstance(value, float) and _math.isinf(value):
            return "Number.PositiveInfinity" if value > 0 else "Number.NegativeInfinity"
        return str(int(value)) if isinstance(value, float) and value.is_integer() else str(value)
    if isinstance(value, str):
        return '"' + value.replace('"', '""') + '"'
    if isinstance(value, _dt.datetime):
        if value.tzinfo is not None:
            offset = value.utcoffset() or _dt.timedelta()
            total_minutes = int(offset.total_seconds() // 60)
            hours = int(total_minutes / 60)
            minutes = abs(total_minutes) % 60
            return f"#datetimezone({value.year}, {value.month}, {value.day}, {value.hour}, {value.minute}, {value.second}, {hours}, {minutes})"
        return f"#datetime({value.year}, {value.month}, {value.day}, {value.hour}, {value.minute}, {value.second})"
    if isinstance(value, _dt.date):
        return f"#date({value.year}, {value.month}, {value.day})"
    if isinstance(value, _dt.time):
        return f"#time({value.hour}, {value.minute}, {value.second})"
    if isinstance(value, list):
        return "{" + ", ".join(_m_literal_text(item) for item in value) + "}"
    if isinstance(value, Mapping):
        return "[" + ", ".join(f"{_m_identifier_text(str(key))} = {_m_literal_text(item)}" for key, item in value.items()) + "]"
    return '"' + str(value).replace('"', '""') + '"'


def _m_identifier_text(value: str) -> str:
    if re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", value):
        return value
    return '#"' + value.replace('"', '""') + '"'


def _guid_from_value(value: Any) -> str | None:
    if value is None:
        return None
    value = _unwrap_metadata(value)
    if isinstance(value, (bytes, bytearray)):
        return str(_uuid.UUID(bytes=bytes(value)))
    text = str(value).strip()
    if not text:
        return None
    return str(_uuid.UUID(text))


def _as_text_list(value: Any) -> list[str]:
    if isinstance(value, list):
        return ["" if item is None else str(item) for item in value]
    if value is None:
        return []
    return [str(value)]


def _split_from_positions(text: str, positions: Sequence[Any], *, from_end: bool = False) -> list[str]:
    raw_positions = [max(0, int(_to_number(item) or 0)) for item in positions]
    if from_end:
        raw_positions = [max(0, len(text) - position) for position in raw_positions]
    points = sorted({0, len(text), *[min(len(text), position) for position in raw_positions]})
    return [text[points[idx] : points[idx + 1]] for idx in range(len(points) - 1)]


def _split_from_lengths(text: str, lengths: Sequence[Any], *, from_end: bool = False) -> list[str]:
    sizes = [max(0, int(_to_number(item) or 0)) for item in lengths]
    if not sizes:
        return [text]
    if from_end:
        out: list[str] = []
        end = len(text)
        for size in sizes:
            start = max(0, end - size)
            out.append(text[start:end])
            end = start
        if end > 0:
            out.append(text[:end])
        return list(reversed([item for item in out if item != ""]))
    out = []
    offset = 0
    for size in sizes:
        out.append(text[offset : offset + size])
        offset += size
    if offset < len(text):
        out.append(text[offset:])
    return out


def _range_pair(value: Any) -> tuple[int, int] | None:
    if isinstance(value, Mapping):
        start = value.get("Start") if "Start" in value else value.get("start")
        length = value.get("Length") if "Length" in value else value.get("length")
        if start is not None and length is not None:
            return max(0, int(_to_number(start) or 0)), max(0, int(_to_number(length) or 0))
    if isinstance(value, list) and len(value) >= 2:
        return max(0, int(_to_number(value[0]) or 0)), max(0, int(_to_number(value[1]) or 0))
    return None


def _split_from_ranges(text: str, ranges: Sequence[Any], *, from_end: bool = False) -> list[str]:
    out: list[str] = []
    for item in ranges:
        pair = _range_pair(item)
        if pair is None:
            continue
        start, length = pair
        if from_end:
            start = max(0, len(text) - start - length)
        out.append(text[start : start + length])
    return out


def _split_each_delimiter(text: str, delimiters: Sequence[str], *, from_end: bool = False) -> list[str]:
    if not delimiters:
        return [text]
    if from_end:
        tail = text
        out: list[str] = []
        for delimiter in reversed(list(delimiters)):
            if delimiter and delimiter in tail:
                left, right = tail.rsplit(delimiter, 1)
                out.append(right)
                tail = left
            else:
                out.append(tail)
                tail = ""
        out.append(tail)
        return list(reversed(out))
    tail = text
    out = []
    for delimiter in delimiters:
        if delimiter and delimiter in tail:
            head, tail = tail.split(delimiter, 1)
            out.append(head)
        else:
            out.append(tail)
            tail = ""
    out.append(tail)
    return out


def _predicate_accepts_char(predicate: Any, char: str, *, variables: Mapping[str, Any] | None = None) -> bool:
    if isinstance(predicate, (MFunctionValue, MBuiltinFunctionValue)):
        return bool(_invoke_m_function(predicate, [char], variables=variables))
    if isinstance(predicate, list):
        return char in {str(item) for item in predicate}
    if isinstance(predicate, str):
        return char in set(predicate)
    return False


def _apply_splitter(function: MBuiltinFunctionValue, value: Any, *, variables: Mapping[str, Any] | None = None) -> list[str]:
    text = "" if value is None else str(value)
    name = function.name.upper()
    options = function.options
    from_end = bool(options.get("start_at_end"))
    if name == "SPLITTER.SPLITBYNOTHING":
        return [text]
    if name == "SPLITTER.SPLITTEXTBYDELIMITER":
        delimiter = str(options.get("delimiter") or "")
        return text.split(delimiter) if delimiter else [text]
    if name == "SPLITTER.SPLITTEXTBYANYDELIMITER":
        delimiters = [delimiter for delimiter in _as_text_list(options.get("delimiters")) if delimiter]
        if not delimiters:
            return [text]
        return re.split("|".join(re.escape(delimiter) for delimiter in delimiters), text)
    if name == "SPLITTER.SPLITTEXTBYEACHDELIMITER":
        return _split_each_delimiter(text, _as_text_list(options.get("delimiters")), from_end=from_end)
    if name == "SPLITTER.SPLITTEXTBYLENGTHS":
        return _split_from_lengths(text, options.get("lengths") if isinstance(options.get("lengths"), list) else [], from_end=from_end)
    if name == "SPLITTER.SPLITTEXTBYPOSITIONS":
        return _split_from_positions(text, options.get("positions") if isinstance(options.get("positions"), list) else [], from_end=from_end)
    if name == "SPLITTER.SPLITTEXTBYRANGES":
        return _split_from_ranges(text, options.get("ranges") if isinstance(options.get("ranges"), list) else [], from_end=from_end)
    if name == "SPLITTER.SPLITTEXTBYREPEATEDLENGTHS":
        size = max(1, int(_to_number(options.get("length")) or 1))
        return [text[idx : idx + size] for idx in range(0, len(text), size)]
    if name == "SPLITTER.SPLITTEXTBYWHITESPACE":
        stripped = text.strip()
        return [] if not stripped else re.split(r"\s+", stripped)
    if name == "SPLITTER.SPLITTEXTBYCHARACTERTRANSITION":
        before = options.get("before")
        after = options.get("after")
        if not text:
            return [text]
        out: list[str] = []
        start = 0
        for idx in range(1, len(text)):
            if _predicate_accepts_char(before, text[idx - 1], variables=variables) and _predicate_accepts_char(after, text[idx], variables=variables):
                out.append(text[start:idx])
                start = idx
        out.append(text[start:])
        return out
    return [text]


def _combine_by_positions(values: Sequence[str], positions: Sequence[Any]) -> str:
    starts = [max(0, int(_to_number(item) or 0)) for item in positions]
    if not starts:
        return "".join(values)
    width = max([start + len(values[idx]) for idx, start in enumerate(starts[: len(values)])] or [0])
    chars = [" "] * width
    for idx, value in enumerate(values):
        start = starts[idx] if idx < len(starts) else len(chars)
        end = start + len(value)
        if end > len(chars):
            chars.extend(" " for _ in range(end - len(chars)))
        chars[start:end] = list(value)
    return "".join(chars).rstrip()


def _apply_combiner(function: MBuiltinFunctionValue, values: Any) -> str:
    parts = ["" if item is None else str(item) for item in values] if isinstance(values, list) else ["" if values is None else str(values)]
    name = function.name.upper()
    options = function.options
    if name == "COMBINER.COMBINETEXTBYDELIMITER":
        return str(options.get("delimiter") or "").join(parts)
    if name == "COMBINER.COMBINETEXTBYEACHDELIMITER":
        delimiters = _as_text_list(options.get("delimiters"))
        if not delimiters:
            return "".join(parts)
        out = parts[0] if parts else ""
        for idx, part in enumerate(parts[1:]):
            delimiter = delimiters[idx] if idx < len(delimiters) else delimiters[-1]
            out += delimiter + part
        return out
    if name == "COMBINER.COMBINETEXTBYLENGTHS":
        lengths = options.get("lengths") if isinstance(options.get("lengths"), list) else []
        out = ""
        for idx, part in enumerate(parts):
            size = max(0, int(_to_number(lengths[idx]) or 0)) if idx < len(lengths) else len(part)
            out += part[:size].ljust(size)
        return out
    if name == "COMBINER.COMBINETEXTBYPOSITIONS":
        return _combine_by_positions(parts, options.get("positions") if isinstance(options.get("positions"), list) else [])
    if name == "COMBINER.COMBINETEXTBYRANGES":
        ranges = options.get("ranges") if isinstance(options.get("ranges"), list) else []
        starts = [(_range_pair(item) or (0, len(parts[idx]) if idx < len(parts) else 0))[0] for idx, item in enumerate(ranges)]
        return _combine_by_positions(parts, starts)
    return "".join(parts)


def _list_percentile(values: Sequence[Any], percentile: Any) -> Any:
    numeric = sorted(_numeric_values(values))
    if not numeric:
        return None
    if isinstance(percentile, list):
        return [_list_percentile(numeric, item) for item in percentile]
    p = _to_number(percentile)
    if p is None:
        return None
    p = min(1.0, max(0.0, float(p)))
    if len(numeric) == 1:
        return numeric[0]
    position = p * (len(numeric) - 1)
    lower = int(_math.floor(position))
    upper = int(_math.ceil(position))
    if lower == upper:
        return numeric[lower]
    fraction = position - lower
    return numeric[lower] + ((numeric[upper] - numeric[lower]) * fraction)


def _wkt_record(value: Any, *, kind: str) -> dict[str, Any]:
    text = str(value or "").strip()
    match = re.match(r"^([A-Za-z]+)\s*\((.*)\)$", text)
    geometry_type = match.group(1).upper() if match else kind.upper()
    coordinates = match.group(2).strip() if match else text
    return {"Kind": kind, "Type": geometry_type, "Coordinates": coordinates, "WellKnownText": text}


def _wkt_from_record(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, Mapping):
        existing = value.get("WellKnownText") or value.get("wellKnownText")
        if existing is not None:
            return str(existing)
        geometry_type = str(value.get("Type") or value.get("type") or "POINT")
        coordinates = str(value.get("Coordinates") or value.get("coordinates") or "")
        return f"{geometry_type.upper()}({coordinates})"
    return str(value)


def _binary_from_list(value: Any) -> bytes:
    values = value if isinstance(value, list) else []
    return bytes(max(0, min(255, int(_to_number(item) or 0))) for item in values)


def _binary_from_value(value: Any) -> bytes | None:
    if value is None:
        return None
    value = _unwrap_metadata(value)
    if isinstance(value, (bytes, bytearray)):
        return bytes(value)
    if isinstance(value, list):
        return _binary_from_list(value)
    if isinstance(value, str):
        return value.encode("utf-8")
    raise ValueError(f"Binary.From preview does not support {type(value).__name__}.")


def _binary_compress(value: Any, compression: Any = None) -> bytes:
    raw = bytes(_unwrap_metadata(value) or b"")
    mode = str(compression or "Compression.Deflate").strip().upper()
    if mode.endswith(".DEFLATE") or mode == "DEFLATE":
        return _zlib.compress(raw)
    if mode.endswith(".GZIP") or mode == "GZIP":
        return _gzip.compress(raw)
    raise ValueError(f"Binary.Compress preview does not support {compression!r}.")


def _binary_range(value: Any, offset: Any = 0, count: Any = None) -> bytes:
    raw = bytes(_unwrap_metadata(value) or b"")
    start = max(0, int(_to_number(offset) or 0))
    if count is None:
        return raw[start:]
    return raw[start : start + max(0, int(_to_number(count) or 0))]


def _binary_split(value: Any, page_size: Any) -> list[bytes]:
    raw = bytes(_unwrap_metadata(value) or b"")
    size = max(1, int(_to_number(page_size) or 1))
    return [raw[idx : idx + size] for idx in range(0, len(raw), size)]


def _binary_decompress(value: Any, compression: Any = None) -> bytes:
    raw = bytes(value or b"")
    mode = str(compression or "Compression.Deflate").strip().upper()
    if mode.endswith(".DEFLATE") or mode == "DEFLATE":
        try:
            return _zlib.decompress(raw, -_zlib.MAX_WBITS)
        except Exception:
            return _zlib.decompress(raw)
    if mode.endswith(".GZIP") or mode == "GZIP":
        return _gzip.decompress(raw)
    raise ValueError(f"Binary.Decompress preview does not support {compression!r}.")


def _binary_source_path(value: Any, function_name: str) -> Path:
    if isinstance(value, Mapping) and str(value.get("kind") or "").lower() == "binary_ref":
        return _binary_ref_path(value, function_name)
    raise ValueError(f"{function_name} preview currently expects a File.Contents or folder [Content] binary reference.")


def _excel_workbook_value(value: Any, use_headers: Any = None) -> list[dict[str, Any]]:
    path = _binary_source_path(value, "Excel.Workbook")
    try:
        from openpyxl import load_workbook
    except Exception as exc:  # pragma: no cover - dependency is present in normal test/runtime envs.
        raise ValueError("Excel.Workbook preview requires openpyxl.") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    promote_headers = bool(use_headers) if use_headers is not None else False
    out: list[dict[str, Any]] = []
    for sheet in workbook.worksheets:
        raw_rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        if promote_headers and raw_rows:
            headers = _unique_csv_headers(raw_rows[0])
            data_rows = raw_rows[1:]
        else:
            width = max((len(row) for row in raw_rows), default=0)
            headers = [f"Column{idx + 1}" for idx in range(width)]
            data_rows = raw_rows
        records = [
            {header: row[idx] if idx < len(row) else None for idx, header in enumerate(headers)}
            for row in data_rows
        ]
        out.append(
            {
                "Name": sheet.title,
                "Data": records,
                "Item": sheet.title,
                "Kind": "Sheet",
                "Hidden": sheet.sheet_state != "visible",
            }
        )
    return out


def _csv_encoding(options: Any) -> str:
    if not isinstance(options, Mapping):
        return "utf-8-sig"
    raw = options.get("Encoding") or options.get("encoding")
    if raw is None:
        return "utf-8-sig"
    code = str(raw).strip().lower()
    if code in {"65001", "utf-8", "utf8"}:
        return "utf-8-sig"
    if code in {"1252", "windows-1252", "cp1252"}:
        return "cp1252"
    if code in {"1200", "utf-16", "utf-16le", "utf-16-le"}:
        return "utf-16-le"
    if code in {"1201", "utf-16be", "utf-16-be"}:
        return "utf-16-be"
    raise ValueError(f"Csv.Document preview does not support Encoding={raw!r} yet.")


def _csv_delimiter(options: Any) -> str:
    delimiter = ","
    if isinstance(options, Mapping):
        raw_delimiter = options.get("Delimiter") or options.get("delimiter")
        if raw_delimiter is not None:
            delimiter = _decode_m_text(str(raw_delimiter))
    if delimiter in {"\\t", "tab"}:
        delimiter = "\t"
    if delimiter == "":
        delimiter = ","
    if len(delimiter) != 1:
        raise ValueError(f"Csv.Document preview requires a one-character delimiter, got {delimiter!r}.")
    return delimiter


def _csv_quote_options(options: Any) -> dict[str, Any]:
    if not isinstance(options, Mapping):
        return {"quotechar": '"', "quoting": _csv.QUOTE_MINIMAL}
    quote_style = str(options.get("QuoteStyle") or options.get("quoteStyle") or "").strip().lower()
    if quote_style.endswith(".none") or quote_style == "none":
        return {"quotechar": None, "quoting": _csv.QUOTE_NONE, "escapechar": "\\"}
    return {"quotechar": '"', "quoting": _csv.QUOTE_MINIMAL}


def _csv_configured_columns(options: Any, rows: Sequence[Sequence[Any]]) -> list[str]:
    if isinstance(options, Mapping):
        raw_columns = options.get("Columns") or options.get("columns")
        if isinstance(raw_columns, list):
            return _unique_csv_headers([str(item) for item in raw_columns])
        if isinstance(raw_columns, int) and raw_columns > 0:
            return [f"Column{idx + 1}" for idx in range(raw_columns)]
        if isinstance(raw_columns, float) and raw_columns > 0:
            return [f"Column{idx + 1}" for idx in range(int(raw_columns))]
    column_count = max((len(row) for row in rows), default=0)
    return [f"Column{idx + 1}" for idx in range(column_count)]


def _csv_document_records(value: Any, options: Any = None, *, promote_headers: bool = False) -> list[dict[str, Any]]:
    """Preview a simple CSV binary reference as records.

    This intentionally covers static File.Contents refs and the common Power BI
    combine-files helper shape: Folder.Files -> [Content] -> Csv.Document ->
    ExpandTableColumn. It is not a full Csv.Document clone yet; unsupported
    binary and option shapes fail loudly so the editor does not pretend they can
    fold or execute.
    """

    path = _binary_ref_path(value, "Csv.Document")

    delimiter = _csv_delimiter(options)
    encoding = _csv_encoding(options)
    quote_options = _csv_quote_options(options)
    with path.open("r", encoding=encoding, newline="") as handle:
        rows = list(_csv.reader(handle, delimiter=delimiter, **quote_options))
    if not rows:
        return []
    headers = _csv_configured_columns(options, rows)
    out: list[dict[str, Any]] = []
    for raw_row in rows:
        record: dict[str, Any] = {}
        for idx, name in enumerate(headers):
            record[name] = raw_row[idx] if idx < len(raw_row) else None
        out.append(record)
    return _promote_headers_records(out) if promote_headers else out


def _m_identifier_name(text: str) -> str:
    stripped = text.strip()
    if stripped.startswith('#"') and stripped.endswith('"'):
        return stripped[2:-1].replace('""', '"')
    return stripped


_M_ACCESS_SENTINEL = object()


def _mapping_get_case_insensitive(value: Mapping[str, Any], key: str) -> Any:
    if key in value:
        return value[key]
    wanted = key.upper()
    for existing_key, existing_value in value.items():
        if str(existing_key).upper() == wanted:
            return existing_value
    raise KeyError(key)


def _mapping_has_value(value: Mapping[str, Any], key: str, expected: Any) -> bool:
    try:
        actual = _mapping_get_case_insensitive(value, key)
    except KeyError:
        return False
    return actual == expected or str(actual) == str(expected)


def _missing_field_mode_from_value(value: Any, *, default: str = "error") -> str:
    text = str(value or "").strip().upper()
    if not text:
        return default
    if text in {"MISSINGFIELD.IGNORE", "IGNORE", "1"}:
        return "ignore"
    if text in {"MISSINGFIELD.USENULL", "USENULL", "USE_NULL", "2"}:
        return "use_null"
    return "error"


def _record_field_names(value: Any) -> list[str]:
    if isinstance(value, list):
        return [str(item) for item in value]
    if value is None:
        return []
    return [str(value)]


def _record_rename_specs(value: Any) -> list[tuple[str, str]]:
    specs = value if isinstance(value, list) else []
    out: list[tuple[str, str]] = []
    for spec in specs:
        if isinstance(spec, list) and len(spec) >= 2:
            out.append((str(spec[0]), str(spec[1])))
    return out


def _record_transform_specs_from_value(value: Any) -> list[tuple[str, Any]]:
    specs = value if isinstance(value, list) else []
    out: list[tuple[str, Any]] = []
    for spec in specs:
        if isinstance(spec, list) and len(spec) >= 2:
            out.append((str(spec[0]), spec[1]))
    return out


def _record_transform_specs_from_text(
    text: str,
    *,
    variables: Mapping[str, Any] | None = None,
) -> list[tuple[str, Any]] | None:
    spec_list = _between_outer(text, "{", "}")
    if spec_list is None:
        value = _eval_m_expression([], [], text, variables=variables)
        return _record_transform_specs_from_value(value)
    out: list[tuple[str, Any]] = []
    for raw_spec in _split_top_level_items(spec_list):
        body = _between_outer(raw_spec, "{", "}")
        if body is None:
            continue
        parts = _split_top_level_items(body)
        if len(parts) < 2:
            continue
        field_name = _eval_m_expression([], [], parts[0], variables=variables)
        out.append((str(field_name), parts[1].strip()))
    return out


def _eval_unary_transform(
    transform: Any,
    value: Any,
    *,
    variables: Mapping[str, Any] | None = None,
) -> Any:
    if isinstance(transform, (MFunctionValue, MBuiltinFunctionValue)):
        return _invoke_m_function(transform, [value], variables=variables)
    if isinstance(transform, str):
        return _eval_list_lambda(transform, value, variables=variables)
    raise ValueError("Record.TransformFields preview needs unary transform functions.")


def _record_select_fields(
    record_value: Any,
    fields_value: Any,
    missing_field: str = "error",
) -> dict[str, Any]:
    record = dict(record_value) if isinstance(record_value, Mapping) else {}
    out: dict[str, Any] = {}
    for field in _record_field_names(fields_value):
        if field in record:
            out[field] = record[field]
        elif missing_field == "use_null":
            out[field] = None
        elif missing_field != "ignore":
            raise KeyError(field)
    return out


def _record_remove_fields(
    record_value: Any,
    fields_value: Any,
    missing_field: str = "error",
) -> dict[str, Any]:
    record = dict(record_value) if isinstance(record_value, Mapping) else {}
    fields = _record_field_names(fields_value)
    missing = [field for field in fields if field not in record]
    if missing and missing_field == "error":
        raise KeyError(missing[0])
    remove = set(fields)
    return {key: value for key, value in record.items() if key not in remove}


def _record_rename_fields(
    record_value: Any,
    renames_value: Any,
    missing_field: str = "error",
) -> dict[str, Any]:
    record = dict(record_value) if isinstance(record_value, Mapping) else {}
    specs = _record_rename_specs(renames_value)
    rename = {old: new for old, new in specs}
    missing = [(old, new) for old, new in specs if old not in record]
    if missing and missing_field == "error":
        raise KeyError(missing[0][0])
    out = {rename.get(key, key): value for key, value in record.items()}
    if missing_field == "use_null":
        for _old, new in missing:
            out[new] = None
    return out


def _record_transform_fields(
    record_value: Any,
    specs: Sequence[tuple[str, Any]],
    *,
    missing_field: str = "error",
    variables: Mapping[str, Any] | None = None,
) -> dict[str, Any]:
    out = dict(record_value) if isinstance(record_value, Mapping) else {}
    for field, transform in specs:
        if field in out:
            out[field] = _eval_unary_transform(transform, out[field], variables=variables)
        elif missing_field == "use_null":
            out[field] = _eval_unary_transform(transform, None, variables=variables)
        elif missing_field != "ignore":
            raise KeyError(field)
    return out


def _access_tokenize_suffix(suffix: str) -> list[tuple[str, str]]:
    tokens: list[tuple[str, str]] = []
    i = 0
    while i < len(suffix):
        ch = suffix[i]
        if ch not in "{[":
            if ch.isspace():
                i += 1
                continue
            return []
        closer = "}" if ch == "{" else "]"
        depth = 1
        in_string = False
        j = i + 1
        while j < len(suffix):
            cur = suffix[j]
            nxt = suffix[j + 1] if j + 1 < len(suffix) else ""
            if in_string:
                if cur == '"' and nxt == '"':
                    j += 2
                    continue
                if cur == '"':
                    in_string = False
                j += 1
                continue
            if cur == '"':
                in_string = True
            elif cur == ch:
                depth += 1
            elif cur == closer:
                depth -= 1
                if depth == 0:
                    tokens.append(("item" if ch == "{" else "field", suffix[i + 1 : j].strip()))
                    i = j + 1
                    break
            j += 1
        else:
            return []
    return tokens


def _eval_item_access(value: Any, selector_text: str) -> Any:
    selector = selector_text.strip()
    if re.fullmatch(r"-?\d+", selector):
        idx = int(selector)
        if isinstance(value, Sequence) and not isinstance(value, (str, bytes, bytearray)):
            return value[idx]
        raise TypeError("M item access with numeric selector requires a list/table value.")
    if isinstance(value, Mapping):
        return _mapping_get_case_insensitive(value, _m_identifier_name(selector))
    selector_value = _parse_m_value(selector)
    if isinstance(selector_value, Mapping) and isinstance(value, Sequence) and not isinstance(value, (str, bytes, bytearray)):
        for item in value:
            if isinstance(item, Mapping) and all(_mapping_has_value(item, str(key), expected) for key, expected in selector_value.items()):
                return item
        raise KeyError(selector_text)
    raise TypeError("M item access preview supports numeric indexes and record selectors only.")


def _eval_field_access(value: Any, field_text: str) -> Any:
    field = _m_identifier_name(field_text.strip())
    if isinstance(value, Mapping):
        return _mapping_get_case_insensitive(value, field)
    raise TypeError("M field access requires a record value.")


def _eval_variable_access_chain(text: str, variables: Mapping[str, Any] | None) -> Any:
    if not variables:
        return _M_ACCESS_SENTINEL
    match = re.match(r'^(#"(?:[^"]|"")*"|[A-Za-z_][A-Za-z0-9_]*)\s*(.+)$', text, flags=re.DOTALL)
    if not match:
        return _M_ACCESS_SENTINEL
    base = _m_identifier_name(match.group(1))
    if base not in variables:
        return _M_ACCESS_SENTINEL
    tokens = _access_tokenize_suffix(match.group(2).strip())
    if not tokens:
        return _M_ACCESS_SENTINEL
    value = variables[base]
    for kind, body in tokens:
        if kind == "item":
            value = _eval_item_access(value, body)
        else:
            value = _eval_field_access(value, body)
    return value


def _eval_row_access_chain(text: str, columns: Sequence[Mapping[str, Any]], row: Sequence[Any]) -> Any:
    if not text.strip().startswith("["):
        return _M_ACCESS_SENTINEL
    tokens: list[tuple[str, bool]] = []
    i = 0
    source = text.strip()
    while i < len(source):
        if source[i].isspace():
            i += 1
            continue
        if source[i] != "[":
            return _M_ACCESS_SENTINEL
        end = source.find("]", i + 1)
        if end < 0:
            return _M_ACCESS_SENTINEL
        field = source[i + 1 : end].strip()
        if not field or "=" in field:
            return _M_ACCESS_SENTINEL
        i = end + 1
        optional = False
        if i < len(source) and source[i] == "?":
            optional = True
            i += 1
        tokens.append((_m_identifier_name(field), optional))
    if not tokens:
        return _M_ACCESS_SENTINEL
    first, optional = tokens[0]
    try:
        idx = _find_column_index(columns, first)
        value = row[idx] if idx < len(row) else None
    except Exception:
        if optional:
            value = None
        else:
            raise
    for field, optional in tokens[1:]:
        if value is None and optional:
            return None
        if not isinstance(value, Mapping):
            if optional:
                return None
            raise TypeError("M optional field access requires a record value.")
        try:
            value = _mapping_get_case_insensitive(value, field)
        except KeyError:
            if optional:
                return None
            raise
    return value


def _eval_list_lambda(lambda_expression: str, item: Any, *, variables: Mapping[str, Any] | None = None) -> Any:
    expr = lambda_expression.strip()
    if not expr.lower().startswith("each ") and re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*(?:\.[A-Za-z0-9_]+)+", expr):
        expr = f"{expr}(_)"
    zero_arg = re.match(r"^\(\s*\)\s*=>\s*(.+)$", expr, flags=re.DOTALL)
    if zero_arg:
        return _eval_m_expression([], [], zero_arg.group(1), variables=variables)
    arrow = re.match(r"^\(?\s*([A-Za-z_][A-Za-z0-9_]*)(?:\s+as\s+[^)=]+)?\s*\)?\s*(?:as\s+[^=]+)?=>\s*(.+)$", expr, flags=re.DOTALL | re.IGNORECASE)
    if arrow:
        param, body = arrow.groups()
        local_vars: dict[str, Any] = dict(variables or {})
        local_vars[param] = item
        return _eval_m_expression([], [], body, current_value=item, variables=local_vars)
    return _eval_m_expression([], [], expr, current_value=item, variables=variables)


def _eval_list_generate(
    raw_args: Sequence[str],
    *,
    variables: Mapping[str, Any] | None = None,
    max_iterations: int = 10000,
) -> list[Any]:
    if len(raw_args) < 3:
        raise ValueError("List.Generate preview needs initial, condition, and next functions.")
    candidate = _eval_list_lambda(raw_args[0], None, variables=variables)
    out: list[Any] = []
    iterations = 0
    while bool(_eval_list_lambda(raw_args[1], candidate, variables=variables)):
        out.append(_eval_list_lambda(raw_args[3], candidate, variables=variables) if len(raw_args) > 3 else candidate)
        candidate = _eval_list_lambda(raw_args[2], candidate, variables=variables)
        iterations += 1
        if iterations > max_iterations:
            raise ValueError("List.Generate preview exceeded the deterministic iteration guard.")
    return out


def _eval_binary_lambda(
    lambda_expression: str,
    first: Any,
    second: Any,
    *,
    variables: Mapping[str, Any] | None = None,
) -> Any:
    expr = lambda_expression.strip()
    arrow_idx = _top_level_arrow_index(expr)
    if arrow_idx is None:
        raise ValueError("List.Accumulate preview needs a two-argument accumulator function.")
    params = _function_parameter_names(expr[:arrow_idx])
    if len(params) < 2:
        raise ValueError("List.Accumulate preview needs accumulator parameters for state and item.")
    local_vars: dict[str, Any] = dict(variables or {})
    local_vars[params[0]] = first
    local_vars[params[1]] = second
    return _eval_m_expression([], [], expr[arrow_idx + 2 :].strip(), current_value=second, variables=local_vars)


def _eval_list_accumulate(
    raw_args: Sequence[str],
    *,
    columns: Sequence[Mapping[str, Any]],
    row: Sequence[Any],
    current_value: Any = None,
    variables: Mapping[str, Any] | None = None,
) -> Any:
    if len(raw_args) < 3:
        raise ValueError("List.Accumulate preview needs list, seed, and accumulator function.")
    values = _eval_m_expression(columns, row, raw_args[0], current_value=current_value, variables=variables)
    state = _eval_m_expression(columns, row, raw_args[1], current_value=current_value, variables=variables)
    for item in values if isinstance(values, list) else []:
        state = _eval_binary_lambda(raw_args[2], state, item, variables=variables)
    return state


def _is_m_lambda_text(text: str) -> bool:
    value = text.strip()
    return value.lower().startswith("each ") or _top_level_arrow_index(value) is not None


def _eval_list_predicate_window(
    fn_name: str,
    raw_args: Sequence[str],
    *,
    columns: Sequence[Mapping[str, Any]],
    row: Sequence[Any],
    current_value: Any = None,
    variables: Mapping[str, Any] | None = None,
) -> list[Any]:
    if len(raw_args) < 2:
        raise ValueError(f"{fn_name} preview needs a list and count/condition.")
    values = _eval_m_expression(columns, row, raw_args[0], current_value=current_value, variables=variables)
    if not isinstance(values, list):
        return []
    predicate = raw_args[1]
    if not _is_m_lambda_text(predicate):
        return _eval_function_call(
            fn_name,
            [
                values,
                _eval_m_expression(columns, row, predicate, current_value=current_value, variables=variables),
            ],
            variables=variables,
        )
    upper = fn_name.upper()
    if upper in {"LIST.FIRSTN", "LIST.LASTN"}:
        source = list(reversed(values)) if upper == "LIST.LASTN" else values
        out: list[Any] = []
        for item in source:
            if not bool(_eval_list_lambda(predicate, item, variables=variables)):
                break
            out.append(item)
        return list(reversed(out)) if upper == "LIST.LASTN" else out
    if upper in {"LIST.SKIP", "LIST.REMOVEFIRSTN"}:
        offset = 0
        for item in values:
            if not bool(_eval_list_lambda(predicate, item, variables=variables)):
                break
            offset += 1
        return values[offset:]
    if upper == "LIST.REMOVELASTN":
        suffix = 0
        for item in reversed(values):
            if not bool(_eval_list_lambda(predicate, item, variables=variables)):
                break
            suffix += 1
        return values[: len(values) - suffix] if suffix else list(values)
    return list(values)


def _eval_list_matches(
    fn_name: str,
    raw_args: Sequence[str],
    *,
    columns: Sequence[Mapping[str, Any]],
    row: Sequence[Any],
    current_value: Any = None,
    variables: Mapping[str, Any] | None = None,
) -> bool:
    if len(raw_args) < 2:
        raise ValueError(f"{fn_name} preview needs a list and predicate.")
    values = _eval_m_expression(columns, row, raw_args[0], current_value=current_value, variables=variables)
    if not isinstance(values, list):
        values = []
    results = [bool(_eval_list_lambda(raw_args[1], item, variables=variables)) for item in values]
    return all(results) if fn_name.upper() == "LIST.MATCHESALL" else any(results)


def _comparison_key(
    item: Any,
    criteria_value: Any = None,
    criteria_text: str | None = None,
    *,
    variables: Mapping[str, Any] | None = None,
) -> Any:
    if criteria_text and _is_m_lambda_text(criteria_text):
        return _eval_list_lambda(criteria_text, item, variables=variables)
    if isinstance(criteria_value, (MFunctionValue, MBuiltinFunctionValue)):
        return _invoke_m_function(criteria_value, [item], variables=variables)
    if isinstance(criteria_value, str) and isinstance(item, Mapping):
        try:
            return _mapping_get_case_insensitive(item, criteria_value)
        except KeyError:
            return None
    return item


def _sortable_key(value: Any) -> Any:
    if isinstance(value, (_dt.datetime, _dt.date, _dt.time)):
        return value.isoformat()
    if isinstance(value, _dt.timedelta):
        return value.total_seconds()
    return value


def _sort_by_comparison(
    values: Sequence[Any],
    *,
    criteria_value: Any = None,
    criteria_text: str | None = None,
    descending: bool = False,
    variables: Mapping[str, Any] | None = None,
) -> list[Any]:
    def key(item: Any) -> Any:
        selected = _comparison_key(item, criteria_value, criteria_text, variables=variables)
        sortable = _sortable_key(selected)
        return (sortable is None, sortable)

    try:
        return sorted(values, key=key, reverse=descending)
    except Exception:
        return sorted(values, key=lambda item: str(_value_key(key(item))), reverse=descending)


def _eval_list_minmaxn(
    fn_name: str,
    raw_args: Sequence[str],
    *,
    columns: Sequence[Mapping[str, Any]],
    row: Sequence[Any],
    current_value: Any = None,
    variables: Mapping[str, Any] | None = None,
) -> list[Any]:
    if len(raw_args) < 2:
        raise ValueError(f"{fn_name} preview needs a list and count/condition.")
    values = _eval_m_expression(columns, row, raw_args[0], current_value=current_value, variables=variables)
    values = list(values) if isinstance(values, list) else []
    condition_text = raw_args[1] if _is_m_lambda_text(raw_args[1]) else None
    criteria_text = raw_args[2] if len(raw_args) > 2 and _is_m_lambda_text(raw_args[2]) else None
    criteria_value = None if criteria_text else (
        _eval_m_expression(columns, row, raw_args[2], current_value=current_value, variables=variables)
        if len(raw_args) > 2
        else None
    )
    descending = fn_name.upper() == "LIST.MAXN"
    ordered = _sort_by_comparison(
        values,
        criteria_value=criteria_value,
        criteria_text=criteria_text,
        descending=descending,
        variables=variables,
    )
    if condition_text:
        out: list[Any] = []
        for item in ordered:
            if not bool(_eval_list_lambda(condition_text, item, variables=variables)):
                break
            out.append(item)
        return out
    count_value = _eval_m_expression(columns, row, raw_args[1], current_value=current_value, variables=variables)
    count = 1 if count_value is None else max(0, int(_to_number(count_value) or 0))
    return ordered[:count]


def _eval_table_minmax(
    fn_name: str,
    raw_args: Sequence[str],
    *,
    columns: Sequence[Mapping[str, Any]],
    row: Sequence[Any],
    current_value: Any = None,
    variables: Mapping[str, Any] | None = None,
) -> Any:
    if len(raw_args) < 2:
        raise ValueError(f"{fn_name} preview needs a table and comparison criteria.")
    table_value = _eval_m_expression(columns, row, raw_args[0], current_value=current_value, variables=variables)
    records = [dict(record) for record in table_value if isinstance(record, Mapping)] if isinstance(table_value, list) else []
    default = (
        _eval_m_expression(columns, row, raw_args[2], current_value=current_value, variables=variables)
        if len(raw_args) > 2
        else None
    )
    if not records:
        return default
    criteria_text = raw_args[1] if _is_m_lambda_text(raw_args[1]) else None
    criteria_value = None if criteria_text else _eval_m_expression(columns, row, raw_args[1], current_value=current_value, variables=variables)
    ordered = _sort_by_comparison(
        records,
        criteria_value=criteria_value,
        criteria_text=criteria_text,
        descending=fn_name.upper() == "TABLE.MAX",
        variables=variables,
    )
    return ordered[0] if ordered else default


def _eval_table_minmaxn(
    fn_name: str,
    raw_args: Sequence[str],
    *,
    columns: Sequence[Mapping[str, Any]],
    row: Sequence[Any],
    current_value: Any = None,
    variables: Mapping[str, Any] | None = None,
) -> list[dict[str, Any]]:
    if len(raw_args) < 3:
        raise ValueError(f"{fn_name} preview needs a table, comparison criteria, and count/condition.")
    table_value = _eval_m_expression(columns, row, raw_args[0], current_value=current_value, variables=variables)
    records = [dict(record) for record in table_value if isinstance(record, Mapping)] if isinstance(table_value, list) else []
    criteria_text = raw_args[1] if _is_m_lambda_text(raw_args[1]) else None
    criteria_value = None if criteria_text else _eval_m_expression(columns, row, raw_args[1], current_value=current_value, variables=variables)
    descending = fn_name.upper() == "TABLE.MAXN"
    ordered = _sort_by_comparison(
        records,
        criteria_value=criteria_value,
        criteria_text=criteria_text,
        descending=descending,
        variables=variables,
    )
    count_or_condition = raw_args[2]
    if _is_m_lambda_text(count_or_condition):
        out: list[dict[str, Any]] = []
        for record in ordered:
            if not bool(_eval_list_lambda(count_or_condition, record, variables=variables)):
                break
            out.append(record)
        return out
    count_value = _eval_m_expression(columns, row, count_or_condition, current_value=current_value, variables=variables)
    count = max(0, int(_to_number(count_value) or 0))
    return ordered[:count]


def _table_records(columns: Sequence[Mapping[str, Any]], rows: Sequence[Sequence[Any]]) -> list[dict[str, Any]]:
    names = [str(col.get("name") or "") for col in columns]
    return [
        {name: row[idx] if idx < len(row) else None for idx, name in enumerate(names)}
        for row in rows
    ]


def _table_value_columns_rows(value: Any) -> tuple[list[dict[str, Any]], list[list[Any]]]:
    return _constant_table_from_records(value if isinstance(value, list) else [])


def _table_schema_records(value: Any) -> list[dict[str, Any]]:
    columns, rows = _table_value_columns_rows(value)
    out: list[dict[str, Any]] = []
    for idx, col in enumerate(columns):
        type_text = str(col.get("m_type") or "").strip()
        if not type_text:
            sample = next((row[idx] for row in rows if idx < len(row) and row[idx] is not None), None)
            type_text = _m_type_text(sample) if sample is not None else "type any"
        kind = _m_kind_from_type(type_text)
        out.append({
            "Name": str(col.get("name") or ""),
            "Position": idx,
            "TypeName": _m_type_name_from_kind(kind),
            "Kind": kind,
            "IsNullable": any(idx >= len(row) or row[idx] is None for row in rows),
            "NumericPrecisionBase": 10 if kind == "Number" else None,
            "NumericPrecision": None,
            "NumericScale": None,
            "DateTimePrecision": None,
            "MaxLength": None,
            "IsVariableLength": kind in {"Text", "Binary"},
            "NativeTypeName": None,
            "NativeDefaultExpression": None,
            "Description": None,
        })
    return out


def _table_profile_records(value: Any) -> list[dict[str, Any]]:
    columns, rows = _table_value_columns_rows(value)
    out: list[dict[str, Any]] = []
    for idx, col in enumerate(columns):
        values = [row[idx] if idx < len(row) else None for row in rows]
        non_null = [item for item in values if item is not None and item != ""]
        numeric = _numeric_values(non_null)
        if len(numeric) == len(non_null):
            minimum: Any = min(numeric) if numeric else None
            maximum: Any = max(numeric) if numeric else None
            average: Any = (sum(numeric) / len(numeric)) if numeric else None
            if len(numeric) > 1:
                mean = sum(numeric) / len(numeric)
                stddev: Any = _math.sqrt(sum((item - mean) ** 2 for item in numeric) / (len(numeric) - 1))
            else:
                stddev = None
        else:
            comparable = list(non_null)
            try:
                ordered = sorted(comparable)
            except Exception:
                ordered = sorted(comparable, key=lambda item: str(_value_key(item)))
            minimum = ordered[0] if ordered else None
            maximum = ordered[-1] if ordered else None
            average = None
            stddev = None
        out.append({
            "Column": str(col.get("name") or ""),
            "Min": minimum,
            "Max": maximum,
            "Average": average,
            "StandardDeviation": stddev,
            "Count": len(values),
            "NullCount": len(values) - len(non_null),
            "DistinctCount": len({_value_key(item) for item in non_null}),
        })
    return out


def _invoke_m_function(function: MFunctionValue | MBuiltinFunctionValue, args: Sequence[Any], variables: Mapping[str, Any] | None = None) -> Any:
    if isinstance(function, MBuiltinFunctionValue):
        upper = function.name.upper()
        if upper.startswith("SPLITTER."):
            return _apply_splitter(function, args[0] if args else None, variables=variables)
        if upper.startswith("COMBINER."):
            return _apply_combiner(function, args[0] if args else [])
        if upper.startswith("BINARYFORMAT."):
            value, _offset = _binary_format_parse(function, bytes(_unwrap_metadata(args[0] if args else b"") or b""), 0)
            return value
        raise ValueError(f"Unsupported built-in M function helper: {function.name}")
    local_vars: dict[str, Any] = dict(variables or {})
    for idx, name in enumerate(function.parameters):
        local_vars[name] = args[idx] if idx < len(args) else None
    body = function.body.strip()
    if re.match(r"^let\b", body, flags=re.IGNORECASE):
        result = evaluate_preview_table(raw_m=body, columns=[], rows=[], initial_variables=local_vars)
        if result.blocked_steps:
            first = result.blocked_steps[0]
            raise ValueError(str(first.get("reason") or "Function helper preview failed."))
        return _table_records(result.columns, result.rows)
    return _eval_m_expression([], [], body, variables=local_vars)


def _evaluate_m_query_value(value: MQueryValue, *, variables: Mapping[str, Any] | None = None) -> Any:
    result = evaluate_preview_table(raw_m=value.raw_m, columns=[], rows=[], initial_variables=variables)
    if result.blocked_steps:
        first = result.blocked_steps[0]
        raise ValueError(str(first.get("reason") or "Helper query preview failed."))
    return _table_records(result.columns, result.rows)


def _eval_m_expression(
    columns: Sequence[Mapping[str, Any]],
    row: Sequence[Any],
    expression: str,
    *,
    current_value: Any = None,
    variables: Mapping[str, Any] | None = None,
) -> Any:
    text = _strip_outer_parens(expression.strip())
    if text.lower().startswith("each "):
        text = _strip_outer_parens(text[5:].strip())
    meta_split = _split_top_level_keyword(text, "meta")
    if meta_split:
        value_expr, metadata_expr = meta_split
        metadata_value = _eval_m_expression(columns, row, metadata_expr, current_value=current_value, variables=variables)
        return MMetadataValue(
            value=_eval_m_expression(columns, row, value_expr, current_value=current_value, variables=variables),
            metadata=dict(metadata_value) if isinstance(metadata_value, Mapping) else {},
        )
    if text == "_":
        return current_value
    identifier = _m_identifier_name(text)
    if variables and identifier in variables:
        value = variables[identifier]
        if isinstance(value, MQueryValue):
            return _evaluate_m_query_value(value, variables=variables)
        return value
    constant_value = _m_constant_value(identifier)
    if constant_value is not None:
        return constant_value
    function_literal = _parse_m_function_value(text)
    if function_literal is not None:
        return function_literal
    access_value = _eval_variable_access_chain(text, variables)
    if access_value is not _M_ACCESS_SENTINEL:
        return access_value
    current_record_ref = re.fullmatch(r"\[([^\]]+)\]", text)
    if current_record_ref and isinstance(current_value, Mapping) and "=" not in current_record_ref.group(1):
        return _mapping_get_case_insensitive(current_value, current_record_ref.group(1))
    row_access_value = _eval_row_access_chain(text, columns, row)
    if row_access_value is not _M_ACCESS_SENTINEL:
        return row_access_value
    if text.lower().startswith("not "):
        return not bool(_eval_m_expression(columns, row, text[4:].strip(), current_value=current_value, variables=variables))
    if text.lower().startswith("error "):
        error_value = _eval_m_expression(columns, row, text[6:].strip(), current_value=current_value, variables=variables)
        raise MExpressionError(error_value if isinstance(error_value, Mapping) else _error_record("Expression.Error", str(error_value), error_value))
    if text.lower().startswith("try "):
        rest = text[4:].strip()
        otherwise = _split_top_level_keyword(rest, "otherwise")
        if otherwise:
            try_expr, fallback_expr = otherwise
            try:
                return _eval_m_expression(columns, row, try_expr, current_value=current_value, variables=variables)
            except Exception:
                return _eval_m_expression(columns, row, fallback_expr, current_value=current_value, variables=variables)
        try:
            return {
                "HasError": False,
                "Value": _eval_m_expression(columns, row, rest, current_value=current_value, variables=variables),
                "Error": None,
            }
        except MExpressionError as exc:
            return {"HasError": True, "Error": exc.error_record, "Value": None}
        except Exception as exc:
            return {"HasError": True, "Error": _error_record("Expression.Error", str(exc), None), "Value": None}
    if text.lower().startswith("if "):
        rest = text[3:].strip()
        then_split = _split_top_level_keyword(rest, "then")
        if not then_split:
            raise ValueError("Malformed if expression; missing then.")
        condition, tail = then_split
        else_split = _split_top_level_keyword(tail, "else")
        if not else_split:
            raise ValueError("Malformed if expression; missing else.")
        true_expr, false_expr = else_split
        return _eval_m_expression(columns, row, true_expr, current_value=current_value, variables=variables) if bool(_eval_m_expression(columns, row, condition, current_value=current_value, variables=variables)) else _eval_m_expression(columns, row, false_expr, current_value=current_value, variables=variables)
    for keyword, reducer in (("or", any), ("and", all)):
        split = _split_top_level_keyword(text, keyword)
        if split:
            left, right = split
            return reducer([
                bool(_eval_m_expression(columns, row, left, current_value=current_value, variables=variables)),
                bool(_eval_m_expression(columns, row, right, current_value=current_value, variables=variables)),
            ])
    comp = _split_top_level_operator(text, [">=", "<=", "<>", "!=", "=", ">", "<"])
    if comp:
        left, op, right = comp
        return _compare(
            _eval_m_expression(columns, row, left, current_value=current_value, variables=variables),
            op,
            _eval_m_expression(columns, row, right, current_value=current_value, variables=variables),
        )
    concat = _split_top_level_operator(text, ["&"])
    if concat:
        left, _op, right = concat
        lval = _eval_m_expression(columns, row, left, current_value=current_value, variables=variables)
        rval = _eval_m_expression(columns, row, right, current_value=current_value, variables=variables)
        if isinstance(lval, list) and isinstance(rval, list):
            return list(lval) + list(rval)
        if isinstance(lval, Mapping) and isinstance(rval, Mapping):
            out = dict(lval)
            out.update(rval)
            return out
        return ("" if lval is None else str(lval)) + ("" if rval is None else str(rval))
    add = _split_top_level_operator(text, ["+", "-"])
    if add:
        left, op, right = add
        lval = _eval_m_expression(columns, row, left, current_value=current_value, variables=variables)
        rval = _eval_m_expression(columns, row, right, current_value=current_value, variables=variables)
        return _eval_add_subtract(lval, op, rval)
    mul = _split_top_level_operator(text, ["*", "/"])
    if mul:
        left, op, right = mul
        lval = _eval_m_expression(columns, row, left, current_value=current_value, variables=variables)
        rval = _eval_m_expression(columns, row, right, current_value=current_value, variables=variables)
        if lval is None or rval is None:
            return None
        return float(lval) * float(rval) if op == "*" else (None if float(rval) == 0 else float(lval) / float(rval))
    column_ref = re.fullmatch(r"\[([^\]]+)\]", text)
    if column_ref:
        body = column_ref.group(1)
        if "=" not in body:
            if isinstance(current_value, Mapping):
                return _mapping_get_case_insensitive(current_value, body)
            idx = _find_column_index(columns, body)
            return row[idx] if idx < len(row) else None
    range_split = _split_top_level_range(text)
    if range_split is not None:
        left, right = range_split
        return _m_range_values(
            _eval_m_expression(columns, row, left, current_value=current_value, variables=variables),
            _eval_m_expression(columns, row, right, current_value=current_value, variables=variables),
        )
    list_body = _between_outer(text, "{", "}")
    if list_body is not None:
        if not list_body.strip():
            return []
        values: list[Any] = []
        for item in _split_top_level_items(list_body):
            range_item = _split_top_level_range(item)
            if range_item is not None:
                left, right = range_item
                values.extend(_m_range_values(
                    _eval_m_expression(columns, row, left, current_value=current_value, variables=variables),
                    _eval_m_expression(columns, row, right, current_value=current_value, variables=variables),
                ))
            else:
                values.append(_eval_m_expression(columns, row, item, current_value=current_value, variables=variables))
        return values
    record_body = _between_outer(text, "[", "]")
    if record_body is not None:
        record: dict[str, Any] = {}
        for item in _split_top_level_items(record_body):
            split = _split_top_level_equals(item)
            if split is None:
                return _parse_m_value(text)
            key, value_expr = split
            record[_m_identifier_name(key)] = _eval_m_expression(
                columns,
                row,
                value_expr,
                current_value=current_value,
                variables=variables,
            )
        return record
    quoted_call = re.match(r'^#"((?:""|[^"])*)"\s*\((.*)\)$', text, flags=re.DOTALL)
    if quoted_call:
        fn_label = quoted_call.group(1).replace('""', '"')
        raw_args = _function_args(text)
        function_value = variables.get(fn_label) if variables else None
        if isinstance(function_value, (MFunctionValue, MBuiltinFunctionValue)):
            args = [
                _eval_m_expression(columns, row, arg, current_value=current_value, variables=variables)
                for arg in raw_args
            ]
            return _invoke_m_function(function_value, args, variables=variables)
        if fn_label.strip().lower().startswith("transform file") and raw_args:
            binary_value = _eval_m_expression(columns, row, raw_args[0], current_value=current_value, variables=variables)
            return _csv_document_records(binary_value, promote_headers=True)
        raise ValueError(f"Unsupported quoted M function invocation: {fn_label}")
    simple_call = re.match(r"^([A-Za-z_][A-Za-z0-9_]*)\s*\(", text)
    if simple_call and text.endswith(")") and variables:
        function_value = variables.get(simple_call.group(1))
        if isinstance(function_value, (MFunctionValue, MBuiltinFunctionValue)):
            args = [
                _eval_m_expression(columns, row, arg, current_value=current_value, variables=variables)
                for arg in _function_args(text)
            ]
            return _invoke_m_function(function_value, args, variables=variables)
    call_match = re.match(r"^([A-Za-z_][A-Za-z0-9_]*(?:\.[A-Za-z0-9_]+)+)\s*\(", text)
    if call_match and text.endswith(")"):
        fn_name = call_match.group(1).upper()
        raw_args = _function_args(text)
        if fn_name == "LIST.GENERATE":
            return _eval_list_generate(raw_args, variables=variables)
        if fn_name == "LIST.ACCUMULATE":
            return _eval_list_accumulate(
                raw_args,
                columns=columns,
                row=row,
                current_value=current_value,
                variables=variables,
            )
        if fn_name in {"LIST.MINN", "LIST.MAXN"}:
            return _eval_list_minmaxn(
                fn_name,
                raw_args,
                columns=columns,
                row=row,
                current_value=current_value,
                variables=variables,
            )
        if fn_name in {"LIST.FIRSTN", "LIST.LASTN", "LIST.SKIP", "LIST.REMOVEFIRSTN", "LIST.REMOVELASTN"} and len(raw_args) > 1 and _is_m_lambda_text(raw_args[1]):
            return _eval_list_predicate_window(
                fn_name,
                raw_args,
                columns=columns,
                row=row,
                current_value=current_value,
                variables=variables,
            )
        if fn_name in {"LIST.MATCHESALL", "LIST.MATCHESANY"}:
            return _eval_list_matches(
                fn_name,
                raw_args,
                columns=columns,
                row=row,
                current_value=current_value,
                variables=variables,
            )
        if fn_name in {"TABLE.MIN", "TABLE.MAX"}:
            return _eval_table_minmax(
                fn_name,
                raw_args,
                columns=columns,
                row=row,
                current_value=current_value,
                variables=variables,
            )
        if fn_name in {"TABLE.MINN", "TABLE.MAXN"}:
            return _eval_table_minmaxn(
                fn_name,
                raw_args,
                columns=columns,
                row=row,
                current_value=current_value,
                variables=variables,
            )
        if fn_name in {"LIST.TRANSFORM", "LIST.SELECT"}:
            if len(raw_args) < 2:
                raise ValueError(f"{call_match.group(1)} needs a list and lambda.")
            values = _eval_m_expression(columns, row, raw_args[0], current_value=current_value, variables=variables)
            if not isinstance(values, list):
                return []
            if fn_name == "LIST.TRANSFORM":
                return [_eval_list_lambda(raw_args[1], item, variables=variables) for item in values]
            return [item for item in values if bool(_eval_list_lambda(raw_args[1], item, variables=variables))]
        if fn_name == "LIST.TRANSFORMMANY":
            if len(raw_args) < 2:
                raise ValueError("List.TransformMany preview needs a list and collection transform.")
            values = _eval_m_expression(columns, row, raw_args[0], current_value=current_value, variables=variables)
            if not isinstance(values, list):
                return []
            result_function = _parse_m_function_value(raw_args[2]) if len(raw_args) > 2 else None
            out: list[Any] = []
            for item in values:
                children = _eval_list_lambda(raw_args[1], item, variables=variables)
                child_values = children if isinstance(children, list) else []
                for child in child_values:
                    out.append(_invoke_m_function(result_function, [item, child], variables=variables) if result_function else child)
            return out
        if fn_name == "TABLE.SELECTROWS":
            if len(raw_args) < 2:
                raise ValueError("Table.SelectRows preview needs a table and predicate.")
            table_value = _eval_m_expression(columns, row, raw_args[0], current_value=current_value, variables=variables)
            table_cols, table_rows = _table_value_columns_rows(table_value)
            out_cols, out_rows = _select_rows(table_cols, table_rows, text, variables=variables)
            return _table_records(out_cols, out_rows)
        if fn_name in {"TABLE.MATCHESALLROWS", "TABLE.MATCHESANYROWS"}:
            if len(raw_args) < 2:
                raise ValueError(f"{call_match.group(1)} preview needs a table and predicate.")
            table_value = _eval_m_expression(columns, row, raw_args[0], current_value=current_value, variables=variables)
            records = [dict(record) for record in table_value if isinstance(record, Mapping)] if isinstance(table_value, list) else []
            results = [bool(_eval_list_lambda(raw_args[1], record, variables=variables)) for record in records]
            return all(results) if fn_name == "TABLE.MATCHESALLROWS" else any(results)
        if fn_name == "TABLE.TRANSFORMROWS":
            if len(raw_args) < 2:
                raise ValueError("Table.TransformRows preview needs a table and transform function.")
            table_value = _eval_m_expression(columns, row, raw_args[0], current_value=current_value, variables=variables)
            records = [dict(record) for record in table_value if isinstance(record, Mapping)] if isinstance(table_value, list) else []
            return [_eval_list_lambda(raw_args[1], record, variables=variables) for record in records]
        if fn_name == "RECORD.TRANSFORMFIELDS":
            if len(raw_args) < 2:
                raise ValueError("Record.TransformFields preview needs a record and transform operations.")
            record_value = _eval_m_expression(columns, row, raw_args[0], current_value=current_value, variables=variables)
            missing_field_value = (
                _eval_m_expression(columns, row, raw_args[2], current_value=current_value, variables=variables)
                if len(raw_args) > 2
                else None
            )
            missing_field = _missing_field_mode_from_value(missing_field_value)
            specs = _record_transform_specs_from_text(raw_args[1], variables=variables)
            return _record_transform_fields(
                record_value,
                specs or [],
                missing_field=missing_field,
                variables=variables,
            )
        args = [
            _eval_m_expression(columns, row, arg, current_value=current_value, variables=variables)
            for arg in raw_args
        ]
        if fn_name == "FILE.CONTENTS":
            return _file_contents_binary_ref(args[0] if args else "", variables)
        return _eval_function_call(call_match.group(1), args, variables=variables)
    return _parse_m_value(text)


def _constant_table_from_records(records: list[Any]) -> tuple[list[dict[str, Any]], list[list[Any]]]:
    names: list[str] = []
    seen: set[str] = set()
    clean_records: list[Mapping[str, Any]] = []
    for record in records:
        if not isinstance(record, Mapping):
            continue
        clean_records.append(record)
        for name in record.keys():
            key = str(name).upper()
            if key not in seen:
                seen.add(key)
                names.append(str(name))
    return [{"name": name} for name in names], [[record.get(name) for name in names] for record in clean_records]


def _constant_table_from_rows(rows_value: Any, columns_value: Any) -> tuple[list[dict[str, Any]], list[list[Any]]]:
    rows = rows_value if isinstance(rows_value, list) else []
    names = [str(item) for item in columns_value] if isinstance(columns_value, list) else []
    if not names and rows and isinstance(rows[0], list):
        names = [f"Column{i + 1}" for i in range(len(rows[0]))]
    return [{"name": name} for name in names], [list(row) if isinstance(row, list) else [row] for row in rows]


def _constant_table_from_columns(columns_value: Any, names_value: Any) -> tuple[list[dict[str, Any]], list[list[Any]]]:
    columns = columns_value if isinstance(columns_value, list) else []
    names = [str(item) for item in names_value] if isinstance(names_value, list) else [f"Column{i + 1}" for i in range(len(columns))]
    max_len = max((len(col) for col in columns if isinstance(col, list)), default=0)
    rows: list[list[Any]] = []
    for row_idx in range(max_len):
        row: list[Any] = []
        for col in columns:
            values = col if isinstance(col, list) else []
            row.append(values[row_idx] if row_idx < len(values) else None)
        rows.append(row)
    return [{"name": name} for name in names], rows


def _constant_table_from_list(values: Any, names_value: Any) -> tuple[list[dict[str, Any]], list[list[Any]]]:
    items = values if isinstance(values, list) else []
    names = [str(item) for item in names_value] if isinstance(names_value, list) and names_value else ["Column1"]
    if len(names) == 1:
        return [{"name": names[0]}], [[item] for item in items]
    rows: list[list[Any]] = []
    for item in items:
        if isinstance(item, list):
            rows.append([item[idx] if idx < len(item) else None for idx, _name in enumerate(names)])
        elif isinstance(item, Mapping):
            rows.append([item.get(name) for name in names])
        else:
            rows.append([item] + [None for _name in names[1:]])
    return [{"name": name} for name in names], rows


def _type_table_column_names(expression: str) -> list[str]:
    match = re.search(r"\btype\s+table\s*\[(.*)\]", expression, flags=re.IGNORECASE | re.DOTALL)
    if not match:
        return []
    names: list[str] = []
    for item in _split_top_level_items(match.group(1)):
        if "=" not in item:
            continue
        raw_name = item.split("=", 1)[0].strip()
        if raw_name:
            names.append(_m_identifier_name(raw_name))
    return names


def _constant_table(expression: str, *, variables: Mapping[str, Any] | None = None) -> tuple[list[dict[str, Any]], list[list[Any]]]:
    args = _function_args(expression)
    op = expression.strip().split("(", 1)[0].strip().upper()
    if op == "#TABLE":
        if len(args) < 2:
            raise ValueError("#table preview needs column names and row values.")
        return _constant_table_from_rows(_parse_m_value(args[1]), _parse_m_value(args[0]))
    if op == "TABLE.FROMRECORDS":
        if not args:
            raise ValueError("Table.FromRecords preview needs record values.")
        records = _eval_m_expression([], [], args[0], variables=variables)
        return _constant_table_from_records(records if isinstance(records, list) else [])
    if op == "TABLE.FROMROWS":
        if not args:
            raise ValueError("Table.FromRows preview needs row values.")
        rows_value = _eval_m_expression([], [], args[0], variables=variables)
        columns_value = _parse_m_value(args[1]) if len(args) > 1 else []
        if not isinstance(columns_value, list) and len(args) > 1:
            columns_value = _type_table_column_names(args[1])
        return _constant_table_from_rows(rows_value, columns_value)
    if op == "TABLE.FROMCOLUMNS":
        if not args:
            raise ValueError("Table.FromColumns preview needs column values.")
        columns_value = _eval_m_expression([], [], args[0], variables=variables)
        names_value = _parse_m_value(args[1]) if len(args) > 1 else []
        return _constant_table_from_columns(columns_value, names_value)
    if op == "TABLE.FROMLIST":
        if not args:
            raise ValueError("Table.FromList preview needs list values.")
        values = _eval_m_expression([], [], args[0], variables=variables)
        names_value = _parse_m_value(args[2]) if len(args) > 2 else []
        return _constant_table_from_list(values, names_value)
    raise ValueError(f"Unsupported constant table expression: {op}")


def _value_preview_table(value: Any) -> tuple[list[dict[str, Any]], list[list[Any]]]:
    value = _unwrap_metadata(value)
    if isinstance(value, list):
        if value and all(isinstance(item, Mapping) for item in value):
            return _constant_table_from_records(value)
        return [{"name": "Value"}], [[item] for item in value]
    if isinstance(value, Mapping):
        return [{"name": "Name"}, {"name": "Value"}], [[str(key), item] for key, item in value.items()]
    return [{"name": "Value"}], [[value]]


def _compare(left: Any, op: str, right: Any) -> bool:
    left = _unwrap_metadata(left)
    right = _unwrap_metadata(right)
    if left is None:
        if op in {"=", "=="}:
            return right is None
        if op in {"<>", "!="}:
            return right is not None
        return False
    try:
        if isinstance(right, (int, float)) and not isinstance(left, (int, float)):
            left = float(left)
    except Exception:
        pass
    if op in {"=", "=="}:
        return left == right
    if op in {"<>", "!="}:
        return left != right
    try:
        if op == ">":
            return left > right
        if op == ">=":
            return left >= right
        if op == "<":
            return left < right
        if op == "<=":
            return left <= right
    except Exception:
        return False
    return False


def _select_rows(
    columns: list[dict[str, Any]],
    rows: list[list[Any]],
    expression: str,
    *,
    variables: Mapping[str, Any] | None = None,
) -> tuple[list[dict[str, Any]], list[list[Any]]]:
    each_match = re.search(r",\s*each\s+(.+?)\)\s*$", expression, flags=re.IGNORECASE | re.DOTALL)
    if not each_match:
        raise ValueError("Table.SelectRows preview supports each predicates only.")
    predicate = "each " + each_match.group(1).strip()
    return columns, [row for row in rows if bool(_eval_m_expression(columns, row, predicate, variables=variables))]


def _distinct_rows(columns: list[dict[str, Any]], rows: list[list[Any]]) -> tuple[list[dict[str, Any]], list[list[Any]]]:
    seen: set[tuple[Any, ...]] = set()
    out: list[list[Any]] = []
    for row in rows:
        key = tuple(row)
        if key in seen:
            continue
        seen.add(key)
        out.append(row)
    return columns, out


def _sort_specs_from_criteria(criteria: Any) -> list[tuple[str, bool]]:
    if isinstance(criteria, str):
        return [(criteria, False)]
    specs: list[tuple[str, bool]] = []
    if isinstance(criteria, list):
        if criteria and isinstance(criteria[0], str) and (len(criteria) == 1 or not isinstance(criteria[1], (list, tuple))):
            order = str(criteria[1]).strip().upper() if len(criteria) > 1 else ""
            return [(str(criteria[0]), order.endswith("DESCENDING"))]
        for item in criteria:
            if isinstance(item, str):
                specs.append((item, False))
            elif isinstance(item, (list, tuple)) and item:
                order = str(item[1]).strip().upper() if len(item) > 1 else ""
                specs.append((str(item[0]), order.endswith("DESCENDING")))
    return specs


def _sort_rows_by_criteria(
    columns: list[dict[str, Any]],
    rows: list[list[Any]],
    criteria: Any,
) -> tuple[list[dict[str, Any]], list[list[Any]]]:
    specs = _sort_specs_from_criteria(criteria)
    if not specs:
        raise ValueError("Table.Sort preview needs at least one sort column.")
    sorted_rows = [list(row) for row in rows]
    for name, descending in reversed(specs):
        idx = _find_column_index(columns, name)
        sorted_rows.sort(
            key=lambda row: (
                row[idx] is None if idx < len(row) else True,
                _sortable_key(row[idx]) if idx < len(row) else None,
            ),
            reverse=descending,
        )
    return columns, sorted_rows


def _sort_rows(columns: list[dict[str, Any]], rows: list[list[Any]], expression: str) -> tuple[list[dict[str, Any]], list[list[Any]]]:
    args = _function_args(expression)
    criteria = _parse_m_value(args[1]) if len(args) > 1 else []
    return _sort_rows_by_criteria(columns, rows, criteria)


def _criteria_field_names(criteria: Any, wanted: Mapping[str, Any]) -> list[str]:
    if isinstance(criteria, str) and criteria:
        return [criteria]
    if isinstance(criteria, list):
        fields: list[str] = []
        for item in criteria:
            if isinstance(item, str):
                fields.append(item)
            elif isinstance(item, (list, tuple)) and item:
                fields.append(str(item[0]))
        if fields:
            return fields
    return [str(key) for key in wanted.keys()]


def _table_row_matches(record: Mapping[str, Any], wanted: Mapping[str, Any], criteria: Any = None) -> bool:
    for field in _criteria_field_names(criteria, wanted):
        try:
            left = _mapping_get_case_insensitive(record, field)
            right = _mapping_get_case_insensitive(wanted, field)
        except KeyError:
            return False
        if _value_key(left) != _value_key(right):
            return False
    return True


def _occurrence_text(value: Any = None) -> str:
    text = str(value or "").strip().upper()
    if not text:
        return "FIRST"
    if text.endswith(".ALL") or text == "ALL":
        return "ALL"
    if text.endswith(".LAST") or text == "LAST":
        return "LAST"
    if text.endswith(".FIRST") or text == "FIRST":
        return "FIRST"
    return text


def _table_positions(
    records: Sequence[Mapping[str, Any]],
    wanted_rows: Sequence[Mapping[str, Any]],
    occurrence: Any = None,
    criteria: Any = None,
) -> Any:
    matches = [
        idx for idx, record in enumerate(records)
        if any(_table_row_matches(record, wanted, criteria) for wanted in wanted_rows)
    ]
    occurrence_text = _occurrence_text(occurrence)
    if occurrence_text == "ALL":
        return matches
    if occurrence_text == "LAST":
        return matches[-1] if matches else -1
    try:
        ordinal = int(_to_number(occurrence) or 0) if occurrence is not None else 0
        return matches[ordinal] if 0 <= ordinal < len(matches) else -1
    except Exception:
        return matches[0] if matches else -1


def _fill_down(columns: list[dict[str, Any]], rows: list[list[Any]], names: list[str]) -> list[list[Any]]:
    indexes = [_find_column_index(columns, name) for name in names]
    last_seen: dict[int, Any] = {}
    out: list[list[Any]] = []
    for row in rows:
        out_row = list(row)
        for idx in indexes:
            value = out_row[idx] if idx < len(out_row) else None
            if value is None or value == "":
                if idx in last_seen and idx < len(out_row):
                    out_row[idx] = last_seen[idx]
            else:
                last_seen[idx] = value
        out.append(out_row)
    return out


def _fill_up(columns: list[dict[str, Any]], rows: list[list[Any]], names: list[str]) -> list[list[Any]]:
    indexes = [_find_column_index(columns, name) for name in names]
    next_seen: dict[int, Any] = {}
    out_reversed: list[list[Any]] = []
    for row in reversed(rows):
        out_row = list(row)
        for idx in indexes:
            value = out_row[idx] if idx < len(out_row) else None
            if value is None or value == "":
                if idx in next_seen and idx < len(out_row):
                    out_row[idx] = next_seen[idx]
            else:
                next_seen[idx] = value
        out_reversed.append(out_row)
    return list(reversed(out_reversed))


def _replace_values(columns: list[dict[str, Any]], rows: list[list[Any]], expression: str) -> list[list[Any]]:
    match = re.search(
        r"Table\.ReplaceValue\([^,]+,\s*(.+?)\s*,\s*(.+?)\s*,\s*Replacer\.(ReplaceValue|ReplaceText)\s*,\s*\{(.+)\}\s*\)",
        expression,
        flags=re.IGNORECASE | re.DOTALL,
    )
    if not match:
        raise ValueError("Table.ReplaceValue preview supports literal old/new values and quoted target columns only.")
    old_value = _parse_literal(match.group(1))
    new_value = _parse_literal(match.group(2))
    replacer = match.group(3).lower()
    target_columns = _string_literals(match.group(4))
    indexes = [_find_column_index(columns, name) for name in target_columns]
    out: list[list[Any]] = []
    for row in rows:
        out_row = list(row)
        for idx in indexes:
            if idx >= len(out_row):
                continue
            value = out_row[idx]
            if replacer == "replacetext" and value is not None:
                out_row[idx] = str(value).replace(str(old_value), str(new_value))
            elif value == old_value:
                out_row[idx] = new_value
        out.append(out_row)
    return out


def _is_error_value(value: Any) -> bool:
    return isinstance(value, MErrorValue)


def _replace_error_values(columns: list[dict[str, Any]], rows: list[list[Any]], expression: str) -> list[list[Any]]:
    args = _function_args(expression)
    if len(args) < 2:
        raise ValueError("Table.ReplaceErrorValues preview needs replacement pairs.")
    specs = _parse_m_value(args[1])
    if not isinstance(specs, list):
        raise ValueError("Table.ReplaceErrorValues preview needs a list of {column, replacement} pairs.")
    replacements: list[tuple[int, Any]] = []
    for spec in specs:
        if not isinstance(spec, list) or len(spec) < 2:
            continue
        idx = _find_column_index(columns, str(spec[0]))
        replacements.append((idx, spec[1]))
    out: list[list[Any]] = []
    for row in rows:
        out_row = list(row)
        for idx, replacement in replacements:
            if idx < len(out_row) and _is_error_value(out_row[idx]):
                out_row[idx] = replacement
        out.append(out_row)
    return out


def _error_target_indexes(columns: list[dict[str, Any]], expression: str) -> list[int]:
    args = _function_args(expression)
    if len(args) < 2:
        return list(range(len(columns)))
    names = _parse_m_value(args[1])
    if isinstance(names, str):
        names = [names]
    if not isinstance(names, list):
        raise ValueError("Rows-with-errors preview needs a column-name list.")
    return [_find_column_index(columns, str(name)) for name in names]


def _rows_with_errors(columns: list[dict[str, Any]], rows: list[list[Any]], expression: str, *, keep_errors: bool) -> list[list[Any]]:
    indexes = _error_target_indexes(columns, expression)
    out: list[list[Any]] = []
    for row in rows:
        has_error = any(idx < len(row) and _is_error_value(row[idx]) for idx in indexes)
        if has_error == keep_errors:
            out.append(list(row))
    return out


def _eval_simple_each(
    columns: Sequence[Mapping[str, Any]],
    row: Sequence[Any],
    expression: str,
    *,
    variables: Mapping[str, Any] | None = None,
) -> Any:
    return _eval_m_expression(columns, row, expression, variables=variables)


def _add_column(
    columns: list[dict[str, Any]],
    rows: list[list[Any]],
    expression: str,
    *,
    variables: Mapping[str, Any] | None = None,
) -> tuple[list[dict[str, Any]], list[list[Any]]]:
    args = _function_args(expression)
    names = _string_literals_after_first_arg(expression)
    if len(args) < 3 or not names:
        raise ValueError("Table.AddColumn preview needs a quoted new column name.")
    new_name = names[0]
    each_expression = args[2].strip()
    if not each_expression.lower().startswith("each "):
        raise ValueError("Table.AddColumn preview supports simple each expressions only.")
    out_cols = [dict(col) for col in columns] + [{"name": new_name}]
    out_rows = []
    for row in rows:
        try:
            value = _eval_simple_each(columns, row, each_expression, variables=variables)
        except Exception as exc:
            value = MErrorValue(str(exc))
        out_rows.append(list(row) + [value])
    return out_cols, out_rows


def _transform_columns(
    columns: list[dict[str, Any]],
    rows: list[list[Any]],
    expression: str,
    *,
    variables: Mapping[str, Any] | None = None,
) -> tuple[list[dict[str, Any]], list[list[Any]]]:
    args = _function_args(expression)
    if len(args) < 2:
        raise ValueError("Table.TransformColumns preview needs a transform operation list.")
    spec_list = _between_outer(args[1], "{", "}")
    if spec_list is None:
        raise ValueError("Table.TransformColumns preview needs a transform operation list.")
    specs: list[tuple[int, str, str | None]] = []
    out_cols = [dict(col) for col in columns]
    for raw_spec in _split_top_level_items(spec_list):
        body = _between_outer(raw_spec, "{", "}")
        if body is None:
            continue
        parts = _split_top_level_items(body)
        if len(parts) < 2:
            continue
        col_name = str(_parse_m_value(parts[0]))
        idx = _find_column_index(columns, col_name)
        transform_expr = parts[1].strip()
        if not transform_expr.lower().startswith("each ") and re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*(?:\.[A-Za-z_][A-Za-z0-9_]*)+", transform_expr):
            transform_expr = f"{transform_expr}(_)"
        type_text = parts[2].strip() if len(parts) > 2 else None
        if type_text and type_text.lower().startswith("type "):
            out_cols[idx]["m_type"] = type_text
        specs.append((idx, transform_expr, type_text))
    out_rows: list[list[Any]] = []
    for row in rows:
        out_row = list(row)
        for idx, transform_expr, _type_text in specs:
            value = out_row[idx] if idx < len(out_row) else None
            try:
                out_row[idx] = _eval_m_expression(columns, out_row, transform_expr, current_value=value, variables=variables)
            except Exception as exc:
                out_row[idx] = MErrorValue(str(exc))
        out_rows.append(out_row)
    return out_cols, out_rows


def _transform_column_names(
    columns: list[dict[str, Any]],
    rows: list[list[Any]],
    expression: str,
    *,
    variables: Mapping[str, Any] | None = None,
) -> tuple[list[dict[str, Any]], list[list[Any]]]:
    args = _function_args(expression)
    if len(args) < 2:
        raise ValueError("Table.TransformColumnNames preview needs a name transform function.")
    transform_expr = args[1].strip()
    if not transform_expr.lower().startswith("each ") and re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*(?:\.[A-Za-z_][A-Za-z0-9_]*)+", transform_expr):
        transform_expr = f"{transform_expr}(_)"
    out_cols = []
    for col in columns:
        current = str(col.get("name") or "")
        next_col = dict(col)
        next_col["name"] = str(_eval_m_expression([], [], transform_expr, current_value=current, variables=variables))
        out_cols.append(next_col)
    return out_cols, rows


def _add_index_column(columns: list[dict[str, Any]], rows: list[list[Any]], expression: str) -> tuple[list[dict[str, Any]], list[list[Any]]]:
    args = _function_args(expression)
    if len(args) < 2:
        raise ValueError("Table.AddIndexColumn preview needs an index column name.")
    name = str(_parse_m_value(args[1]))
    initial = int(_parse_m_value(args[2])) if len(args) > 2 else 0
    increment = int(_parse_m_value(args[3])) if len(args) > 3 else 1
    out_cols = [dict(col) for col in columns] + [{"name": name, "m_type": "type number"}]
    out_rows = [list(row) + [initial + idx * increment] for idx, row in enumerate(rows)]
    return out_cols, out_rows


def _promote_headers(columns: list[dict[str, Any]], rows: list[list[Any]]) -> tuple[list[dict[str, Any]], list[list[Any]]]:
    if not rows:
        return columns, rows
    first = rows[0]
    out_cols = [
        {"name": str(first[idx]) if idx < len(first) and first[idx] not in {None, ""} else str(col.get("name") or f"Column{idx + 1}")}
        for idx, col in enumerate(columns)
    ]
    return out_cols, [list(row) for row in rows[1:]]


def _demote_headers(columns: list[dict[str, Any]], rows: list[list[Any]]) -> tuple[list[dict[str, Any]], list[list[Any]]]:
    out_cols = [{"name": f"Column{idx + 1}"} for idx, _col in enumerate(columns)]
    header_row = [str(col.get("name") or f"Column{idx + 1}") for idx, col in enumerate(columns)]
    return out_cols, [header_row] + [list(row) for row in rows]


def _split_column(columns: list[dict[str, Any]], rows: list[list[Any]], expression: str) -> tuple[list[dict[str, Any]], list[list[Any]]]:
    args = _function_args(expression)
    if len(args) < 3:
        raise ValueError("Table.SplitColumn preview needs a source column and splitter.")
    source_name = str(_parse_m_value(args[1]))
    splitter = _eval_m_expression([], [], args[2])
    if not isinstance(splitter, MBuiltinFunctionValue):
        raise ValueError("Table.SplitColumn preview needs a supported Splitter.* helper.")
    output_names = (
        [str(item) for item in _parse_m_value(args[3])]
        if len(args) > 3 and isinstance(_parse_m_value(args[3]), list)
        else [f"{source_name}.1", f"{source_name}.2"]
    )
    source_idx = _find_column_index(columns, source_name)
    out_cols: list[dict[str, Any]] = []
    for idx, col in enumerate(columns):
        if idx == source_idx:
            out_cols.extend({"name": name} for name in output_names)
        else:
            out_cols.append(dict(col))
    out_rows: list[list[Any]] = []
    for row in rows:
        value = "" if source_idx >= len(row) or row[source_idx] is None else str(row[source_idx])
        parts = _apply_splitter(splitter, value)
        split_values = [parts[idx] if idx < len(parts) else None for idx, _name in enumerate(output_names)]
        out_row: list[Any] = []
        for idx, value in enumerate(row):
            if idx == source_idx:
                out_row.extend(split_values)
            else:
                out_row.append(value)
        out_rows.append(out_row)
    return out_cols, out_rows


def _combiner_text_delimiter(expression: str) -> str:
    match = re.search(r"Combiner\.CombineTextByDelimiter\s*\((.*)\)", expression, flags=re.IGNORECASE | re.DOTALL)
    if not match:
        raise ValueError("Table.CombineColumns preview supports Combiner.CombineTextByDelimiter only.")
    args = _split_top_level_items(match.group(1))
    if not args:
        return ""
    return str(_parse_m_value(args[0]))


def _combine_columns(columns: list[dict[str, Any]], rows: list[list[Any]], expression: str) -> tuple[list[dict[str, Any]], list[list[Any]]]:
    args = _function_args(expression)
    if len(args) < 4:
        raise ValueError("Table.CombineColumns preview needs source columns, combiner, and output column.")
    source_columns = _parse_m_value(args[1])
    if isinstance(source_columns, str):
        source_names = [source_columns]
    elif isinstance(source_columns, list):
        source_names = [str(name) for name in source_columns]
    else:
        raise ValueError("Table.CombineColumns preview needs a source-column list.")
    if not source_names:
        raise ValueError("Table.CombineColumns preview needs at least one source column.")
    combiner = _eval_m_expression([], [], args[2])
    if not isinstance(combiner, MBuiltinFunctionValue):
        delimiter = _combiner_text_delimiter(args[2])
        combiner = MBuiltinFunctionValue("Combiner.CombineTextByDelimiter", {"delimiter": delimiter})
    output_name = str(_parse_m_value(args[3]))
    source_indexes = [_find_column_index(columns, name) for name in source_names]
    source_index_set = set(source_indexes)
    insert_at = min(source_indexes)
    out_cols: list[dict[str, Any]] = []
    for idx, col in enumerate(columns):
        if idx == insert_at:
            out_cols.append({"name": output_name})
        if idx not in source_index_set:
            out_cols.append(dict(col))
    out_rows: list[list[Any]] = []
    for row in rows:
        combined = _apply_combiner(combiner, [row[idx] if idx < len(row) else None for idx in source_indexes])
        out_row: list[Any] = []
        for idx, value in enumerate(row):
            if idx == insert_at:
                out_row.append(combined)
            if idx not in source_index_set:
                out_row.append(value)
        if insert_at >= len(row):
            out_row.append(combined)
        out_rows.append(out_row)
    return out_cols, out_rows


def _numeric_values(values: Sequence[Any]) -> list[float]:
    out: list[float] = []
    for value in values:
        if value is None or value == "":
            continue
        try:
            out.append(float(value))
        except Exception:
            continue
    return out


def _value_key(value: Any) -> Any:
    value = _unwrap_metadata(value)
    if isinstance(value, Mapping):
        return ("record", tuple(sorted((str(key), _value_key(item)) for key, item in value.items())))
    if isinstance(value, list):
        return ("list", tuple(_value_key(item) for item in value))
    if isinstance(value, (_dt.datetime, _dt.date, _dt.time)):
        return (type(value).__name__, value.isoformat())
    if isinstance(value, _dt.timedelta):
        return ("duration", value.total_seconds())
    try:
        hash(value)
        return ("scalar", value)
    except Exception:
        return ("repr", repr(value))


def _list_median(values: Sequence[Any]) -> Any:
    non_null = [value for value in values if value is not None]
    if not non_null:
        return None
    numeric = _numeric_values(non_null)
    if len(numeric) == len(non_null):
        ordered = sorted(numeric)
        midpoint = len(ordered) // 2
        if len(ordered) % 2:
            return ordered[midpoint]
        return (ordered[midpoint - 1] + ordered[midpoint]) / 2
    try:
        ordered_values = sorted(non_null)
    except Exception:
        ordered_values = sorted(non_null, key=lambda item: str(_value_key(item)))
    midpoint = len(ordered_values) // 2
    return ordered_values[midpoint] if len(ordered_values) % 2 else ordered_values[midpoint - 1]


def _list_mode(values: Sequence[Any], *, return_all: bool) -> Any:
    if not values:
        raise ValueError("List.Mode preview needs a non-empty list.")
    counts: OrderedDict[Any, tuple[Any, int]] = OrderedDict()
    for item in values:
        key = _value_key(item)
        original, count = counts.get(key, (item, 0))
        counts[key] = (original, count + 1)
    max_count = max(count for _item, count in counts.values())
    modes = [item for item, count in counts.values() if count == max_count]
    return modes if return_all else modes[-1]


def _evaluate_group_aggregate(expression: str, group_columns: list[dict[str, Any]], group_rows: list[list[Any]]) -> Any:
    text = expression.strip()
    if text.lower().startswith("each "):
        text = text[5:].strip()
    if text == "_":
        return [_row_record(group_columns, row) for row in group_rows]
    if re.search(r"Table\.RowCount\(\s*_\s*\)", text, flags=re.IGNORECASE):
        return len(group_rows)
    list_match = re.search(r"List\.(Sum|Min|Max|Average|Count|NonNullCount)\(\s*\[([^\]]+)\]\s*\)", text, flags=re.IGNORECASE)
    if not list_match:
        raise ValueError(f"Unsupported Table.Group aggregate expression: {expression}")
    func = list_match.group(1).lower()
    idx = _find_column_index(group_columns, list_match.group(2))
    values = [row[idx] if idx < len(row) else None for row in group_rows]
    if func == "count":
        return len(values)
    if func == "nonnullcount":
        return sum(1 for value in values if value is not None and value != "")
    numeric = _numeric_values(values)
    if func == "sum":
        return sum(numeric)
    if func == "min":
        return min(numeric) if numeric else None
    if func == "max":
        return max(numeric) if numeric else None
    if func == "average":
        return sum(numeric) / len(numeric) if numeric else None
    return None


def _table_group_specs(expression: str) -> tuple[list[str], list[tuple[str, str]]]:
    args = _function_args(expression)
    if len(args) < 3:
        raise ValueError("Table.Group preview needs group columns and aggregate specs.")
    raw_groups = _eval_m_expression([], [], args[1])
    group_names = [str(item) for item in raw_groups] if isinstance(raw_groups, list) else []
    spec_body = _between_outer(args[2], "{", "}")
    if not spec_body:
        raise ValueError("Table.Group preview needs aggregate specs.")
    aggregate_specs: list[tuple[str, str]] = []
    for raw_spec in _split_top_level_items(spec_body):
        body = _between_outer(raw_spec, "{", "}")
        if body is None:
            continue
        parts = _split_top_level_items(body)
        if len(parts) < 2:
            continue
        name = _eval_m_expression([], [], parts[0])
        aggregate_specs.append((str(name), parts[1].strip()))
    return group_names, aggregate_specs


def _evaluate_list_aggregate(function_name: str, values: Sequence[Any]) -> Any:
    func = str(function_name or "").strip().upper()
    if func.startswith("EACH "):
        func = func[5:].strip().upper()
    if func.endswith("(_)"):
        func = func[:-3]
    if func == "LIST.COUNT":
        return len(values)
    if func == "LIST.NONNULLCOUNT":
        return sum(1 for value in values if value is not None and value != "")
    numeric = _numeric_values(values)
    if func == "LIST.SUM":
        return sum(numeric) if numeric else None
    if func == "LIST.MIN":
        return min(numeric) if numeric else None
    if func == "LIST.MAX":
        return max(numeric) if numeric else None
    if func == "LIST.AVERAGE":
        return sum(numeric) / len(numeric) if numeric else None
    raise ValueError(f"Unsupported Table.AggregateTableColumn aggregate function: {function_name}")


def _aggregate_table_column(
    columns: list[dict[str, Any]],
    rows: list[list[Any]],
    expression: str,
) -> tuple[list[dict[str, Any]], list[list[Any]]]:
    args = _function_args(expression)
    if len(args) < 3:
        raise ValueError("Table.AggregateTableColumn preview needs a nested column and aggregate specs.")
    nested_name = str(_parse_m_value(args[1]))
    nested_idx = _find_column_index(columns, nested_name)
    raw_specs = _parse_m_value(args[2])
    if not isinstance(raw_specs, list) or not raw_specs:
        raise ValueError("Table.AggregateTableColumn preview needs at least one aggregate spec.")
    specs: list[tuple[str, str, str]] = []
    for raw_spec in raw_specs:
        if not isinstance(raw_spec, list) or len(raw_spec) < 2:
            raise ValueError("Table.AggregateTableColumn preview only supports list aggregate specs.")
        source_field = str(raw_spec[0])
        function_name = str(raw_spec[1])
        output_name = str(raw_spec[2]) if len(raw_spec) >= 3 else source_field
        specs.append((source_field, function_name, output_name))
    keep_indexes = [idx for idx, _col in enumerate(columns) if idx != nested_idx]
    out_cols = [dict(columns[idx]) for idx in keep_indexes] + [{"name": output_name} for _src, _fn, output_name in specs]
    out_rows: list[list[Any]] = []
    for row in rows:
        prefix = [row[idx] if idx < len(row) else None for idx in keep_indexes]
        nested_value = row[nested_idx] if nested_idx < len(row) else None
        nested_records = nested_value if isinstance(nested_value, list) else []
        aggregates: list[Any] = []
        for source_field, function_name, _output_name in specs:
            values = [
                record.get(source_field)
                for record in nested_records
                if isinstance(record, Mapping)
            ]
            aggregates.append(_evaluate_list_aggregate(function_name, values))
        out_rows.append(prefix + aggregates)
    return out_cols, out_rows


def _group_table(columns: list[dict[str, Any]], rows: list[list[Any]], expression: str) -> tuple[list[dict[str, Any]], list[list[Any]]]:
    group_names, aggregate_specs = _table_group_specs(expression)
    if not group_names:
        raise ValueError("Table.Group preview could not infer group columns.")
    group_indexes = [_find_column_index(columns, name) for name in group_names]
    groups: "OrderedDict[tuple[Any, ...], list[list[Any]]]" = OrderedDict()
    for row in rows:
        key = tuple(row[idx] if idx < len(row) else None for idx in group_indexes)
        groups.setdefault(key, []).append(row)
    out_cols = [{"name": name} for name in group_names] + [{"name": name} for name, _expr in aggregate_specs]
    out_rows: list[list[Any]] = []
    for key, group_rows in groups.items():
        out_row = list(key)
        for _name, aggregate_expression in aggregate_specs:
            out_row.append(_evaluate_group_aggregate(aggregate_expression, columns, group_rows))
        out_rows.append(out_row)
    return out_cols, out_rows


def _combine_tables(tables: Sequence[tuple[list[dict[str, Any]], list[list[Any]]]]) -> tuple[list[dict[str, Any]], list[list[Any]]]:
    names: list[str] = []
    seen: set[str] = set()
    for cols, _rows in tables:
        for col in cols:
            name = str(col.get("name") or "")
            key = name.upper()
            if key and key not in seen:
                seen.add(key)
                names.append(name)
    out_cols = [{"name": name} for name in names]
    out_rows: list[list[Any]] = []
    for cols, rows in tables:
        for row in rows:
            record = _row_record(cols, row)
            out_rows.append([record.get(name) for name in names])
    return out_cols, out_rows


def _unpivot_other_columns(columns: list[dict[str, Any]], rows: list[list[Any]], expression: str) -> tuple[list[dict[str, Any]], list[list[Any]]]:
    literals = _string_literals_after_first_arg(expression)
    if len(literals) < 3:
        raise ValueError("Table.UnpivotOtherColumns preview needs preserved columns plus attribute/value names.")
    attribute_name = literals[-2]
    value_name = literals[-1]
    preserve_names = literals[:-2]
    preserve_indexes = [_find_column_index(columns, name) for name in preserve_names]
    preserve_keys = {name.upper() for name in preserve_names}
    unpivot_indexes = [
        idx for idx, col in enumerate(columns)
        if str(col.get("name") or "").upper() not in preserve_keys
    ]
    out_cols = [{"name": name} for name in preserve_names] + [{"name": attribute_name}, {"name": value_name}]
    out_rows: list[list[Any]] = []
    for row in rows:
        prefix = [row[idx] if idx < len(row) else None for idx in preserve_indexes]
        for idx in unpivot_indexes:
            out_rows.append(prefix + [str(columns[idx].get("name") or ""), row[idx] if idx < len(row) else None])
    return out_cols, out_rows


def _unpivot_columns(columns: list[dict[str, Any]], rows: list[list[Any]], expression: str) -> tuple[list[dict[str, Any]], list[list[Any]]]:
    literals = _string_literals_after_first_arg(expression)
    if len(literals) < 3:
        raise ValueError("Table.Unpivot preview needs target columns plus attribute/value names.")
    attribute_name = literals[-2]
    value_name = literals[-1]
    target_names = literals[:-2]
    target_keys = {name.upper() for name in target_names}
    target_indexes = [_find_column_index(columns, name) for name in target_names]
    preserve_indexes = [
        idx for idx, col in enumerate(columns)
        if str(col.get("name") or "").upper() not in target_keys
    ]
    preserve_names = [str(columns[idx].get("name") or "") for idx in preserve_indexes]
    out_cols = [{"name": name} for name in preserve_names] + [{"name": attribute_name}, {"name": value_name}]
    out_rows: list[list[Any]] = []
    for row in rows:
        prefix = [row[idx] if idx < len(row) else None for idx in preserve_indexes]
        for idx in target_indexes:
            out_rows.append(prefix + [str(columns[idx].get("name") or ""), row[idx] if idx < len(row) else None])
    return out_cols, out_rows


def _pivot_table(columns: list[dict[str, Any]], rows: list[list[Any]], expression: str) -> tuple[list[dict[str, Any]], list[list[Any]]]:
    literals = _string_literals_after_first_arg(expression)
    if len(literals) < 2:
        raise ValueError("Table.Pivot preview needs attribute and value column names.")
    attribute_name = literals[-2]
    value_name = literals[-1]
    requested_values = literals[:-2]
    attr_idx = _find_column_index(columns, attribute_name)
    value_idx = _find_column_index(columns, value_name)
    key_indexes = [
        idx for idx, col in enumerate(columns)
        if idx not in {attr_idx, value_idx}
    ]
    key_names = [str(columns[idx].get("name") or "") for idx in key_indexes]
    pivot_values = requested_values or []
    if not pivot_values:
        seen: set[Any] = set()
        for row in rows:
            value = row[attr_idx] if attr_idx < len(row) else None
            if value not in seen:
                seen.add(value)
                pivot_values.append(str(value))
    groups: "OrderedDict[tuple[Any, ...], dict[str, list[Any]]]" = OrderedDict()
    for row in rows:
        key = tuple(row[idx] if idx < len(row) else None for idx in key_indexes)
        bucket = groups.setdefault(key, {name: [] for name in pivot_values})
        attr = str(row[attr_idx] if attr_idx < len(row) else "")
        if attr not in bucket:
            bucket[attr] = []
            pivot_values.append(attr)
        bucket[attr].append(row[value_idx] if value_idx < len(row) else None)
    use_sum = "LIST.SUM" in expression.upper()
    out_cols = [{"name": name} for name in key_names + pivot_values]
    out_rows: list[list[Any]] = []
    for key, bucket in groups.items():
        out_row = list(key)
        for name in pivot_values:
            values = bucket.get(name) or []
            numeric = _numeric_values(values)
            if use_sum:
                out_row.append(sum(numeric) if numeric else None)
            else:
                out_row.append(values[-1] if values else None)
        out_rows.append(out_row)
    return out_cols, out_rows


def _unique_output_name(existing: set[str], desired: str) -> str:
    base = desired or "Column"
    candidate = base
    idx = 1
    while candidate.upper() in existing:
        candidate = f"{base}.{idx}"
        idx += 1
    existing.add(candidate.upper())
    return candidate


def _first_two_tables(
    dependencies: Sequence[str],
    env: Mapping[str, tuple[list[dict[str, Any]], list[list[Any]]]],
) -> tuple[tuple[list[dict[str, Any]], list[list[Any]]], tuple[list[dict[str, Any]], list[list[Any]]]]:
    tables = [env[dep] for dep in dependencies if dep in env]
    if len(tables) < 2:
        raise ValueError("Join preview needs two table dependencies.")
    return tables[0], tables[1]


def _env_table_by_name(
    env: Mapping[str, tuple[list[dict[str, Any]], list[list[Any]]]],
    raw_name: str,
) -> tuple[list[dict[str, Any]], list[list[Any]]] | None:
    name = _m_identifier_name(raw_name.strip())
    if name in env:
        return env[name]
    wanted = name.upper()
    for existing_name, table in env.items():
        if str(existing_name).upper() == wanted:
            return table
    return None


def _join_tables_from_expression(
    dependencies: Sequence[str],
    env: Mapping[str, tuple[list[dict[str, Any]], list[list[Any]]]],
    expression: str,
) -> tuple[tuple[list[dict[str, Any]], list[list[Any]]], tuple[list[dict[str, Any]], list[list[Any]]]]:
    args = _function_args(expression)
    ordered: list[tuple[list[dict[str, Any]], list[list[Any]]]] = []
    if len(args) >= 4:
        for arg_idx in (0, 2):
            table = _env_table_by_name(env, args[arg_idx])
            if table is not None:
                ordered.append(table)
    if len(ordered) == 2:
        return ordered[0], ordered[1]
    return _first_two_tables(dependencies, env)


def _join_kind(expression: str) -> str:
    return _join_kind_value(_function_args(expression)[4] if len(_function_args(expression)) > 4 else None, default="left")


def _join_kind_value(value: Any, *, default: str) -> str:
    text = str(value or "").strip().strip('"').strip("'").upper()
    if not text:
        return default
    if text in {"JOINKIND.INNER", "INNER", "0"}:
        return "inner"
    if text in {"JOINKIND.LEFTOUTER", "LEFTOUTER", "LEFT", "1"}:
        return "left"
    if text in {"JOINKIND.RIGHTOUTER", "RIGHTOUTER", "RIGHT", "2"}:
        return "right"
    if text in {"JOINKIND.FULLOUTER", "FULLOUTER", "FULL", "3"}:
        return "full"
    if text in {"JOINKIND.LEFTANTI", "LEFTANTI", "LEFT_ANTI", "4"}:
        return "left_anti"
    if text in {"JOINKIND.RIGHTANTI", "RIGHTANTI", "RIGHT_ANTI", "5"}:
        return "right_anti"
    raise ValueError(f"Unsupported M join kind: {value}")


def _join_key_names(value: Any) -> list[str]:
    parsed = _parse_m_value(str(value).strip()) if isinstance(value, str) else value
    if isinstance(parsed, list):
        names = [str(item) for item in parsed if str(item)]
    elif parsed is not None:
        names = [str(parsed)]
    else:
        names = []
    if not names:
        raise ValueError("Join preview needs at least one key column.")
    return names


def _fuzzy_option(options: Mapping[str, Any], name: str, default: Any = None) -> Any:
    try:
        return _mapping_get_case_insensitive(options, name)
    except KeyError:
        return default


def _fuzzy_transformations(options: Mapping[str, Any]) -> dict[str, str]:
    table = _fuzzy_option(options, "TransformationTable", [])
    if not isinstance(table, list):
        return {}
    out: dict[str, str] = {}
    for record in table:
        if not isinstance(record, Mapping):
            continue
        try:
            source = str(_mapping_get_case_insensitive(record, "From"))
            target = str(_mapping_get_case_insensitive(record, "To"))
        except KeyError:
            continue
        out[source.upper()] = target
    return out


def _fuzzy_normalize(value: Any, options: Mapping[str, Any]) -> str:
    text = "" if value is None else str(value)
    transforms = _fuzzy_transformations(options)
    if transforms:
        text = transforms.get(text.upper(), text)
    if bool(_fuzzy_option(options, "IgnoreSpace", False)):
        text = re.sub(r"\s+", "", text)
    if bool(_fuzzy_option(options, "IgnoreCase", True)):
        text = text.casefold()
    return text


def _fuzzy_similarity(left: Any, right: Any, options: Mapping[str, Any]) -> float:
    lval = _fuzzy_normalize(left, options)
    rval = _fuzzy_normalize(right, options)
    if lval == rval:
        return 1.0
    if not lval or not rval:
        return 0.0
    return _difflib.SequenceMatcher(None, lval, rval).ratio()


def _fuzzy_row_similarity(
    left_cols: Sequence[Mapping[str, Any]],
    left_row: Sequence[Any],
    right_cols: Sequence[Mapping[str, Any]],
    right_row: Sequence[Any],
    left_keys: Sequence[str],
    right_keys: Sequence[str],
    options: Mapping[str, Any],
) -> float:
    scores: list[float] = []
    for left_key, right_key in zip(left_keys, right_keys):
        left_idx = _find_column_index(left_cols, left_key)
        right_idx = _find_column_index(right_cols, right_key)
        scores.append(_fuzzy_similarity(
            left_row[left_idx] if left_idx < len(left_row) else None,
            right_row[right_idx] if right_idx < len(right_row) else None,
            options,
        ))
    return min(scores) if scores else 0.0


def _row_key(row: Sequence[Any], indexes: Sequence[int]) -> tuple[Any, ...]:
    return tuple(row[idx] if idx < len(row) else None for idx in indexes)


def _join_args(expression: str, *, nested: bool) -> tuple[list[str], list[str], str | None, str]:
    args = _function_args(expression)
    if len(args) < (5 if nested else 4):
        raise ValueError("Join preview needs left/right tables and key columns.")
    left_keys = _join_key_names(args[1])
    right_keys = _join_key_names(args[3])
    if len(left_keys) != len(right_keys):
        raise ValueError("Join preview needs the same number of left and right key columns.")
    if nested:
        nested_name = str(_parse_m_value(args[4]))
        return left_keys, right_keys, nested_name, _join_kind_value(args[5] if len(args) > 5 else None, default="left")
    return left_keys, right_keys, None, _join_kind_value(args[4] if len(args) > 4 else None, default="inner")


def _join_tables(
    left: tuple[list[dict[str, Any]], list[list[Any]]],
    right: tuple[list[dict[str, Any]], list[list[Any]]],
    expression: str,
) -> tuple[list[dict[str, Any]], list[list[Any]]]:
    left_cols, left_rows = left
    right_cols, right_rows = right
    left_keys, right_keys, _nested_name, kind = _join_args(expression, nested=False)
    left_indexes = [_find_column_index(left_cols, key) for key in left_keys]
    right_indexes = [_find_column_index(right_cols, key) for key in right_keys]

    if kind == "left_anti":
        right_lookup = {_row_key(row, right_indexes) for row in right_rows}
        return [dict(col) for col in left_cols], [
            list(row) for row in left_rows
            if _row_key(row, left_indexes) not in right_lookup
        ]
    if kind == "right_anti":
        left_lookup = {_row_key(row, left_indexes) for row in left_rows}
        return [dict(col) for col in right_cols], [
            list(row) for row in right_rows
            if _row_key(row, right_indexes) not in left_lookup
        ]

    existing = {str(col.get("name") or "").upper() for col in left_cols}
    right_key_index_set = set(right_indexes)
    right_output_indexes: list[int] = []
    right_output_names: list[str] = []
    for idx, col in enumerate(right_cols):
        if idx in right_key_index_set:
            continue
        right_output_indexes.append(idx)
        right_output_names.append(_unique_output_name(existing, str(col.get("name") or "")))

    lookup: dict[tuple[Any, ...], list[tuple[int, list[Any]]]] = {}
    for right_row_idx, row in enumerate(right_rows):
        lookup.setdefault(_row_key(row, right_indexes), []).append((right_row_idx, row))

    out_cols = [dict(col) for col in left_cols] + [{"name": name} for name in right_output_names]
    out_rows: list[list[Any]] = []
    matched_right_indexes: set[int] = set()
    for left_row in left_rows:
        key = _row_key(left_row, left_indexes)
        matches = lookup.get(key) or []
        if not matches and kind in {"inner", "right"}:
            continue
        if not matches:
            out_rows.append(list(left_row) + [None for _idx in right_output_indexes])
            continue
        for right_row_idx, right_row in matches:
            matched_right_indexes.add(right_row_idx)
            out_rows.append(list(left_row) + [right_row[idx] if idx < len(right_row) else None for idx in right_output_indexes])
    if kind in {"right", "full"}:
        for right_row_idx, right_row in enumerate(right_rows):
            if right_row_idx in matched_right_indexes:
                continue
            out_rows.append([None for _col in left_cols] + [right_row[idx] if idx < len(right_row) else None for idx in right_output_indexes])
    return out_cols, out_rows


def _nested_join_tables(
    left: tuple[list[dict[str, Any]], list[list[Any]]],
    right: tuple[list[dict[str, Any]], list[list[Any]]],
    expression: str,
) -> tuple[list[dict[str, Any]], list[list[Any]]]:
    left_cols, left_rows = left
    right_cols, right_rows = right
    left_keys, right_keys, nested_name, kind = _join_args(expression, nested=True)
    left_indexes = [_find_column_index(left_cols, key) for key in left_keys]
    right_indexes = [_find_column_index(right_cols, key) for key in right_keys]
    lookup: dict[tuple[Any, ...], list[tuple[int, dict[str, Any]]]] = {}
    for right_row_idx, row in enumerate(right_rows):
        lookup.setdefault(_row_key(row, right_indexes), []).append((right_row_idx, _row_record(right_cols, row)))
    out_cols = [dict(col) for col in left_cols] + [{"name": nested_name, "m_type": "table"}]
    out_rows: list[list[Any]] = []
    matched_right_indexes: set[int] = set()
    for left_row in left_rows:
        key = _row_key(left_row, left_indexes)
        matches = lookup.get(key) or []
        if not matches and kind in {"inner", "right"}:
            continue
        if not matches:
            if kind in {"left", "full", "left_anti"}:
                out_rows.append(list(left_row) + [[]])
            continue
        for right_row_idx, _record in matches:
            matched_right_indexes.add(right_row_idx)
        if kind not in {"left_anti", "right_anti"}:
            out_rows.append(list(left_row) + [[record for _idx, record in matches]])
    if kind in {"right", "full", "right_anti"}:
        for right_row_idx, right_row in enumerate(right_rows):
            if right_row_idx in matched_right_indexes:
                continue
            out_rows.append([None for _col in left_cols] + [[_row_record(right_cols, right_row)]])
    return out_cols, out_rows


def _fuzzy_join_args(expression: str, *, nested: bool) -> tuple[list[str], list[str], str | None, str, dict[str, Any]]:
    args = _function_args(expression)
    if len(args) < (5 if nested else 4):
        raise ValueError("Fuzzy join preview needs left/right tables and key columns.")
    left_keys = _join_key_names(args[1])
    right_keys = _join_key_names(args[3])
    if len(left_keys) != len(right_keys):
        raise ValueError("Fuzzy join preview needs the same number of left and right key columns.")
    options_idx = 6 if nested else 5
    options = _eval_m_expression([], [], args[options_idx]) if len(args) > options_idx else {}
    options_record = dict(options) if isinstance(options, Mapping) else {}
    if nested:
        nested_name = str(_parse_m_value(args[4]))
        return left_keys, right_keys, nested_name, _join_kind_value(args[5] if len(args) > 5 else None, default="left"), options_record
    return left_keys, right_keys, None, _join_kind_value(args[4] if len(args) > 4 else None, default="inner"), options_record


def _fuzzy_matches_for_row(
    left_cols: Sequence[Mapping[str, Any]],
    left_row: Sequence[Any],
    right_cols: Sequence[Mapping[str, Any]],
    right_rows: Sequence[Sequence[Any]],
    left_keys: Sequence[str],
    right_keys: Sequence[str],
    options: Mapping[str, Any],
) -> list[tuple[int, list[Any], float]]:
    threshold = float(_to_number(_fuzzy_option(options, "Threshold", 0.8)) or 0.8)
    number_of_matches = _fuzzy_option(options, "NumberOfMatches", None)
    limit = int(_to_number(number_of_matches) or 0) if number_of_matches is not None else 0
    matches: list[tuple[int, list[Any], float]] = []
    for right_idx, right_row in enumerate(right_rows):
        score = _fuzzy_row_similarity(left_cols, left_row, right_cols, right_row, left_keys, right_keys, options)
        if score >= threshold:
            matches.append((right_idx, list(right_row), score))
    matches.sort(key=lambda item: item[2], reverse=True)
    return matches[:limit] if limit > 0 else matches


def _fuzzy_join_tables(
    left: tuple[list[dict[str, Any]], list[list[Any]]],
    right: tuple[list[dict[str, Any]], list[list[Any]]],
    expression: str,
) -> tuple[list[dict[str, Any]], list[list[Any]]]:
    left_cols, left_rows = left
    right_cols, right_rows = right
    left_keys, right_keys, _nested_name, kind, options = _fuzzy_join_args(expression, nested=False)
    existing = {str(col.get("name") or "").upper() for col in left_cols}
    right_output_indexes = [idx for idx, _col in enumerate(right_cols)]
    right_output_names = [_unique_output_name(existing, str(col.get("name") or "")) for col in right_cols]
    similarity_name = _fuzzy_option(options, "SimilarityColumnName", None)
    out_cols = [dict(col) for col in left_cols] + [{"name": name} for name in right_output_names]
    if similarity_name:
        out_cols.append({"name": str(similarity_name)})
    out_rows: list[list[Any]] = []
    matched_right_indexes: set[int] = set()
    for left_row in left_rows:
        matches = _fuzzy_matches_for_row(left_cols, left_row, right_cols, right_rows, left_keys, right_keys, options)
        if not matches and kind in {"inner", "right"}:
            continue
        if not matches:
            if kind in {"left", "full"}:
                out_rows.append(list(left_row) + [None for _idx in right_output_indexes] + ([None] if similarity_name else []))
            elif kind == "left_anti":
                out_rows.append(list(left_row) + [None for _idx in right_output_indexes] + ([None] if similarity_name else []))
            continue
        if kind == "left_anti":
            continue
        for right_idx, right_row, score in matches:
            matched_right_indexes.add(right_idx)
            out_rows.append(
                list(left_row)
                + [right_row[idx] if idx < len(right_row) else None for idx in right_output_indexes]
                + ([score] if similarity_name else [])
            )
    if kind in {"right", "full", "right_anti"}:
        for right_idx, right_row in enumerate(right_rows):
            if right_idx in matched_right_indexes:
                continue
            out_rows.append(
                [None for _col in left_cols]
                + [right_row[idx] if idx < len(right_row) else None for idx in right_output_indexes]
                + ([None] if similarity_name else [])
            )
    return out_cols, out_rows


def _fuzzy_nested_join_tables(
    left: tuple[list[dict[str, Any]], list[list[Any]]],
    right: tuple[list[dict[str, Any]], list[list[Any]]],
    expression: str,
) -> tuple[list[dict[str, Any]], list[list[Any]]]:
    left_cols, left_rows = left
    right_cols, right_rows = right
    left_keys, right_keys, nested_name, kind, options = _fuzzy_join_args(expression, nested=True)
    similarity_name = _fuzzy_option(options, "SimilarityColumnName", None)
    out_cols = [dict(col) for col in left_cols] + [{"name": nested_name, "m_type": "table"}]
    out_rows: list[list[Any]] = []
    matched_right_indexes: set[int] = set()
    for left_row in left_rows:
        matches = _fuzzy_matches_for_row(left_cols, left_row, right_cols, right_rows, left_keys, right_keys, options)
        if not matches and kind in {"inner", "right"}:
            continue
        if not matches:
            if kind in {"left", "full", "left_anti"}:
                out_rows.append(list(left_row) + [[]])
            continue
        for right_idx, _right_row, _score in matches:
            matched_right_indexes.add(right_idx)
        if kind not in {"left_anti", "right_anti"}:
            nested_records: list[dict[str, Any]] = []
            for _right_idx, right_row, score in matches:
                record = _row_record(right_cols, right_row)
                if similarity_name:
                    record[str(similarity_name)] = score
                nested_records.append(record)
            out_rows.append(list(left_row) + [nested_records])
    if kind in {"right", "full", "right_anti"}:
        for right_idx, right_row in enumerate(right_rows):
            if right_idx in matched_right_indexes:
                continue
            record = _row_record(right_cols, right_row)
            if similarity_name:
                record[str(similarity_name)] = None
            out_rows.append([None for _col in left_cols] + [[record]])
    return out_cols, out_rows


def _fuzzy_cluster_value(values: Sequence[Any], value: Any, options: Mapping[str, Any]) -> Any:
    threshold = float(_to_number(_fuzzy_option(options, "Threshold", 0.8)) or 0.8)
    for candidate in values:
        if _fuzzy_similarity(candidate, value, options) >= threshold:
            return candidate
    return value


def _add_fuzzy_cluster_column(columns: list[dict[str, Any]], rows: list[list[Any]], expression: str) -> tuple[list[dict[str, Any]], list[list[Any]]]:
    args = _function_args(expression)
    if len(args) < 3:
        raise ValueError("Table.AddFuzzyClusterColumn preview needs source column and output column.")
    source_name = str(_parse_m_value(args[1]))
    output_name = str(_parse_m_value(args[2]))
    options_value = _eval_m_expression([], [], args[3]) if len(args) > 3 else {}
    options = dict(options_value) if isinstance(options_value, Mapping) else {}
    source_idx = _find_column_index(columns, source_name)
    clusters: list[Any] = []
    out_rows: list[list[Any]] = []
    for row in rows:
        value = row[source_idx] if source_idx < len(row) else None
        cluster = _fuzzy_cluster_value(clusters, value, options)
        if cluster == value and all(_fuzzy_similarity(existing, value, options) < float(_to_number(_fuzzy_option(options, "Threshold", 0.8)) or 0.8) for existing in clusters):
            clusters.append(value)
        out_rows.append(list(row) + [cluster])
    return [dict(col) for col in columns] + [{"name": output_name}], out_rows


def _fuzzy_group_table(columns: list[dict[str, Any]], rows: list[list[Any]], expression: str) -> tuple[list[dict[str, Any]], list[list[Any]]]:
    args = _function_args(expression)
    if len(args) < 3:
        raise ValueError("Table.FuzzyGroup preview needs key columns and aggregate specs.")
    key_names = _join_key_names(_eval_m_expression([], [], args[1]))
    options_value = _eval_m_expression([], [], args[3]) if len(args) > 3 else {}
    options = dict(options_value) if isinstance(options_value, Mapping) else {}
    key_indexes = [_find_column_index(columns, name) for name in key_names]
    groups: list[tuple[list[Any], list[list[Any]]]] = []
    for row in rows:
        key = [row[idx] if idx < len(row) else None for idx in key_indexes]
        target_group: list[list[Any]] | None = None
        for existing_key, group_rows in groups:
            scores = [_fuzzy_similarity(existing, current, options) for existing, current in zip(existing_key, key)]
            if scores and min(scores) >= float(_to_number(_fuzzy_option(options, "Threshold", 0.8)) or 0.8):
                target_group = group_rows
                break
        if target_group is None:
            target_group = []
            groups.append((key, target_group))
        target_group.append(list(row))
    regex_specs = re.findall(
        r'\{\s*"([^"]+)"\s*,\s*each\s+(.+?)(?:,\s*(?:type\s+[A-Za-z0-9_ ]+|[A-Za-z0-9_.]+))?\s*\}',
        args[2],
        flags=re.IGNORECASE | re.DOTALL,
    )
    out_cols = [{"name": name} for name in key_names]
    out_rows: list[list[Any]] = []
    if regex_specs:
        out_cols += [{"name": name} for name, _expr in regex_specs]
        for key, group_rows in groups:
            out_row = list(key)
            for _name, aggregate_expression in regex_specs:
                out_row.append(_evaluate_group_aggregate(aggregate_expression, columns, group_rows))
            out_rows.append(out_row)
        return out_cols, out_rows

    aggregate_specs = _eval_m_expression([], [], args[2])
    specs: list[tuple[str, Any, str]] = []
    if isinstance(aggregate_specs, list):
        for item in aggregate_specs:
            if isinstance(item, list) and len(item) >= 2:
                specs.append((str(item[0]), item[1], str(item[2]) if len(item) > 2 else str(item[0])))
    out_cols += [{"name": output_name} for _source, _func, output_name in specs]
    for key, group_rows in groups:
        records = _table_records(columns, group_rows)
        out_row = list(key)
        for _source, aggregate_fn, _output in specs:
            if isinstance(aggregate_fn, (MFunctionValue, MBuiltinFunctionValue)):
                out_row.append(_invoke_m_function(aggregate_fn, [records]))
            else:
                text = str(aggregate_fn)
                if text.upper() == "TABLE.ROWCOUNT":
                    out_row.append(len(records))
                else:
                    out_row.append(_eval_m_expression(columns, group_rows[0] if group_rows else [], text, current_value=records))
        out_rows.append(out_row)
    return out_cols, out_rows


def _filter_with_data_table(columns: list[dict[str, Any]], rows: list[list[Any]], expression: str) -> tuple[list[dict[str, Any]], list[list[Any]]]:
    args = _function_args(expression)
    if len(args) < 2:
        raise ValueError("Table.FilterWithDataTable preview needs a source table and data table.")
    data_value = _eval_m_expression([], [], args[1])
    filter_records = [dict(record) for record in data_value if isinstance(record, Mapping)] if isinstance(data_value, list) else []
    out_rows: list[list[Any]] = []
    for row in rows:
        record = _row_record(columns, row)
        keep = False
        for filter_record in filter_records:
            if {"Column", "Value"}.issubset({str(key) for key in filter_record.keys()}):
                field = str(_mapping_get_case_insensitive(filter_record, "Column"))
                wanted = _mapping_get_case_insensitive(filter_record, "Value")
                keep = _mapping_get_case_insensitive(record, field) == wanted
            else:
                common = [key for key in filter_record if any(str(col.get("name") or "").upper() == str(key).upper() for col in columns)]
                keep = bool(common) and all(_mapping_get_case_insensitive(record, str(key)) == filter_record[key] for key in common)
            if keep:
                break
        if keep:
            out_rows.append(list(row))
    return [dict(col) for col in columns], out_rows


def _expand_table_column(
    columns: list[dict[str, Any]],
    rows: list[list[Any]],
    expression: str,
    *,
    variables: Mapping[str, Any] | None = None,
) -> tuple[list[dict[str, Any]], list[list[Any]]]:
    args = _function_args(expression)
    if len(args) < 3:
        raise ValueError("Table.ExpandTableColumn preview needs a nested column and at least one field.")
    nested_name = str(_parse_m_value(args[1]))
    nested_idx = _find_column_index(columns, nested_name)
    field_value = _eval_m_expression([], [], args[2], variables=variables)
    field_names = [str(item) for item in field_value] if isinstance(field_value, list) else _string_literals(args[2])
    if not field_names:
        raise ValueError("Table.ExpandTableColumn preview needs at least one field.")
    if len(args) >= 4:
        output_value = _eval_m_expression([], [], args[3], variables=variables)
        output_names = [str(item) for item in output_value] if isinstance(output_value, list) else _string_literals(args[3])
    else:
        output_names = field_names
    keep_indexes = [idx for idx, _col in enumerate(columns) if idx != nested_idx]
    out_cols = [dict(columns[idx]) for idx in keep_indexes] + [{"name": name} for name in output_names]
    out_rows: list[list[Any]] = []
    for row in rows:
        prefix = [row[idx] if idx < len(row) else None for idx in keep_indexes]
        nested_value = row[nested_idx] if nested_idx < len(row) else None
        nested_records = nested_value if isinstance(nested_value, list) else []
        if not nested_records:
            out_rows.append(prefix + [None for _name in output_names])
            continue
        for record in nested_records:
            if not isinstance(record, Mapping):
                continue
            out_rows.append(prefix + [record.get(field) for field in field_names])
    return out_cols, out_rows


def _expand_record_column(columns: list[dict[str, Any]], rows: list[list[Any]], expression: str) -> tuple[list[dict[str, Any]], list[list[Any]]]:
    names = _string_literals_after_first_arg(expression)
    if len(names) < 2:
        raise ValueError("Table.ExpandRecordColumn preview needs a record column and at least one field.")
    record_name = names[0]
    record_idx = _find_column_index(columns, record_name)
    remaining = names[1:]
    half = len(remaining) // 2
    if half and len(remaining) % 2 == 0:
        field_names = remaining[:half]
        output_names = remaining[half:]
    else:
        field_names = remaining
        output_names = remaining
    keep_indexes = [idx for idx, _col in enumerate(columns) if idx != record_idx]
    out_cols = [dict(columns[idx]) for idx in keep_indexes] + [{"name": name} for name in output_names]
    out_rows = []
    for row in rows:
        prefix = [row[idx] if idx < len(row) else None for idx in keep_indexes]
        record = row[record_idx] if record_idx < len(row) and isinstance(row[record_idx], Mapping) else {}
        out_rows.append(prefix + [record.get(field) for field in field_names])
    return out_cols, out_rows


def _expand_list_column(columns: list[dict[str, Any]], rows: list[list[Any]], expression: str) -> tuple[list[dict[str, Any]], list[list[Any]]]:
    names = _string_literals_after_first_arg(expression)
    if not names:
        raise ValueError("Table.ExpandListColumn preview needs a list column name.")
    list_idx = _find_column_index(columns, names[0])
    out_cols = [dict(col) for col in columns]
    out_rows: list[list[Any]] = []
    for row in rows:
        value = row[list_idx] if list_idx < len(row) else None
        values = value if isinstance(value, list) else [value]
        if not values:
            out_rows.append(list(row))
            continue
        for item in values:
            out_row = list(row)
            if list_idx < len(out_row):
                out_row[list_idx] = item
            out_rows.append(out_row)
    return out_cols, out_rows


def _first_number_arg(expression: str) -> int | None:
    numbers = re.findall(r",\s*(\d+)\s*(?:\)|,)", expression)
    if not numbers:
        return None
    return int(numbers[0])


def _range_args(expression: str) -> tuple[int, int] | None:
    match = re.search(r"Table\.Range\([^,]+,\s*(\d+)\s*,\s*(\d+)\s*\)", expression, flags=re.IGNORECASE | re.DOTALL)
    if not match:
        return None
    return int(match.group(1)), int(match.group(2))


def _row_offset_count_args(expression: str, *, default_count: int | None) -> tuple[int, int | None] | None:
    args = _function_args(expression)
    if len(args) < 2:
        return None
    try:
        offset = max(0, int(_to_number(_parse_m_value(args[1])) or 0))
        count = default_count if len(args) <= 2 or args[2].strip().lower() == "null" else max(0, int(_to_number(_parse_m_value(args[2])) or 0))
    except Exception:
        return None
    return offset, count


def _alternate_args(expression: str) -> tuple[int, int, int] | None:
    args = _function_args(expression)
    if len(args) < 4:
        return None
    try:
        return (
            max(0, int(_to_number(_parse_m_value(args[1])) or 0)),
            max(0, int(_to_number(_parse_m_value(args[2])) or 0)),
            max(0, int(_to_number(_parse_m_value(args[3])) or 0)),
        )
    except Exception:
        return None


def _alternate_sequence(values: Sequence[Any], offset: int, skip: int, take: int) -> list[Any]:
    kept = list(values[:offset])
    idx = offset
    while idx < len(values):
        idx += skip
        if take <= 0:
            if skip <= 0:
                break
            continue
        kept.extend(values[idx : idx + take])
        idx += take
        if skip <= 0 and take <= 0:
            break
    return kept


def _replace_rows_from_records(
    columns: list[dict[str, Any]],
    rows: list[list[Any]],
    expression: str,
    *,
    variables: Mapping[str, Any] | None = None,
) -> tuple[list[dict[str, Any]], list[list[Any]]]:
    args = _function_args(expression)
    if len(args) < 4:
        raise ValueError("Table.ReplaceRows preview needs offset, count, and replacement rows.")
    offset = max(0, int(_to_number(_parse_m_value(args[1])) or 0))
    count = max(0, int(_to_number(_parse_m_value(args[2])) or 0))
    replacements = _eval_m_expression([], [], args[3], variables=variables)
    replacement_records = [record for record in replacements if isinstance(record, Mapping)] if isinstance(replacements, list) else []
    out_cols = [dict(col) for col in columns]
    existing = {str(col.get("name") or "").upper() for col in out_cols}
    for record in replacement_records:
        for key in record.keys():
            key_text = str(key)
            if key_text.upper() not in existing:
                existing.add(key_text.upper())
                out_cols.append({"name": key_text})
    record_rows = [_row_record(columns, row) for row in rows]
    out_records = record_rows[:offset] + [dict(record) for record in replacement_records] + record_rows[offset + count :]
    out_rows: list[list[Any]] = []
    for record in out_records:
        out_rows.append([
            _mapping_get_case_insensitive(record, str(col.get("name") or ""))
            if any(str(key).upper() == str(col.get("name") or "").upper() for key in record.keys())
            else None
            for col in out_cols
        ])
    return out_cols, out_rows


def _is_scalar_helper_operation(op: str) -> bool:
    upper = (op or "").strip().upper()
    if upper in {"EXPRESSION", "", "#BINARY", "#DATE", "#DATETIME", "#DATETIMEZONE", "#TIME", "#DURATION"}:
        return True
    if upper in {
        "TABLE.COLUMN",
        "TABLE.COLUMNNAMES",
        "TABLE.ROWCOUNT",
        "TABLE.APPROXIMATEROWCOUNT",
        "TABLE.HASCOLUMNS",
        "TABLE.ISDISTINCT",
        "TABLE.FIRSTVALUE",
        "TABLE.SINGLEROW",
        "TABLE.PREFIXCOLUMNS",
        "TABLE.FROMVALUE",
        "TABLE.VIEW",
        "TABLE.SPLIT",
        "TABLE.SPLITAT",
        "TABLE.PARTITION",
        "TABLES.GETRELATIONSHIPS",
        "TABLE.VIEWERROR",
        "TABLE.VIEWFUNCTION",
        "TABLE.KEYS",
        "TABLE.REPLACEKEYS",
        "TABLE.PARTIONKEY",
        "TABLE.REPLACEPARTITIONKEY",
        "TABLE.REPLACERELATIONSHIPIDENTITY",
        "TABLE.PARTITIONVALUES",
        "TABLE.CONFORMTOPAGEREADER",
        "TABLE.WITHERRORCONTEXT",
        "TABLE.MATCHESALLROWS",
        "TABLE.MATCHESANYROWS",
        "TABLE.COLUMNCOUNT",
        "TABLE.COLUMNSOFTYPE",
        "TABLE.SCHEMA",
        "TABLE.PROFILE",
        "TABLE.ISEMPTY",
        "TABLE.CONTAINS",
        "TABLE.CONTAINSANY",
        "TABLE.CONTAINSALL",
        "TABLE.POSITIONOF",
        "TABLE.POSITIONOFANY",
        "TABLE.FIRST",
        "TABLE.LAST",
        "TABLE.TORECORDS",
        "TABLE.TOROWS",
        "TABLE.TOCOLUMNS",
        "TABLE.TRANSFORMROWS",
        "TABLE.TOLIST",
        "TABLE.POSITIONOF",
        "TABLE.POSITIONOFANY",
        "ACCESSCONTROLENTRY.CONDITIONTOIDENTITIES",
        "HTML.TABLE",
        "IDENTITY.FROM",
        "IDENTITY.ISMEMBEROF",
        "IDENTITYPROVIDER.DEFAULT",
        "JSON.FROMVALUE",
        "ODBC.INFEROPTIONS",
        "WEB.HEADERS",
        "XML.DOCUMENT",
    }:
        return True
    return upper.startswith((
        "LIST.",
        "RECORD.",
        "DATE.",
        "DATETIME.",
        "DATETIMEZONE.",
        "DURATION.",
        "TIME.",
        "NUMBER.",
        "TEXT.",
        "GUID.",
        "LOGICAL.",
        "VALUE.",
        "TYPE.",
        "ERROR.",
        "CHARACTER.",
        "LINES.",
        "URI.",
        "BINARY.",
        "BINARYFORMAT.",
        "BYTE.",
        "INT8.",
        "INT16.",
        "INT32.",
        "INT64.",
        "CURRENCY.",
        "DECIMAL.",
        "DOUBLE.",
        "SINGLE.",
        "PERCENTAGE.",
        "COMPARER.",
        "REPLACER.",
        "SPLITTER.",
        "COMBINER.",
        "ACTION.",
        "DIAGNOSTICS.",
        "DIRECTQUERYCAPABILITIES.",
        "EMBEDDED.",
        "EXCEL.",
        "FUNCTION.",
        "ITEMEXPRESSION.",
        "ROWEXPRESSION.",
        "GEOGRAPHY.",
        "GEOGRAPHYPOINT.",
        "GEOMETRY.",
        "GEOMETRYPOINT.",
        "GRAPH.",
        "MODULE.",
        "PROGRESS.",
        "SQLEXPRESSION.",
        "VARIABLE.",
    ))


def evaluate_preview_table(
    *,
    raw_m: str,
    columns: Sequence[Mapping[str, Any]],
    rows: Sequence[Sequence[Any]],
    source_sql: str = "",
    target_step_id: str | None = None,
    helper_queries: Mapping[str, str] | None = None,
    initial_variables: Mapping[str, Any] | None = None,
    table_variables: Mapping[str, tuple[Sequence[Mapping[str, Any]], Sequence[Sequence[Any]]]] | None = None,
    project_root: str | Path | None = None,
) -> PreviewEvaluationResult:
    """Apply a safe subset of M table operations over preview rows."""

    current_columns = [dict(col) for col in columns]
    current_rows = [list(row) for row in rows]
    steps, result_expression, _functions, diagnostics = parse_m_query(raw_m)
    env: dict[str, tuple[list[dict[str, Any]], list[list[Any]]]] = {}
    value_env: dict[str, Any] = dict(initial_variables or {})
    if project_root is not None:
        value_env[_STATIC_FILE_BASE_PATH_KEY] = str(project_root)
    for helper_name, helper_raw_m in (helper_queries or {}).items():
        name = str(helper_name or "").strip()
        if not name or name in value_env:
            continue
        helper_text = str(helper_raw_m or "").strip()
        function_value = _parse_m_function_value(helper_text)
        if function_value is not None:
            value_env[name] = function_value
            continue
        expr = _strip_top_level_meta(helper_text)
        if re.match(r"^let\b", expr, flags=re.IGNORECASE):
            value_env[name] = MQueryValue(expr)
            continue
        if expr and not re.match(r"^let\b", expr, flags=re.IGNORECASE):
            try:
                value_env[name] = _eval_m_expression([], [], expr, variables=value_env)
            except Exception:
                value_env[name] = MQueryValue(expr)
    table_env: dict[str, tuple[list[dict[str, Any]], list[list[Any]]]] = {}
    for table_name, table_value in (table_variables or {}).items():
        name = _m_identifier_name(str(table_name or "").strip())
        if not name:
            continue
        table_env[name.upper()] = (
            [dict(col) for col in table_value[0]],
            [list(row) for row in table_value[1]],
        )
    applied: list[str] = []
    blocked: list[dict[str, str]] = []
    target_key = (target_step_id or "").strip().upper()
    result_step_id: str | None = None

    for step in steps:
        op = str(step.operation or "").strip().upper()
        produces_table = True
        if step.dependencies and op not in {"TABLE.COMBINE", "TABLE.JOIN", "TABLE.NESTEDJOIN"}:
            dep = step.dependencies[-1]
            if dep in env:
                current_columns, current_rows = ([dict(col) for col in env[dep][0]], [list(row) for row in env[dep][1]])
        try:
            table_ref = table_env.get(_m_identifier_name(step.expression.strip()).upper())
            if table_ref is not None:
                current_columns, current_rows = ([dict(col) for col in table_ref[0]], [list(row) for row in table_ref[1]])
                applied.append(step.id)
            elif op == "FILE.CONTENTS":
                value_env[step.id] = _eval_m_expression([], [], step.expression, variables=value_env)
                produces_table = False
                if target_key and step.id.upper() == target_key:
                    current_columns, current_rows = _value_preview_table(value_env[step.id])
            elif op in {"FOLDER.FILES", "FOLDER.CONTENTS"}:
                pass
            elif op in {"CSV.DOCUMENT", "JSON.DOCUMENT", "XML.TABLES"}:
                try:
                    value = _eval_m_expression([], [], step.expression, variables=value_env)
                    current_columns, current_rows = _value_preview_table(value)
                    applied.append(step.id)
                except Exception:
                    if current_columns or current_rows:
                        pass
                    else:
                        raise
            elif op == "EXCEL.WORKBOOK":
                try:
                    value = _eval_m_expression([], [], step.expression, variables=value_env)
                    current_columns, current_rows = _value_preview_table(value)
                    applied.append(step.id)
                except Exception:
                    if current_columns or current_rows:
                        pass
                    else:
                        raise
            elif op in {"XML.TABLES", "PARQUET.DOCUMENT"}:
                pass
            elif step.operation == "Navigation":
                try:
                    value = _eval_m_expression([], [], step.expression, variables=value_env)
                    if isinstance(value, list):
                        current_columns, current_rows = _value_preview_table(value)
                    else:
                        value_env[step.id] = value
                        produces_table = False
                        if target_key and step.id.upper() == target_key:
                            current_columns, current_rows = _value_preview_table(value)
                    applied.append(step.id)
                except Exception:
                    # Source navigation against external connector handles remains
                    # preserve-only until a connector adapter supplies table values.
                    pass
            elif op in {"#TABLE", "TABLE.FROMRECORDS", "TABLE.FROMROWS", "TABLE.FROMCOLUMNS", "TABLE.FROMLIST"}:
                current_columns, current_rows = _constant_table(step.expression, variables=value_env)
                applied.append(step.id)
            elif op == "TABLE.TRANSFORMCOLUMNTYPES":
                current_columns, current_rows = _transform_column_types(current_columns, current_rows, step.expression)
                applied.append(step.id)
            elif op == "TABLE.TRANSFORMCOLUMNS":
                current_columns, current_rows = _transform_columns(current_columns, current_rows, step.expression, variables=value_env)
                applied.append(step.id)
            elif op == "TABLE.TRANSFORMCOLUMNNAMES":
                current_columns, current_rows = _transform_column_names(current_columns, current_rows, step.expression, variables=value_env)
                applied.append(step.id)
            elif op == "TABLE.SELECTROWS":
                current_columns, current_rows = _select_rows(current_columns, current_rows, step.expression, variables=value_env)
                applied.append(step.id)
            elif op == "TABLE.SELECTCOLUMNS":
                current_columns, current_rows = _select_columns(
                    current_columns,
                    current_rows,
                    _string_literals_after_first_arg(step.expression),
                    missing_field=_missing_field_mode(step.expression),
                )
                applied.append(step.id)
            elif op == "TABLE.REMOVECOLUMNS":
                current_columns, current_rows = _remove_columns(
                    current_columns,
                    current_rows,
                    _string_literals_after_first_arg(step.expression),
                    missing_field=_missing_field_mode(step.expression),
                )
                applied.append(step.id)
            elif op == "TABLE.RENAMECOLUMNS":
                current_columns, current_rows = _rename_columns(current_columns, current_rows, step.expression)
                applied.append(step.id)
            elif op == "TABLE.DUPLICATECOLUMN":
                current_columns, current_rows = _duplicate_column(current_columns, current_rows, step.expression)
                applied.append(step.id)
            elif op == "TABLE.REORDERCOLUMNS":
                current_columns, current_rows = _reorder_columns(current_columns, current_rows, step.expression)
                applied.append(step.id)
            elif op == "TABLE.DISTINCT":
                current_columns, current_rows = _distinct_rows(current_columns, current_rows)
                applied.append(step.id)
            elif op == "TABLE.SORT":
                current_columns, current_rows = _sort_rows(current_columns, current_rows, step.expression)
                applied.append(step.id)
            elif op == "TABLE.FINDTEXT":
                value = _eval_m_expression([], [], step.expression, variables=value_env)
                current_columns, current_rows = _value_preview_table(value)
                applied.append(step.id)
            elif op == "TABLE.FIRSTN":
                n = _first_number_arg(step.expression)
                if n is None:
                    raise ValueError("Table.FirstN preview supports numeric row counts only.")
                current_rows = current_rows[:n]
                applied.append(step.id)
            elif op == "TABLE.LASTN":
                n = _first_number_arg(step.expression)
                if n is None:
                    raise ValueError("Table.LastN preview supports numeric row counts only.")
                current_rows = current_rows[-n:] if n > 0 else []
                applied.append(step.id)
            elif op in {"TABLE.SKIP", "TABLE.REMOVEFIRSTN"}:
                n = _first_number_arg(step.expression)
                if n is None:
                    raise ValueError(f"{step.operation} preview supports numeric row counts only.")
                current_rows = current_rows[n:]
                applied.append(step.id)
            elif op == "TABLE.REMOVELASTN":
                n = _first_number_arg(step.expression)
                if n is None:
                    raise ValueError("Table.RemoveLastN preview supports numeric row counts only.")
                current_rows = current_rows[:-n] if n > 0 else current_rows
                applied.append(step.id)
            elif op == "TABLE.RANGE":
                args = _range_args(step.expression)
                if args is None:
                    raise ValueError("Table.Range preview supports numeric offset/count only.")
                start, count = args
                current_rows = current_rows[start : start + count]
                applied.append(step.id)
            elif op == "TABLE.REMOVEROWS":
                args = _row_offset_count_args(step.expression, default_count=1)
                if args is None:
                    raise ValueError("Table.RemoveRows preview supports numeric offset/count only.")
                offset, count = args
                count = 1 if count is None else count
                current_rows = current_rows[:offset] + current_rows[offset + count :]
                applied.append(step.id)
            elif op == "TABLE.ALTERNATEROWS":
                args = _alternate_args(step.expression)
                if args is None:
                    raise ValueError("Table.AlternateRows preview supports numeric offset/skip/take only.")
                offset, skip, take = args
                current_rows = _alternate_sequence(current_rows, offset, skip, take)
                applied.append(step.id)
            elif op == "TABLE.REPLACEROWS":
                current_columns, current_rows = _replace_rows_from_records(current_columns, current_rows, step.expression, variables=value_env)
                applied.append(step.id)
            elif op in {"TABLE.INSERTROWS", "TABLE.REMOVEMATCHINGROWS", "TABLE.REPLACEMATCHINGROWS", "TABLE.REVERSEROWS"}:
                value = _eval_m_expression([], [], step.expression, variables=value_env)
                current_columns, current_rows = _value_preview_table(value)
                applied.append(step.id)
            elif op in {"TABLE.REPEAT", "TABLE.TRANSPOSE"}:
                value = _eval_m_expression([], [], step.expression, variables=value_env)
                current_columns, current_rows = _value_preview_table(value)
                applied.append(step.id)
            elif op == "TABLE.ADDCOLUMN":
                current_columns, current_rows = _add_column(current_columns, current_rows, step.expression, variables=value_env)
                applied.append(step.id)
            elif op == "TABLE.ADDINDEXCOLUMN":
                current_columns, current_rows = _add_index_column(current_columns, current_rows, step.expression)
                applied.append(step.id)
            elif op == "TABLE.REPLACEVALUE":
                current_rows = _replace_values(current_columns, current_rows, step.expression)
                applied.append(step.id)
            elif op == "TABLE.REPLACEERRORVALUES":
                current_rows = _replace_error_values(current_columns, current_rows, step.expression)
                applied.append(step.id)
            elif op == "TABLE.REMOVEROWSWITHERRORS":
                current_rows = _rows_with_errors(current_columns, current_rows, step.expression, keep_errors=False)
                applied.append(step.id)
            elif op == "TABLE.SELECTROWSWITHERRORS":
                current_rows = _rows_with_errors(current_columns, current_rows, step.expression, keep_errors=True)
                applied.append(step.id)
            elif op == "TABLE.FILLDOWN":
                current_rows = _fill_down(current_columns, current_rows, _string_literals_after_first_arg(step.expression))
                applied.append(step.id)
            elif op == "TABLE.FILLUP":
                current_rows = _fill_up(current_columns, current_rows, _string_literals_after_first_arg(step.expression))
                applied.append(step.id)
            elif op == "TABLE.PROMOTEHEADERS":
                try:
                    value = _eval_m_expression([], [], step.expression, variables=value_env)
                    current_columns, current_rows = _value_preview_table(value)
                except Exception:
                    current_columns, current_rows = _promote_headers(current_columns, current_rows)
                applied.append(step.id)
            elif op == "TABLE.DEMOTEHEADERS":
                current_columns, current_rows = _demote_headers(current_columns, current_rows)
                applied.append(step.id)
            elif op == "TABLE.SPLITCOLUMN":
                current_columns, current_rows = _split_column(current_columns, current_rows, step.expression)
                applied.append(step.id)
            elif op == "TABLE.COMBINECOLUMNS":
                current_columns, current_rows = _combine_columns(current_columns, current_rows, step.expression)
                applied.append(step.id)
            elif op == "TABLE.ADDFUZZYCLUSTERCOLUMN":
                current_columns, current_rows = _add_fuzzy_cluster_column(current_columns, current_rows, step.expression)
                applied.append(step.id)
            elif op == "TABLE.FUZZYGROUP":
                current_columns, current_rows = _fuzzy_group_table(current_columns, current_rows, step.expression)
                applied.append(step.id)
            elif op == "TABLE.FILTERWITHDATATABLE":
                current_columns, current_rows = _filter_with_data_table(current_columns, current_rows, step.expression)
                applied.append(step.id)
            elif op == "TABLE.FUZZYJOIN":
                current_columns, current_rows = _fuzzy_join_tables(*_join_tables_from_expression(step.dependencies, env, step.expression), step.expression)
                applied.append(step.id)
            elif op == "TABLE.FUZZYNESTEDJOIN":
                current_columns, current_rows = _fuzzy_nested_join_tables(*_join_tables_from_expression(step.dependencies, env, step.expression), step.expression)
                applied.append(step.id)
            elif op in {"TABLE.ADDJOINCOLUMN", "TABLE.ADDRANKCOLUMN", "TABLE.COMBINECOLUMNSTORECORD", "TABLE.FROMPARTITIONS"}:
                value = _eval_m_expression([], [], step.expression, variables=value_env)
                current_columns, current_rows = _value_preview_table(value)
                applied.append(step.id)
            elif op == "TABLE.ADDKEY":
                args = _function_args(step.expression)
                source_columns, source_rows = current_columns, current_rows
                dep_table = _table_env_lookup(env, args[0]) if args else None
                if dep_table is not None:
                    source_columns, source_rows = dep_table
                current_columns, current_rows = (
                    _with_table_key_metadata(
                        source_columns,
                        _table_key_metadata(source_columns) + _table_key_records_from_args(
                            _eval_m_expression([], [], args[1], variables=value_env) if len(args) > 1 else [],
                            _eval_m_expression([], [], args[2], variables=value_env) if len(args) > 2 else False,
                        ),
                    ),
                    [list(row) for row in source_rows],
                )
                applied.append(step.id)
            elif op == "TABLE.REPLACEKEYS":
                args = _function_args(step.expression)
                source_columns, source_rows = current_columns, current_rows
                dep_table = _table_env_lookup(env, args[0]) if args else None
                if dep_table is not None:
                    source_columns, source_rows = dep_table
                keys = _eval_m_expression([], [], args[1], variables=value_env) if len(args) > 1 else []
                current_columns, current_rows = (
                    _with_table_key_metadata(source_columns, keys if isinstance(keys, list) else []),
                    [list(row) for row in source_rows],
                )
                applied.append(step.id)
            elif op == "TABLE.KEYS":
                args = _function_args(step.expression)
                source_columns = current_columns
                dep_table = _table_env_lookup(env, args[0]) if args else None
                if dep_table is not None:
                    source_columns = dep_table[0]
                current_columns, current_rows = _value_preview_table(_table_key_records_from_columns(source_columns))
                applied.append(step.id)
            elif op in {"TABLE.BUFFER", "TABLE.STOPFOLDING"}:
                applied.append(step.id)
            elif op == "VALUE.REPLACETYPE":
                current_columns, current_rows = _apply_value_replace_type(current_columns, current_rows, step.expression)
                applied.append(step.id)
            elif op == "TABLE.GROUP":
                current_columns, current_rows = _group_table(current_columns, current_rows, step.expression)
                applied.append(step.id)
            elif op == "TABLE.COMBINE":
                source_tables = [
                    env[dep] for dep in step.dependencies
                    if dep in env
                ]
                if not source_tables:
                    source_tables = [(current_columns, current_rows)]
                current_columns, current_rows = _combine_tables(source_tables)
                applied.append(step.id)
            elif op == "TABLE.UNPIVOTOTHERCOLUMNS":
                current_columns, current_rows = _unpivot_other_columns(current_columns, current_rows, step.expression)
                applied.append(step.id)
            elif op == "TABLE.UNPIVOT":
                current_columns, current_rows = _unpivot_columns(current_columns, current_rows, step.expression)
                applied.append(step.id)
            elif op == "TABLE.PIVOT":
                current_columns, current_rows = _pivot_table(current_columns, current_rows, step.expression)
                applied.append(step.id)
            elif op == "TABLE.JOIN":
                current_columns, current_rows = _join_tables(*_join_tables_from_expression(step.dependencies, env, step.expression), step.expression)
                applied.append(step.id)
            elif op == "TABLE.NESTEDJOIN":
                current_columns, current_rows = _nested_join_tables(*_join_tables_from_expression(step.dependencies, env, step.expression), step.expression)
                applied.append(step.id)
            elif op == "TABLE.AGGREGATETABLECOLUMN":
                current_columns, current_rows = _aggregate_table_column(current_columns, current_rows, step.expression)
                applied.append(step.id)
            elif op == "TABLE.EXPANDTABLECOLUMN":
                current_columns, current_rows = _expand_table_column(current_columns, current_rows, step.expression, variables=value_env)
                applied.append(step.id)
            elif op == "TABLE.EXPANDRECORDCOLUMN":
                current_columns, current_rows = _expand_record_column(current_columns, current_rows, step.expression)
                applied.append(step.id)
            elif op == "TABLE.EXPANDLISTCOLUMN":
                current_columns, current_rows = _expand_list_column(current_columns, current_rows, step.expression)
                applied.append(step.id)
            elif _is_scalar_helper_operation(op) or not step.functions:
                value_env[step.id] = _eval_m_expression([], [], step.expression, variables=value_env)
                applied.append(step.id)
                produces_table = False
                if target_key and step.id.upper() == target_key:
                    current_columns, current_rows = _value_preview_table(value_env[step.id])
            else:
                blocked.append({"step_id": step.id, "operation": step.operation, "reason": "Preview evaluator does not support this M operation yet."})
        except Exception as exc:
            blocked.append({"step_id": step.id, "operation": step.operation, "reason": str(exc)})
        if produces_table:
            env[step.id] = ([dict(col) for col in current_columns], [list(row) for row in current_rows])
            value_env[step.id] = _table_records(current_columns, current_rows)
        if target_key and step.id.upper() == target_key:
            result_step_id = step.id
            break

    result_id = (result_expression or "").strip()
    if result_id.startswith('#"') and result_id.endswith('"'):
        result_id = result_id[2:-1].replace('""', '"')
    if target_key:
        pass
    elif result_id in env:
        current_columns, current_rows = env[result_id]
    elif result_id in value_env:
        current_columns, current_rows = _value_preview_table(value_env[result_id])

    return PreviewEvaluationResult(
        columns=current_columns,
        rows=current_rows,
        applied_steps=applied,
        blocked_steps=blocked,
        diagnostics=diagnostics,
        result_step_id=result_step_id,
    )
